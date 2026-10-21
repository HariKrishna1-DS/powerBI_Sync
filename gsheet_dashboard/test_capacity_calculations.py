from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
from capacity_calculations import capacity_date, capacity_days, capacity_details
from report_workspace import capacity_report, imported_snapshot, report_frames, register_report_routes, ReportWorkspace
from preview_store import PreviewStore
from tracker_sync import sheet_reports, FULL


def order(number, status='QC in Progress', received='10/01/2026 09:00 AM', completed=''):
    return {'Order Number':number, 'Product':'Full Title', 'Status':status,
            'In-Time':received, 'Out Time':completed, 'SLA Expiration':'10/20/2026 09:00 AM'}


class CapacityCalculationTests(unittest.TestCase):
    def setUp(self):
        clock = patch('report_workspace.current_capacity_month', return_value='2026-10')
        clock.start()
        self.addCleanup(clock.stop)

    def test_summary_only_uses_exact_daily_counts_without_carryover(self):
        reports={'daily':[
            {'Date':'2026-10-05','Received':31,'Completed':22,'Clarification':3,'Vendor Pending':1},
            {'Date':'2026-10-02','Received':414,'Completed':297,'Clarification':50,'Vendor Pending':30},
            {'Date':'2026-10-01','Received':323,'Completed':275,'Clarification':11,'Vendor Pending':11}], 'monthly':[]}
        saved=deepcopy(reports)
        result=capacity_report(reports, {'mode':'excel'})
        self.assertEqual([row['Received'] for row in result['daily']], [323,414,31])
        self.assertEqual(result['daily'][1]['Rollover Received'],0)
        self.assertEqual(result['monthly'][0]['Received'],768)
        self.assertEqual(result['monthly'][0]['Clarification'],64)
        self.assertEqual(result['monthly'][0]['Vendor Pending'],42)
        self.assertEqual([row['Completed'] for row in result['daily']], [275,297,22])
        self.assertEqual(reports,saved)
        self.assertEqual(capacity_report(reports, {'mode':'excel'}),result)

    def test_previous_month_pending_orders_do_not_increase_current_receipts(self):
        reports=sheet_reports({FULL:[order('A','Awaiting for Clarification',received='09/30/2026 09:00 AM'),
                                   order('B','Assign to ABS',received='09/30/2026 09:00 AM'),
                                   order('C',received='10/02/2026 09:00 AM')]})
        result=capacity_report(reports, {'selected_date':'2026-10-02'})
        self.assertEqual(result['daily'][0]['Received'],1)
        self.assertEqual(result['daily'][0]['Raw Received'],1)
        self.assertEqual(result['monthly'][1]['Received'],1)

    def test_in_house_order_counts_on_actual_completion_day_once(self):
        pending=order('P')
        before=capacity_report(sheet_reports({FULL:[pending]}), {'mode':'tracker'})
        self.assertEqual(before['daily'][0]['In-House Pending'],1)
        completed=dict(pending, Status='Completed and Delivered', **{'Out Time':'10/03/2026 01:30 PM'})
        reports=sheet_reports({FULL:[completed,order('Q','Cancelled',completed='10/03/2026 02:00 PM')]})
        result=capacity_report(reports, {'mode':'tracker'})
        self.assertEqual([(r['Date'],r['Received'],r['Completed']) for r in result['daily']],
                         [('2026-10-01',2,0),('2026-10-02',0,2)])
        self.assertEqual(result['monthly'][0]['Completed'],2)
        self.assertEqual(result['monthly'][0]['Capacity'],1400)
        self.assertEqual(reports['daily'][0]['Completed'],2)

    def test_existing_completion_day_moves_counts_without_double_counting(self):
        reports=sheet_reports({FULL:[order('A','Completed and Delivered',completed='10/02/2026 10:00 AM'),
                                   order('B','Completed and Delivered',received='10/02/2026 09:00 AM',completed='10/02/2026 11:00 AM'),
                                   order('C','Search In Progress',completed='10/02/2026 12:00 PM')]})
        result=capacity_report(reports, {'mode':'tracker'})
        self.assertEqual([r['Completed'] for r in result['daily']], [0,3])
        self.assertEqual(sum(r['Completed'] for r in result['daily']),3)

    def test_completion_date_only_rows_serials_and_missing_date_fallback(self):
        rows=[order('A','Completed and Delivered',completed=str((datetime(2026,10,4,12)-datetime(1899,12,30)).total_seconds()/86400)),
              order('B','Completed and Delivered',completed='invalid'),
              order('C','Completed and Delivered')]
        result=capacity_report(sheet_reports({FULL:rows}), {'mode':'tracker'})
        self.assertEqual([(r['Date'],r['Completed']) for r in result['daily']], [('2026-10-01',2),('2026-10-02',1)])
        self.assertEqual(result['missing_completion_date_count'],2)

    def test_imported_completion_in_next_year_stays_on_report_date_in_pr_excel(self):
        rows=[order('A','Completed and Delivered',received='12/31/2026 09:00 AM',completed='01/02/2027 01:00 PM')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[], 'created':'2026-12-31'})
        frames,report=report_frames(snapshot, {'mode':'excel','capacity_month':'2027-01'})
        self.assertEqual([(r['Date'],r['Received'],r['Completed']) for r in report['monthly']],
                         [('2026-12',1,1)])
        self.assertEqual(report['month'],'2027-01')
        self.assertEqual(frames['PR Excel'][2][:3],['2027 YTD Total',0,0])
        self.assertEqual(report['daily'],[])
        self.assertEqual(report['missing_completion_date_count'],0)

    def test_live_daily_table_uses_current_month_and_keeps_previous_month_totals(self):
        reports=sheet_reports({FULL:[order('A','Awaiting for Clarification',received='08/20/2026 09:00 AM'),
                                   order('B','Assign to ABS',received='09/30/2026 09:00 AM'),
                                   order('C',received='10/01/2026 09:00 AM'),
                                   order('D',received='10/02/2026 09:00 AM')]})
        for mode in ('excel','tracker'):
            report=capacity_report(reports, {'mode':mode,'selected_date':'2026-09-30'})
            self.assertEqual(report['month'],'2026-10')
            self.assertEqual(report['scope_label'],'October 2026')
            self.assertEqual([row['Date'] for row in report['daily']],['2026-10-01','2026-10-02'])
            self.assertEqual([row['Date'] for row in report['monthly']],['2026-08','2026-09','2026-10'])
            self.assertEqual(report['daily'][0]['Rollover Received'],0)
            self.assertEqual(report['monthly'][0]['Received'],1)

    def test_sheet_frames_show_old_months_only_in_summary_and_october_days_in_detail(self):
        rows=[order('A',received='08/20/2026 09:00 AM'),order('B',received='09/30/2026 09:00 AM'),
              order('C',received='10/01/2026 09:00 AM'),order('D',received='10/02/2026 09:00 AM')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[],'created':'2026-10-15'})
        frames,report=report_frames(snapshot,{'mode':'excel','selected_date':'2026-09-30'})
        values=frames['PR Excel']
        self.assertEqual([row[0] for row in values[1:4]],['2026-08','2026-09','2026-10'])
        self.assertEqual([row[0] for row in values[6:-1]],['2026-10-01','2026-10-02'])
        self.assertEqual(values[-1][1],2)
        self.assertEqual(values[4][1],4)

    def test_default_excel_has_only_october_details_and_explicit_month_export_is_preserved(self):
        rows=[order('A',received='09/30/2026 09:00 AM'),order('B',received='10/01/2026 09:00 AM')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[],'created':'2026-10-15'})
        with tempfile.TemporaryDirectory() as folder:
            workspace=ReportWorkspace(PreviewStore(Path(folder)/'previews'))
            workspace.save(mode='excel',selected_date='2026-09-30')
            app=Flask(__name__)
            register_report_routes(app,workspace,threading.Lock(),lambda **kwargs:snapshot,lambda:None,lambda:None)
            client=app.test_client()
            for query,expected in (('', '2026-10-01'),('&month=2026-09', '2026-09-30')):
                response=client.get('/api/export/report?sheet=PR%20Excel'+query)
                self.assertEqual(response.status_code,200)
                book=load_workbook(BytesIO(response.data),data_only=True)
                sheet=book['PR Excel']
                self.assertEqual(sheet['A6'].value.date().isoformat(),expected)
                self.assertEqual(sheet['A6'].number_format,'dd-mm-yy')

    def test_completion_duplicate_archive_copy_counted_once(self):
        completed=order('A','Completed and Delivered',completed='10/02/2026 10:00 AM')
        reports={'daily':[{'Date':'2026-10-01','Received':2,'Completed':1,'Clarification':0,'Vendor Pending':0,
            'rows':[dict(order('A'),_month='2026-09'),dict(completed,_month='2026-10')]}], 'monthly':[]}
        days,_=capacity_days(reports)
        self.assertEqual(sum(day['Completed'] for day in days),1)
        self.assertEqual(days[-1]['Date'],'2026-10-02')

    def test_weekend_receipts_completions_and_pending_roll_into_friday_once(self):
        rows=[order('Friday clarification','Awaiting for Clarification',received='10/02/2026'),
              order('Friday vendor','Assign to ABS',received='10/02/2026'),
              order('Saturday','Completed and Delivered',received='10/03/2026',completed='10/03/2026'),
              order('Sunday','Awaiting for Clarification',received='10/04/2026'),
              order('Monday','Completed and Delivered',received='10/05/2026',completed='10/04/2026'),
              order('Tuesday',received='10/06/2026')]
        reports=sheet_reports({FULL:rows})
        original=deepcopy(reports)
        report=capacity_report(reports, {'mode':'tracker'})
        self.assertEqual([(day['Date'],day['Raw Received'],day['Received'],day['Completed']) for day in report['daily']],
                         [('2026-10-02',4,4,2),('2026-10-05',1,1,0),('2026-10-06',1,1,0)])
        self.assertEqual(report['monthly'][0]['Capacity'],2100)
        friday=capacity_details(reports,'2026-10-02')
        groups={group['key']:group['rows'] for group in friday['groups']}
        self.assertEqual(len(groups['received']),4)
        self.assertEqual(len(groups['completed']),2)
        self.assertEqual({row['Source received date'] for row in groups['received']}, {'2026-10-02','2026-10-03','2026-10-04'})
        self.assertEqual({row['Capacity date'] for row in groups['received']}, {'2026-10-02'})
        monday=capacity_details(reports,'2026-10-05')
        self.assertEqual(monday['rollover_received'],0)
        self.assertEqual(len(next(group['rows'] for group in monday['groups'] if group['key']=='received')),1)
        self.assertEqual(reports,original)
        self.assertIn('2026-10-03',[day['Date'] for day in reports['daily']])

    def test_weekend_month_start_stays_in_previous_friday_month_in_summary_detail_and_export(self):
        rows=[order('Friday',received='10/30/2026'),
              order('Saturday','Completed and Delivered',received='10/31/2026',completed='11/01/2026'),
              order('Sunday',received='11/01/2026')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[],'created':'2026-10-31'})
        frames,report=report_frames(snapshot, {'mode':'excel','capacity_month':'2026-10'})
        self.assertEqual([(day['Date'],day['Received'],day['Completed'],day['Capacity']) for day in report['monthly']],
                         [('2026-10',3,1,700)])
        self.assertEqual([day['Date'] for day in report['daily']], ['2026-10-30'])
        self.assertEqual(frames['PR Excel'][4][:3],['2026-10-30',3,1])
        groups={g['key']:g['rows'] for g in capacity_details(snapshot['reports'],'2026-10-30',use_report_date=True)['groups']}
        self.assertEqual({row['In-Time'] for row in groups['received']},{'10/30/2026','10/31/2026','11/01/2026'})
        self.assertEqual(capacity_report(snapshot['reports'],{'mode':'excel','capacity_month':'2026-11'})['daily'],[])

    def test_weekend_mapping_handles_month_and_year_boundaries(self):
        for original,friday in [('2026-10-03','2026-10-02'),('2026-10-04','2026-10-02'),
                                ('2026-02-01','2026-01-30'),('2022-01-01','2021-12-31'),
                                ('2022-01-02','2021-12-31'),('2026-10-05','2026-10-05')]:
            with self.subTest(original=original):
                self.assertEqual(capacity_date(original),friday)

    def test_api_and_export_use_same_direct_receipt_and_completion_counts(self):
        rows=[order('A','Awaiting for Clarification'),order('B','Assign to ABS'),
              order('C','Completed and Delivered',completed='10/02/2026 10:00 AM'),
              order('D',received='10/02/2026 09:00 AM')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[], 'created':'2026-10-02'})
        with tempfile.TemporaryDirectory() as folder:
            workspace=ReportWorkspace(PreviewStore(Path(folder)/'previews'))
            workspace.save(mode='excel')
            app=Flask(__name__)
            register_report_routes(app,workspace,threading.Lock(),lambda **kwargs:snapshot,lambda:None,lambda:None)
            client=app.test_client()
            data=client.get('/api/capacity-report').json
            self.assertEqual([(r['Received'],r['Completed']) for r in data['daily']],[(3,1),(1,0)])
            details=client.get('/api/capacity-report/details?date=2026-10-01').json
            completed=next(group for group in details['groups'] if group['key']=='completed')
            self.assertEqual([row['Order Number'] for row in completed['rows']],['C'])
            self.assertIn('grouped by report Date',completed['description'])
            month=client.get('/api/capacity-report/details?month=2026-10').json
            self.assertEqual([(r['Received'],r['Completed']) for r in month['daily']],[(3,1),(1,0)])
            book=load_workbook(BytesIO(client.get('/api/export/report?sheet=PR%20Excel').data),data_only=True)
            values=list(book['PR Excel'].values)
            self.assertEqual(values[4][1:3],(3,1))
            self.assertEqual(values[5][1:3],(1,0))
            self.assertEqual(values[6][1:3],(4,1))

    def test_previous_pending_and_later_completions_never_create_new_receipts(self):
        rows=[order('clarification','Awaiting for Clarification'), order('vendor','Assign to ABS'),
              order('in-house'),order('cancelled','Cancelled'),
              order('later','Completed and Delivered',completed='10/05/2026 10:00 AM'),
              order('tomorrow','Completed and Delivered',completed='10/02/2026 10:00 AM'),
              order('new-day-2',received='10/02/2026 09:00 AM'),
              order('new-done','Completed and Delivered',received='10/05/2026 09:00 AM',completed='10/05/2026 11:00 AM')]
        reports=sheet_reports({FULL:rows});original=deepcopy(reports)
        result=capacity_report(reports,{'mode':'tracker'})
        self.assertEqual([(r['Received'],r['Completed'],r['Clarification'],r['Vendor Pending'],r['In-House Pending']) for r in result['daily']],
                         [(6,0,1,1,1),(1,1,0,0,1),(1,2,0,0,0)])
        details=capacity_details(reports,'2026-10-02');groups={g['key']:g['rows'] for g in details['groups']}
        self.assertEqual([r['Order Number'] for r in groups['received']],['new-day-2'])
        self.assertEqual([r['Order Number'] for r in groups['completed']],['tomorrow'])
        self.assertEqual([r['Order Number'] for r in groups['in_house']],['new-day-2'])
        self.assertEqual(set(groups), {'received','completed','clarification','cancelled','vendor','in_house'})
        self.assertEqual(reports,original)

    def test_monthly_and_export_pending_totals_count_source_statuses_once(self):
        rows=[order('clarification','Awaiting for Clarification'),
              order('vendor','Assign to ABS',received='10/02/2026'),order('in-house',received='10/05/2026')]
        snapshot=imported_snapshot({'rows':rows,'columns':list(rows[0]),'files':[],'created':'2026-10-05'})
        frames,report=report_frames(snapshot,{'mode':'excel'})
        self.assertEqual(report['monthly'][0]['Raw Received'],3)
        self.assertEqual([report['monthly'][0][c] for c in ('Clarification','Vendor Pending','In-House Pending')],[1,1,1])
        self.assertEqual(frames['PR Excel'][-1][3:7],[1,0,1,1])
        self.assertEqual(frames['PR Excel'][2][3:7],[1,0,1,1])

    def test_shared_order_numbers_in_separate_imports_keep_independent_receipts_and_completions(self):
        from report_workspace import read_report_files
        from test_pdf_report_workspace import excel
        imported=read_report_files([('full.xlsx',excel([order('shared','Completed and Delivered',completed='10/02/2026')])),
                                   ('co_update.xlsx',excel([order('shared','Awaiting for Clarification'),order('next',received='10/05/2026')]))])
        snapshot=imported_snapshot(imported);saved=deepcopy(imported)
        report=capacity_report(snapshot['reports'],{'mode':'excel'})
        self.assertEqual([(r['Received'],r['Completed']) for r in report['daily']],[(2,1),(1,0)])
        groups={g['key']:g['rows'] for g in capacity_details(snapshot['reports'],'2026-10-01',use_report_date=True)['groups']}
        self.assertEqual([r['Source workbook'] for r in groups['clarification']],['co_update.xlsx'])
        self.assertEqual([r['Source workbook'] for r in groups['completed']],['full.xlsx'])
        self.assertEqual(imported,saved)

    def test_excel_import_counts_completed_status_by_report_date_despite_out_time(self):
        from report_workspace import read_report_files
        from test_pdf_report_workspace import excel
        rows=[dict(order('later','  COMPLETED  AND DELIVERED ',received='10/02/2026',completed='11/02/2026'),Date='10/01/2026'),
              dict(order('no-out','Completed and Delivered'),Date='10/01/2026'),
              dict(order('bad-out','Completed and Delivered',completed='invalid'),Date='10/01/2026'),
              dict(order('cancelled','Cancelled',completed='10/01/2026'),Date='10/01/2026'),
              dict(order('pending','Search In Progress',completed='10/01/2026'),Date='10/01/2026',**{'Task Status':'Completed and Delivered'}),
              dict(order('reopened','Completed and Delivered - Reopened'),Date='10/01/2026'),
              dict(order('weekend','Completed and Delivered',completed='10/08/2026'),Date='10/03/2026'),
              dict(order('fallback','Completed and Delivered',received='10/05/2026',completed='10/06/2026'),Date='invalid'),
              dict(order('task-status','',received='10/06/2026',completed='10/07/2026'),Date='10/05/2026',**{'Task Status':'Completed and Delivered'})]
        imported=read_report_files([('full.xlsx',excel(rows))]);saved=deepcopy(imported)
        snapshot=imported_snapshot(imported)
        frames,report=report_frames(snapshot,{'mode':'excel','capacity_month':'2026-10'})
        self.assertEqual([(r['Date'],r['Received'],r['Completed']) for r in report['daily']],
                         [('2026-10-01',6,5),('2026-10-02',1,1),('2026-10-05',2,2)])
        self.assertEqual((report['monthly'][0]['Received'],report['monthly'][0]['Completed']),(9,8))
        self.assertEqual(report['completion_date_basis'],'report_date')
        self.assertEqual(report['missing_completion_date_count'],0)
        self.assertEqual(frames['PR Excel'][-1][1:3],[9,8])
        for day in report['daily']:
            details=capacity_details(snapshot['reports'],day['Date'],use_report_date=True)
            completed=next(g['rows'] for g in details['groups'] if g['key']=='completed')
            self.assertEqual(len(completed),day['Completed'])
        self.assertEqual(imported,saved)

    def test_reference_report_totals_sum_daily_values_in_all_six_columns(self):
        reference=[
            ('2026-10-01',319,252,0,9,3,1), ('2026-10-02',454,363,0,4,2,3),
            ('2026-10-05',328,311,1,6,11,4), ('2026-10-06',293,302,2,2,9,0),
            ('2026-10-07',432,387,1,3,14,2), ('2026-10-08',369,372,0,7,8,11),
            ('2026-10-09',276,405,4,5,13,8), ('2026-10-12',152,228,3,3,6,5),
            ('2026-10-13',252,217,41,0,9,59)]
        columns=['Date','Received','Completed','Clarification','Cancelled','Vendor Pending','In-House Pending']
        reports={'daily':[dict(zip(columns, values)) for values in reference], 'monthly':[]}
        report=capacity_report(reports,{'mode':'excel'})
        self.assertEqual([report['monthly'][0][c] for c in columns[1:]], [2875,2837,52,39,75,93])
        self.assertEqual(report['monthly'][0]['Capacity'],6300)
        self.assertEqual([r['Received'] for r in report['daily']], [r[1] for r in reference])

    def test_capacity_publication_changes_only_pr_excel_and_keeps_chart_style(self):
        from report_workspace import publish_reports
        from test_pdf_report_workspace import ReportBook
        book=ReportBook()
        untouched=book.add('Other report',[{'Order Number':'unchanged','Status':'Cancelled'}])
        retired=book.add('Daily Orders',[{'Order Number':'keep-this-tab'}])
        before=deepcopy(untouched.values)
        retired_values=deepcopy(retired.values)
        snapshot=imported_snapshot({'rows':[order('A'),order('B',received='10/02/2026')],
                                   'columns':list(order('A')),'files':[],'created':'2026-10-02'})
        with tempfile.TemporaryDirectory() as root:
            result=publish_reports(book,snapshot,{'mode':'excel'},root,only_titles={'PR Excel'})
            self.assertEqual(result['tabs'],['PR Excel'])
            pr=book.worksheet('PR Excel');chart=book.charts[pr.id][0]
            chart['spec']['basicChart']['series'][0]['color']={'red':.1,'green':.2,'blue':.3}
            chart['position']['overlayPosition']['widthPixels']=777
            original=deepcopy(chart)
            with patch.object(book,'batch_update',wraps=book.batch_update) as publication:
                publish_reports(book,snapshot,{'mode':'excel'},root,only_titles={'PR Excel'})
            requests=publication.call_args.args[0]['requests']
            self.assertFalse(any('userEnteredFormat' in str(request) or 'ConditionalFormatRule' in str(request)
                                 for request in requests))
        self.assertEqual(book.worksheet('Other report').values,before)
        self.assertEqual(book.worksheet('Daily Orders').id,retired.id)
        self.assertEqual(book.worksheet('Daily Orders').values,retired_values)
        self.assertEqual(chart['position'],original['position'])
        self.assertEqual(chart['spec']['basicChart']['series'][0]['color'],original['spec']['basicChart']['series'][0]['color'])


if __name__=='__main__':
    unittest.main()

from io import BytesIO
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from preview_store import PreviewStore, compare
from server import create_app
import datatrace_sync as sync


class PreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = PreviewStore(self.root / 'previews')

    def frame(self, status='Available'):
        return pd.DataFrame([{'OPON': '0001', 'Arrival Time': '09/22/2026 01:00 PM',
                              'Task Name': 'Search', 'Task Status': status, 'Time Since Arrival': '1h 0m'}])

    def test_durable_sequential_exports_and_no_seed(self):
        self.assertEqual(self.store.list(), [])
        a = self.store.save(self.frame())
        b = PreviewStore(self.root / 'previews').save(self.frame('Completed'))
        self.assertEqual([a['name'], b['name']], ['preview1', 'preview2'])
        self.assertEqual(self.store.get(1)['rows'][0]['Task Status'], 'Available')
        for number in [1, 2]:
            self.assertTrue((self.root / 'previews' / f'preview{number}.xlsx').exists())
        self.assertEqual(self.store.get(1)['rows'][0]['OPON'], '0001')

    def test_changed_added_removed_and_ignored_clock(self):
        a = self.store.save(pd.concat([self.frame(), self.frame().assign(OPON='2')]))
        b = self.store.save(pd.concat([self.frame('Completed').assign(**{'Time Since Arrival': '2h 0m'}), self.frame().assign(OPON='3')]))
        result = compare(a, b)
        self.assertEqual(result['counts'], {'added': 1, 'removed': 1, 'modified': 1, 'unchanged': 0})
        self.assertEqual(result['record_counts'], {'matched': 1, 'missing': 1, 'newly_added': 1, 'unchanged': 0})
        self.assertEqual(result['matched_rows'][0]['Comparison Status'], 'Matched - changed')
        self.assertEqual({row['Comparison Status'] for row in result['unmatched_rows']}, {'Missing', 'Newly Added'})
        modified = [r for r in result['rows'] if r['Change'] == 'Modified']
        self.assertEqual(len(modified), 1)
        self.assertEqual(modified[0]['Column'], 'Task Status')

    def test_duplicates_preserved_and_invalid_keys_rejected(self):
        a = self.store.save(pd.concat([self.frame(), self.frame()]))
        b = self.store.save(self.frame())
        result = compare(a, b)
        self.assertEqual(result['counts']['removed'], 1)
        self.assertEqual(result['counts']['unchanged'], 1)
        self.assertIn('No unique', result['method'])
        with self.assertRaises(ValueError):
            compare(a, b, keys=['OPON'])

    def test_rows_after_previous_last_order(self):
        previous = self.store.save(pd.DataFrame([
            {'Order Number': 'ORD-100', 'Product': 'Full Title'},
            {'Order Number': 'ORD-101', 'Product': 'Current Owner'},
        ]))
        latest = self.store.save(pd.DataFrame([
            {'Order Number': 'ORD-100', 'Product': 'Full Title'},
            {'Order Number': 'ORD-101', 'Product': 'Current Owner'},
            {'Order Number': 'ORD-102', 'Product': 'Full Title'},
            {'Order Number': 'ORD-103', 'Product': 'Two Owner'},
        ]))
        appended = compare(previous, latest)['order_append']
        self.assertTrue(appended['available'])
        self.assertEqual(appended['anchor_order'], 'ORD-101')
        self.assertEqual(appended['count'], 2)
        self.assertEqual([row['Order Number'] for row in appended['rows']], ['ORD-102', 'ORD-103'])

    def test_delete_removes_preview_record_and_exports(self):
        saved = self.store.save(self.frame())
        self.store.delete(saved['id'])
        self.assertEqual(self.store.list(), [])
        self.assertFalse((self.store.root / 'preview1.csv').exists())
        self.assertFalse((self.store.root / 'preview1.xlsx').exists())
        with self.assertRaises(KeyError):
            self.store.get(saved['id'])

    def test_reporting_reads_history_in_one_transaction(self):
        self.store.save(self.frame())
        with patch.object(self.store, 'connect', wraps=self.store.connect) as connect:
            self.assertEqual(len(self.store.history()), 1)
            self.assertEqual(connect.call_count, 1)

    def test_second_real_pipeline_run_creates_preview_without_auto_sync(self):
        calls = []
        def scrape(command, cwd, env, *args):
            calls.append(1)
            Path(env['DATATRACE_OUTPUT_JSON']).write_text(
                self.frame('Available' if len(calls) == 1 else 'Completed').to_json(orient='records'))
            return Mock(returncode=0)
        with patch.object(sync, 'BASE_DIR', self.root), patch.dict(sync.os.environ, {'DATATRACE_USERNAME':'test', 'DATATRACE_PASSWORD':'test'}), \
                patch.object(sync, 'run_extractor', side_effect=scrape), \
                patch.object(sync, 'sync_workbook') as upload:
            a, b = sync.run_sync(auto_sync=False), sync.run_sync(auto_sync=False)
        self.assertEqual(len(calls), 2)
        self.assertEqual([a['preview_id'], b['preview_id']], [1, 2])
        self.assertEqual(b['google_sheet'], 'not_synced')
        upload.assert_not_called()
        self.assertEqual(self.store.get(2)['rows'][0]['Task Status'], 'Completed')

    def test_manual_google_sync_and_daily_schedule(self):
        syncer = Mock(return_value=['Full Title'])
        from datetime import datetime
        from server import IST
        clock = Mock(now=lambda: datetime(2026, 9, 30, 8, tzinfo=IST), snapshot=lambda: {})
        client = create_app(self.root / 'manual-sync', syncer=syncer, time_source=clock).test_client()
        self.assertEqual(client.post('/api/sync').status_code, 422)
        imported = client.post('/api/import', data={'file': (BytesIO(self.frame().to_csv(index=False).encode()), 'queue.csv')})
        self.assertEqual(imported.status_code, 201)
        for _ in range(200):
            state = client.get('/api/state').json
            if not state['job']['running']:
                break
            time.sleep(.01)
        self.assertEqual(state['job']['result']['trigger'], 'import')
        self.assertEqual(syncer.call_count, 1)
        self.assertEqual(syncer.call_args.args[0].iloc[0]['Task Status'], 'Available')
        self.assertEqual(client.post('/api/sync', json={'preview': imported.json['id']}).status_code, 202)
        for _ in range(200):
            if not client.get('/api/state').json['job']['running']:
                break
            time.sleep(.01)
        saved = client.post('/api/sync-schedule', json={'enabled': True, 'time': '14:30'})
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json['time'], '14:30')
        self.assertEqual(client.post('/api/sync-schedule', json={'enabled': True, 'time': '99:00'}).status_code, 422)

    def test_api_import_compare_excel_and_initial_empty(self):
        client = create_app(self.root / 'api', syncer=Mock(return_value=[])).test_client()
        self.assertEqual(client.get('/api/state').json['previews'], [])
        for status in ['Available', 'Completed']:
            response = client.post('/api/import', data={'file': (BytesIO(self.frame(status).assign(**{'Order Number': '0001'}).to_csv(index=False).encode()), 'queue.csv')})
            self.assertEqual(response.status_code, 201)
        for _ in range(200):
            if not client.get('/api/state').json['job']['running']:
                break
            time.sleep(.01)
        self.assertFalse(client.get('/api/state').json['job']['running'])
        result = client.post('/api/compare', json={'previous':1,'latest':2}).json
        self.assertEqual(result['counts']['modified'], 1)
        response = client.post('/api/compare', json={'previous':1,'latest':2,'download':True})
        workbook = load_workbook(BytesIO(response.data))
        self.assertIn('DataTraceChanges', workbook.active.tables)
        with client.get('/api/previews/1/download/xlsx') as downloaded:
            self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(client.delete('/api/previews/1').status_code, 200)
        self.assertEqual([p['id'] for p in client.get('/api/state').json['previews']], [2])
        self.assertEqual(client.post('/api/extract', headers={'Origin':'https://other.example'}).status_code, 403)

    def test_job_lock_released_after_failure_and_repeat_click(self):
        release = threading.Event()
        calls = []
        def runner(on_progress):
            calls.append(1)
            release.wait(timeout=2)
            raise RuntimeError('test failure')
        client = create_app(self.root / 'jobs', runner).test_client()
        self.assertEqual(client.post('/api/extract').status_code, 202)
        self.assertEqual(client.post('/api/extract').status_code, 409)
        release.set()
        for _ in range(100):
            if not client.get('/api/state').json['job']['running']:
                break
            time.sleep(.01)
        self.assertEqual(client.post('/api/extract').status_code, 202)
        for _ in range(100):
            if not client.get('/api/state').json['job']['running']:
                break
            time.sleep(.01)
        self.assertEqual(len(calls), 2)


if __name__ == '__main__':
    unittest.main()

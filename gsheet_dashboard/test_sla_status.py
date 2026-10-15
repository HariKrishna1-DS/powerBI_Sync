from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from sla_status import (build_sla_status, category, excel_report, sheet_summary_values,
                        DETAIL_COLUMNS, SUMMARY_COLUMNS)
from monthly_views import monthly_view_values
from report_workspace import imported_snapshot, read_report_files, report_frames, publish_reports
from test_pdf_report_workspace import ReportBook, excel, row
from tracker_sync import FULL


class SlaStatusTests(unittest.TestCase):
    def setUp(self):
        self.today = datetime(2026, 10, 14, 16, 15)
        self.rows = [row('001', 'Assign to ABS', **{'SLA Expiration': '2026-10-15 00:00:00', 'Typer': 'Maya'}),
                     row('002', 'Assign to ABS', **{'SLA Expiration': '-19d 4h 40m'}),
                     row('003', 'QC in Progress', **{'SLA Expiration': '10/14/2026 11:59 PM'}),
                     row('004', 'QC in Progress'),
                     row('005', 'QC in Progress', **{'SLA Expiration': 'invalid'})]
        self.snapshot = imported_snapshot(read_report_files([('full.xlsx', excel(self.rows))]))

    def test_dates_durations_serials_and_day_boundaries(self):
        cases = [('2026-10-14 00:00:00', 'Today SLA'), ('2026-10-14 23:59:59', 'Today SLA'),
                 ('2026-10-13 23:59:59', 'Negative SLA'), ('2026-10-15', 'Future SLA'),
                 ('-0d 1h', 'Negative SLA'), ('2h', 'Today SLA'), ('1d', 'Future SLA'),
                 ('46309.9', 'Today SLA'), ('', ''), ('invalid', '')]
        for value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(category(value, self.today), expected)
        self.assertEqual(category('2026-10-13T20:00:00Z', self.today), 'Today SLA')
        self.assertEqual(category('2h', self.today, datetime(2026, 10, 13, 16, 15)), 'Negative SLA')

    def test_summary_download_and_exclusions_reconcile(self):
        report = build_sla_status(self.snapshot, self.today)
        self.assertEqual(report['totals'], dict(Status='Grand Total', **{'Future SLA': 1, 'Negative SLA': 1, 'Today SLA': 1, 'Grand Total': 3}))
        self.assertEqual(report['excluded_count'], 2)
        book = load_workbook(excel_report(report), data_only=True)
        self.assertEqual(book.sheetnames, ['Sheet2', 'Sheet1'])
        self.assertEqual([cell.value for cell in book['Sheet2'][4]], SUMMARY_COLUMNS)
        self.assertEqual([cell.value for cell in book['Sheet1'][1]], DETAIL_COLUMNS)
        self.assertEqual(book['Sheet1'].max_row, 4)
        self.assertEqual(book['Sheet1']['A2'].value, '001')
        self.assertIsInstance(book['Sheet1']['B2'].value, datetime)
        self.assertEqual(book['Sheet2']['E7'].value, 3)
        self.assertEqual(len(book['Sheet2']._charts), 1)
        pivot = book['Sheet2']._pivots[0]
        self.assertEqual(pivot.cache.cacheSource.worksheetSource.sheet, 'Sheet1')
        self.assertEqual(pivot.cache.cacheSource.worksheetSource.ref, 'A1:M4')
        self.assertEqual(book['Sheet1']['M2'].value, 'full.xlsx')
        self.assertEqual(pivot.cache.recordCount, 3)
        self.assertEqual(pivot.cache.records.count, 3)
        self.assertEqual(pivot.dataFields[0].subtotal, 'count')
        self.assertEqual(pivot.rowFields[0].x, 11)
        self.assertEqual(pivot.colFields[0].x, 2)
        decoded = Counter((pivot.cache.cacheFields[11].sharedItems._fields[record._fields[11].v].v,
                           pivot.cache.cacheFields[2].sharedItems._fields[record._fields[2].v].v)
                          for record in pivot.cache.records.r)
        self.assertEqual(decoded, Counter((row['Status'], row['SLA Category']) for row in report['rows']))
        self.assertEqual(book['Sheet2']._charts[0].series[0].val.numRef.f, "'Sheet2'!$B$5:$B$6")

    def test_empty_report_exports_both_sheets_and_zero_totals(self):
        report = build_sla_status({'sheets': {'Overview': {'rows': [], 'columns': []}}}, self.today)
        book = load_workbook(excel_report(report), data_only=True)
        self.assertEqual(book.sheetnames, ['Sheet2', 'Sheet1'])
        self.assertEqual(book['Sheet2']['E5'].value, 0)
        self.assertEqual(book['Sheet1'].max_row, 1)
        self.assertEqual(len(book['Sheet2']._pivots), 0)

    def test_column_y_is_typer_in_both_publication_paths(self):
        frames, _ = report_frames(self.snapshot, {})
        for name in ('Full_Search_OCT_2026', 'Remaining_Search_OCT_2026'):
            self.assertEqual(frames[name][0][24], 'Typer')
            self.assertNotIn('SLA Category', frames[name][0])
        self.assertEqual(frames['Full_Search_OCT_2026'][1][24], 'Maya')
        views = monthly_view_values({FULL: self.rows}, headers=['Order Number', 'SLA Category', 'Typer'])
        self.assertEqual(views['Full_Search_OCT_2026'][0][24], 'Typer')
        self.assertEqual(views['Full_Search_OCT_2026'][1][24], 'Maya')

    def test_sheets_publish_table_and_reuse_sla_chart(self):
        book = ReportBook()
        with tempfile.TemporaryDirectory() as root, patch('sla_status.build_sla_status', side_effect=lambda snapshot: build_sla_status(snapshot, self.today)):
            publish_reports(book, self.snapshot, {}, root)
            publish_reports(book, self.snapshot, {}, root)
        sheet = book.worksheet('SLA Status')
        values = sheet.get_all_values(value_render_option='UNFORMATTED_VALUE')
        self.assertEqual(values[0], ['Count of Order Number', 'Column Labels'])
        self.assertEqual(values[1], SUMMARY_COLUMNS)
        self.assertEqual(values[2], ['Assign to ABS', 1, 1, '', 2])
        self.assertEqual(values[-1][0], 'Grand Total')
        self.assertEqual(book.props[sheet.id]['gridProperties']['frozenRowCount'], 2)
        formatting = [r['repeatCell'] for r in book.batches[-1]['requests']
                      if 'repeatCell' in r and r['repeatCell']['range']['sheetId'] == sheet.id]
        self.assertTrue(any(r['cell'] == {} and r['range'] == {'sheetId': sheet.id}
                            and r['fields'] == 'userEnteredFormat' for r in formatting))
        for request in formatting:
            fmt = request['cell'].get('userEnteredFormat', {})
            self.assertNotIn('backgroundColor', fmt)
            self.assertNotIn('foregroundColor', fmt.get('textFormat', {}))
        self.assertEqual(sheet.conditional_formats, [])
        self.assertEqual(len(book.charts[sheet.id]), 1)
        chart = book.charts[sheet.id][0]
        self.assertEqual(chart['spec']['basicChart']['chartType'], 'BAR')
        self.assertEqual(chart['spec']['basicChart']['series'][0]['series']['sourceRange']['sources'][0]['endRowIndex'], len(sheet.get_all_values())-1)
        self.assertEqual(chart['spec']['basicChart']['series'][0]['series']['sourceRange']['sources'][0]['startRowIndex'], 1)
        for series in chart['spec']['basicChart']['series']:
            self.assertEqual(len(set(series['color'].values())), 1)

    def test_sheets_remove_old_color_rules_and_clear_stale_rows(self):
        book = ReportBook()
        sheet = book.add('SLA Status', [])
        sheet.conditional_formats = [{'ranges': [{'sheetId': sheet.id}], 'booleanRule': {}}]
        with tempfile.TemporaryDirectory() as root:
            publish_reports(book, self.snapshot, {}, root)
            empty = imported_snapshot({'rows': [], 'columns': [], 'files': [], 'created': self.today.isoformat()})
            publish_reports(book, empty, {}, root)
        sheet = book.worksheet('SLA Status')
        self.assertEqual(sheet.conditional_formats, [])
        expected = sheet_summary_values(build_sla_status(empty))
        self.assertEqual(sheet.get_all_values(value_render_option='UNFORMATTED_VALUE'),
                         [row[:2] if index == 0 else row for index, row in enumerate(expected)])
        self.assertTrue(any('deleteEmbeddedObject' in r for r in book.batches[-1]['requests']))

    def test_routes_use_current_report_source_and_sample_layout(self):
        from flask import Flask
        from preview_store import PreviewStore
        from report_workspace import ReportWorkspace, register_report_routes
        with tempfile.TemporaryDirectory() as root:
            workspace = ReportWorkspace(PreviewStore(Path(root)))
            app = Flask(__name__)
            register_report_routes(app, workspace, threading.Lock(), lambda: self.snapshot, lambda: None, lambda: None)
            with app.test_client() as client:
                result = client.get('/api/sla-status')
                self.assertEqual(result.status_code, 200)
                self.assertEqual(result.json['source'], 'Imported Excel')
                download = client.get('/api/export/report?sheet=SLA%20Status')
                self.assertEqual(download.status_code, 200)
                book = load_workbook(BytesIO(download.data), data_only=True)
                self.assertEqual(book.sheetnames, ['Sheet2', 'Sheet1'])
                self.assertEqual(book['Sheet1'].max_row-1, result.json['totals']['Grand Total'])


if __name__ == '__main__':
    unittest.main()

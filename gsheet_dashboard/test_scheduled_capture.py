from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from preview_store import PreviewStore
from server import IST, create_app
from test_production_sync import Book, preview, raw
from order_reporting import automatic_sync_frame
from sheets_repository import sync_production


class ScheduledCaptureTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        self.store = PreviewStore(self.root / 'previews')
        (self.root / 'sync_schedule.json').write_text(json.dumps({
            'enabled': True, 'times': ['13:00']}))

    def wait_finished(self, client):
        for _ in range(200):
            job = client.get('/api/state').json['job']
            if not job['running']:
                return job
            time.sleep(.02)
        self.fail('Worker did not finish')

    def test_ist_extract_then_sync_saved_preview_and_once_per_day(self):
        release = threading.Event()
        self.addCleanup(release.set)
        def runner(on_progress):
            release.wait(5)
            preview = self.store.save(pd.DataFrame([{'Order Number': '1', 'Task Status': 'Available'}]))
            return {'preview_id': preview['id'], 'error': None}
        syncer = Mock(return_value=['Sheet1'])
        app = create_app(self.store.root, runner=runner, syncer=syncer)
        tick = app.extensions['schedule_tick']
        client = app.test_client()
        self.assertFalse(tick(datetime(2026, 9, 30, 7, 29, tzinfo=timezone.utc)))
        self.assertTrue(tick(datetime(2026, 9, 30, 7, 30, tzinfo=timezone.utc)))
        syncer.assert_not_called()
        self.assertEqual(client.get('/api/export/google-sheets').status_code, 409)
        release.set()
        job = self.wait_finished(client)
        self.assertIsNone(job['result']['error'])
        self.assertEqual(syncer.call_args.args[0].attrs['preview_name'], 'preview1')
        self.assertFalse(tick(datetime(2026, 9, 30, 8, 0, tzinfo=timezone.utc)))
        syncer.assert_called_once()

    def test_failed_extraction_never_syncs_old_preview(self):
        self.store.save(pd.DataFrame([{'Order Number': 'old', 'Task Status': 'Available'}]))
        syncer = Mock()
        app = create_app(self.store.root, runner=lambda on_progress: {'error': 'Login failed'}, syncer=syncer)
        app.extensions['schedule_tick'](datetime(2026, 9, 30, 7, 30, tzinfo=timezone.utc))
        job = self.wait_finished(app.test_client())
        self.assertIn('Login failed', job['result']['error'])
        syncer.assert_not_called()

    def test_schedule_save_skips_elapsed_times_and_preserves_completed_slots(self):
        app = create_app(self.store.root, runner=lambda on_progress: {'error': 'Test run'})
        client = app.test_client()
        with patch('server.datetime') as clock:
            clock.strptime = datetime.strptime
            clock.now.return_value = datetime(2026, 9, 30, 13, 0, 10, tzinfo=IST)
            saved = client.post('/api/sync-schedule', json={'enabled':True,'times':['13:00','14:00']}).json
            self.assertEqual(saved['triggered_today'], [])
            self.assertEqual(client.get('/api/state').json['schedule']['triggered_today'], ['13:00'])
            paused = client.post('/api/sync-schedule', json={'enabled':False,'times':saved['times']}).json
            resumed = client.post('/api/sync-schedule', json={'enabled':True,'times':paused['times']}).json
        self.assertEqual(resumed['times'], ['13:00','14:00'])
        self.assertEqual(resumed['triggered_today'], ['13:00'])

    def test_saving_a_past_minute_waits_until_next_day(self):
        app = create_app(self.store.root, runner=lambda on_progress: {'error': 'Should not run'})
        with patch('server.datetime') as clock:
            clock.strptime = datetime.strptime
            clock.now.return_value = datetime(2026, 9, 30, 13, 1, 10, tzinfo=IST)
            client = app.test_client()
            client.post('/api/sync-schedule', json={'enabled':True,'times':['13:00']})
        self.assertEqual(client.get('/api/state').json['schedule']['triggered_today'], ['13:00'])
        self.assertFalse(app.extensions['schedule_tick'](datetime(2026,9,30,8,0,tzinfo=timezone.utc)))

    def test_sync_failure_keeps_saved_preview_number(self):
        def runner(on_progress):
            preview = self.store.save(pd.DataFrame([{'Order Number':'1','Task Status':'Available'}]))
            return {'preview_id':preview['id']}
        app = create_app(self.store.root, runner=runner, syncer=Mock(side_effect=RuntimeError('Sync unavailable')))
        app.extensions['schedule_tick'](datetime(2026,9,30,7,30,tzinfo=timezone.utc))
        result = self.wait_finished(app.test_client())['result']
        self.assertEqual(result['preview_name'], 'preview1')
        self.assertIn('Sync unavailable', result['error'])

    def test_busy_run_does_not_consume_due_schedule(self):
        release = threading.Event()
        self.addCleanup(release.set)
        def runner(on_progress):
            release.wait(5)
            return {'error':'Test extraction stopped'}
        app = create_app(self.store.root, runner=runner)
        client = app.test_client()
        client.post('/api/extract')
        self.assertFalse(app.extensions['schedule_tick'](datetime(2026,9,30,7,30,tzinfo=timezone.utc)))
        self.assertNotIn('13:00', client.get('/api/state').json['schedule']['triggered_today'])
        release.set()
        self.wait_finished(client)

    def test_export_reads_all_actual_tabs(self):
        sheets = []
        for title, values in [('Sheet1', [['Order Number'], ['001']]),
                              ('Status Report', [['Preview'], ['preview25']]),
                              ('Extra tab', [['Comment'], ['=not a formula']])]:
            sheet = Mock(title=title)
            sheet.get_all_values.return_value = values
            sheets.append(sheet)
        book = Mock()
        book.worksheets.return_value = sheets
        with patch('server.target_worksheet', return_value=(book, sheets[0])):
            response = create_app(self.store.root).test_client().get('/api/export/google-sheets')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Production_data.xlsx', response.headers['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ['Sheet1', 'Status Report', 'Extra tab'])
        self.assertEqual(workbook['Sheet1']['A2'].value, '001')
        self.assertEqual(workbook['Extra tab']['A2'].data_type, 's')

    def test_preview_name_survives_historical_merge(self):
        book = Book()
        capture = preview(1, [raw('A')])
        capture.update(id=25, name='preview25')
        frame, _ = automatic_sync_frame(capture)
        with patch('datatrace_sync.target_worksheet', return_value=(book, book.sheets[0])), patch('sheets_repository.BASE_DIR', self.root):
            sync_production(frame)
        values = book.worksheet('Status Report').get_all_values()
        offset = values[0].index('Preview')
        self.assertEqual({row[offset] for row in values[1:]}, {'preview25'})

    def test_render_recovers_latest_numbered_preview_from_sheet(self):
        book = Book()
        capture = preview(1, [raw('001')])
        capture.update(id=27, name='preview27')
        frame, _ = automatic_sync_frame(capture)
        with patch('datatrace_sync.target_worksheet', return_value=(book, book.sheets[0])), patch('sheets_repository.BASE_DIR', self.root):
            sync_production(frame)
        with patch('server.BASE_DIR', self.root), patch('server.target_worksheet', return_value=(book, book.sheets[0])):
            client = create_app().test_client()
            state = client.get('/api/state').json
        self.assertEqual(state['previews'][0]['name'], 'preview27')
        self.assertEqual(client.get('/api/previews/27').json['rows'][0]['Order Number'], '001')
        self.assertEqual(self.store.get(27)['created'], capture['created'])

    def test_render_schedule_survives_local_file_loss(self):
        settings = Mock()
        cloud = {'value': json.dumps({'enabled': False, 'times': ['09:00']})}
        settings.acell.side_effect = lambda cell: Mock(value=cloud['value'])
        settings.update_acell.side_effect = lambda cell, value: cloud.update(value=value)
        book = Mock()
        book.worksheet.side_effect = lambda title: settings if title == '__DataTrace_Config' else Mock(get_all_values=Mock(return_value=[]))
        with patch.dict(os.environ, {'RENDER': 'true'}), patch('server.target_worksheet', return_value=(book, Mock())):
            with patch('server.datetime') as clock:
                clock.strptime = datetime.strptime
                clock.now.return_value = datetime(2026, 9, 30, 13, 1, tzinfo=IST)
                client = create_app(self.store.root).test_client()
                saved = client.post('/api/sync-schedule', json={'enabled': True, 'times': ['13:00', '14:00']})
                self.assertEqual(saved.status_code, 200)
                self.assertEqual(json.loads(cloud['value'])['times'], ['13:00', '14:00'])
            (self.root / 'sync_schedule.json').unlink()
            restored = create_app(self.store.root).test_client().get('/api/state').json['schedule']
        self.assertTrue(restored['enabled'])
        self.assertEqual(restored['times'], ['13:00', '14:00'])


if __name__ == '__main__':
    unittest.main()

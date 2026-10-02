from datetime import datetime, timezone
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import Mock, patch

import requests

from indian_clock import IndianClock
from tracker_sync import FULL, REMAINING, merge_trackers, free_site
from tracker_formatting import PALETTE, rules_for, ensure_tracker_formatting, comparable_rule
from test_tracker_v2 import Book, HEADERS, order


class OctoberTrackerTests(unittest.TestCase):
    def test_empty_columns_in_new_existing_and_absent_orders(self):
        rows = [order('A', Comments='old', Assignee='old'), order('B', Comments='absent', Assignee='absent')]
        incoming = [{'Order Number': x, 'Product': 'Full Title', 'Comment': 'raw comment', 'Last User': 'raw user'} for x in ['A', 'C']]
        result, _ = merge_trackers({FULL: rows, REMAINING: []}, incoming, [], datetime(2026, 10, 2))
        self.assertEqual([r['Order Number'] for r in result[FULL]], ['A', 'B', 'C'])
        self.assertTrue(all(not r['Comments'] and not r['Assignee'] for r in result[FULL]))
        self.assertEqual(incoming[0]['Comment'], 'raw comment')

    def test_free_site_all_retained_rows_and_boundary(self):
        rows = [order('A', **{'Out Time': '10/01/2026 10:00 AM', 'SLA Expiration': '10/01/2026 10:00 AM'}),
                order('B', **{'Out Time': '10/01/2026 10:01 AM', 'SLA Expiration': '10/01/2026 10:00 AM'}),
                order('C', **{'Out Time': '', 'Free Site': 'Missing', 'SLA Expiration': '10/01/2026 10:00 AM'})]
        result, _ = merge_trackers({FULL: rows, REMAINING: []}, [], [], datetime(2026, 10, 2))
        self.assertEqual([r['Free Site'] for r in result[FULL]], ['On Time', 'Missing', ''])
        self.assertEqual(free_site({'Out Time': '10/01/2026 10:00 AM', 'SLA Expiration': 'PAUSED'}, datetime(2026, 10, 2))[0], '')

    def test_valid_out_time_completes_current_absent_and_new_orders(self):
        existing = [order('A', 'Awaiting for Clarification', **{'Out Time': '10/01/2026',
                    'SLA Expiration': '10/02/2026 01:00 PM'}),
                    order('B', 'Cancelled', **{'Out Time': '46279',
                    'SLA Expiration': '10/02/2026 01:00 PM'})]
        incoming = [{'Order Number': 'A', 'Product': 'Full Title', 'WorkflowSuspended': True},
                    {'Order Number': 'C', 'Product': 'Full Title', 'Out Time': '10/01/2026',
                     'SLA Expiration': '10/02/2026 01:00 PM'}]
        result, _ = merge_trackers({FULL: existing, REMAINING: []}, incoming, [], datetime(2026, 10, 2))
        self.assertEqual([r['Status'] for r in result[FULL]], ['Completed and Delivered'] * 3)
        self.assertEqual([r['Free Site'] for r in result[FULL]], ['On Time'] * 3)

    def test_exact_sample_colors_cover_free_site(self):
        self.assertEqual(PALETTE['cancelled'], '#c00000')
        self.assertEqual(PALETTE['completed and delivered'], '#ffff99')
        self.assertEqual(PALETTE['awaiting for clarification'], '#a66bd3')
        rules = rules_for(15, HEADERS)
        first = rules[0]
        self.assertEqual(first['ranges'][0]['startColumnIndex'], 0)
        self.assertEqual(first['ranges'][0]['endColumnIndex'], len(HEADERS))
        self.assertNotIn('endRowIndex', first['ranges'][0])
        self.assertFalse(any('TRUE)' in str(rule) for rule in rules))

    def test_excel_serial_dates_are_used_for_sla(self):
        from datatrace_sync import parse_report_datetime
        self.assertEqual(parse_report_datetime('46277'), datetime(2026, 9, 12))
        self.assertEqual(parse_report_datetime('46277.5'), datetime(2026, 9, 12, 12))
        self.assertEqual(free_site({'Out Time': '46277', 'SLA Expiration': '09/14/2026 03:00 PM'}, datetime(2026, 10, 2))[0], 'On Time')
        self.assertEqual(free_site({'Out Time': '46279', 'SLA Expiration': '09/11/2026 03:00 PM'}, datetime(2026, 10, 2))[0], 'Missing')

    def test_formatting_reconciliation_is_idempotent(self):
        book = Book()
        sheet = book.add_worksheet(FULL, 10, len(HEADERS))
        sheet.values = [HEADERS]
        self.assertGreater(ensure_tracker_formatting(book), 0)
        self.assertEqual(ensure_tracker_formatting(book), 0)

    def test_google_normalized_rules_and_growing_row_ranges(self):
        wanted = rules_for(15, HEADERS)[0]
        remote = json.loads(json.dumps(wanted))
        remote['ranges'][0]['endRowIndex'] = 150
        fmt = remote['booleanRule']['format']
        fmt['backgroundColor'] = {k: round(v, 7) for k, v in fmt['backgroundColor'].items() if v}
        fmt['backgroundColorStyle'] = {'rgbColor': fmt['backgroundColor']}
        self.assertEqual(comparable_rule(remote, 150), comparable_rule(wanted, 150))
        self.assertNotEqual(comparable_rule(remote, 200), comparable_rule(wanted, 200))

    def test_excel_export_uses_row_palette_for_free_site(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from server import create_app
        book = Book()
        sheet = book.add_worksheet(FULL, 10, 3)
        sheet.values = [['Order Number', 'Status', 'Free Site'], ['A', 'Cancelled', 'Missing']]
        with tempfile.TemporaryDirectory() as root, patch('server.target_worksheet', return_value=(book, sheet)):
            client = create_app(Path(root) / 'previews').test_client()
            response = client.get('/api/export/google-sheets?short_names=true')
            self.assertEqual(response.status_code, 200)
            ws = load_workbook(BytesIO(response.data))['Full Title']
            self.assertEqual(ws['A2'].fill.fgColor.rgb[-6:], 'C00000')
            self.assertEqual(ws['B2'].fill.fgColor.rgb[-6:], 'C00000')
            self.assertEqual(ws['C2'].fill.fgColor.rgb[-6:], 'C00000')


class IndianClockTests(unittest.TestCase):
    def make_clock(self):
        ticks = [100.0]
        session = Mock()
        session.head.return_value.headers = {'Date': 'Fri, 02 Oct 2026 06:30:00 GMT'}
        clock = IndianClock(session=session, monotonic=lambda: ticks[0])
        return clock, ticks, session

    def test_network_time_advances_without_system_clock(self):
        clock, ticks, _ = self.make_clock()
        self.assertTrue(clock.synchronize())
        with patch('indian_clock.datetime') as wall:
            wall.fromtimestamp = datetime.fromtimestamp
            now = clock.now()
            self.assertEqual((now.hour, now.minute), (12, 0))
            ticks[0] += 60
            self.assertEqual(clock.now().minute, 1)
            wall.now.assert_not_called()

    def test_failed_refresh_retains_monotonic_time(self):
        clock, ticks, session = self.make_clock()
        clock.synchronize()
        session.head.side_effect = requests.Timeout()
        ticks[0] += 10
        self.assertFalse(clock.synchronize())
        self.assertTrue(clock.snapshot()['synchronized'])
        self.assertEqual(clock.now().second, 10)

    def test_no_initial_network_time_blocks_scheduling(self):
        from server import create_app
        source = Mock()
        source.now.return_value = None
        with tempfile.TemporaryDirectory() as root:
            runner = Mock()
            app = create_app(Path(root) / 'previews', runner=runner, time_source=source)
            response = app.test_client().post('/api/sync-schedule', json={'enabled': True, 'times': ['12:00']})
            self.assertEqual(response.status_code, 503)
            self.assertFalse(app.extensions['schedule_tick']())
            runner.assert_not_called()
            self.assertEqual(app.test_client().post('/api/sync-schedule', json={'enabled': False, 'times': ['12:00']}).status_code, 200)

    def test_scheduler_uses_network_ist_for_due_time_and_once_only(self):
        from server import create_app
        clock, _, _ = self.make_clock()
        clock.synchronize()
        with tempfile.TemporaryDirectory() as root:
            (Path(root) / 'sync_schedule.json').write_text(json.dumps({'enabled': True, 'times': ['11:59', '12:00', '12:01']}))
            app = create_app(Path(root) / 'previews', runner=lambda on_progress: {'error': 'fixture'}, time_source=clock)
            with patch('server.datetime') as wall:
                wall.now.side_effect = AssertionError('System wall time used')
                self.assertTrue(app.extensions['schedule_tick']())
            saved = json.loads((Path(root) / 'sync_schedule.json').read_text())
            self.assertEqual(saved['triggered_today'], ['11:59', '12:00'])
            self.assertEqual(saved['last_triggered_date'], '2026-10-02')
            self.assertFalse(app.extensions['schedule_tick']())


if __name__ == '__main__':
    unittest.main()

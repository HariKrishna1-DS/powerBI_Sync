"""Report source isolation, batch imports, selected-date publication and recovery."""
from io import BytesIO
from contextlib import contextmanager
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from report_workspace import (ReportWorkspace, read_report_files, imported_snapshot,
                              report_frames, publish_reports, combine_sources, DAILY_COLUMNS)
from preview_store import PreviewStore
from test_monthly_production import AtomicBook
from tracker_sync import FULL, REMAINING


def row(number, status='Search In Progress', date='10/01/2026 09:00 AM', **extra):
    return {'Order Number': number, 'Product': 'Full Search', 'Status': status,
            'In-Time': date, **extra}


def excel(rows, title='TV Orders'):
    book=Workbook();sheet=book.active;sheet.title=title
    columns=list(dict.fromkeys(c for record in rows for c in record))
    sheet.append(columns)
    for record in rows:
        sheet.append([record.get(c,'') for c in columns])
    stream=BytesIO();book.save(stream)
    return stream.getvalue()


@contextmanager
def report_app(root):
    from server import create_app
    prior_threads=set(threading.enumerate())
    app=create_app(root=root)
    try:
        yield app
    finally:
        app.extensions['stop_scheduler'].set()
        for worker in set(threading.enumerate())-prior_threads:
            if worker.name=='report-publication':
                worker.join(timeout=10)
                if worker.is_alive():
                    raise RuntimeError('Report fixture worker did not stop before cleanup.')


class ReportBook(AtomicBook):
    def __init__(self):
        super().__init__();self.props={};self.charts={}

    def fetch_sheet_metadata(self, params=None):
        result=super().fetch_sheet_metadata(params)
        for sheet in result['sheets']:
            number=sheet['properties']['sheetId']
            sheet['properties'].update(self.props.get(number,{}))
            sheet['charts']=self.charts.get(number,[])
        return result

    def batch_update(self, body):
        super().batch_update(body)
        for item in body['requests']:
            if 'updateSheetProperties' in item:
                props=item['updateSheetProperties']['properties']
                self.props.setdefault(props['sheetId'],{}).update({k:v for k,v in props.items() if k!='title'})
            if 'addChart' in item:
                chart=item['addChart']['chart']
                number=chart['position']['overlayPosition']['anchorCell']['sheetId']
                self.charts.setdefault(number,[]).append(dict(chart,chartId=number+100))
            if 'updateChartSpec' in item:
                change=item['updateChartSpec']
                for charts in self.charts.values():
                    for chart in charts:
                        if chart['chartId']==change['chartId']:
                            chart['spec']=change['spec']


class PdfReportWorkspaceTests(unittest.TestCase):
    def test_batch_reads_multiple_sheets_preserves_ids_and_reports_duplicates(self):
        first=excel([row('0001'),row('0002',Product='Update')])
        second=excel([row('0001','Assign to ABS'),row('0003','Cancelled',Product='Update')])
        result=read_report_files([('full.xlsx',first),('remaining.xlsx',second)])
        self.assertEqual(len(result['rows']),3)
        self.assertEqual(result['rows'][0]['Order Number'],'0001')
        report=imported_snapshot(result)['reports']['daily'][0]
        self.assertEqual([report[c] for c in DAILY_COLUMNS[1:7]],[3,0,0,1,1,1])
        self.assertEqual(result['files'][1]['duplicates'],1)

    def test_daily_counts_partition_received_and_sla_is_eligible_only(self):
        records=[row('1','Completed and Delivered',**{'Out Time':'10/01/2026 10:00 AM','SLA Expiration':'10/01/2026 10:00 AM'}),
                 row('2','Awaiting for Clarification'),row('3','Cancelled'),row('4','Assign to ABS'),row('5','QC in Progress')]
        report=imported_snapshot(read_report_files([('full.xlsx',excel(records))]))['reports']['daily'][0]
        self.assertEqual(sum(report[c] for c in DAILY_COLUMNS[2:7]),report['Received'])
        self.assertEqual((report['SLA OnTime'],report['Missing']),(1,0))

    def test_selected_date_frames_share_schema_and_capacity_uses_real_days(self):
        imported=read_report_files([('full.xlsx',excel([row('1'),row('2',date='10/02/2026 09:00 AM',Product='Update')]))])
        frames,capacity=report_frames(imported_snapshot(imported),{'selected_date':'2026-10-01','capacity':700,'extended_capacity':750})
        self.assertEqual(len(frames),7)
        self.assertEqual(frames['Daily Status Report'][1][0],'2026-10-01')
        self.assertEqual(frames['Full_Search_OCT_2026'][0],frames['Remaining_Search_OCT_2026'][0])
        self.assertEqual(capacity['monthly'][0]['Capacity'],1400)
        self.assertEqual(len(capacity['monthly']),1)

    def test_publish_hides_internal_tabs_preserves_trackers_and_reuses_chart(self):
        book=ReportBook();book.add(FULL,[row('tracker')]);book.add(REMAINING,[])
        legacy=book.add('Capacity Report', [])
        book.props[legacy.id]={'gridProperties':{'rowCount':500,'columnCount':13}}
        original=book.worksheet(FULL).get_all_values()
        snapshot=imported_snapshot(read_report_files([('full.xlsx',excel([row('001')]))]))
        with tempfile.TemporaryDirectory() as path:
            publish_reports(book,snapshot,{},path)
            publish_reports(book,snapshot,{},path)
            self.assertTrue(list(Path(path).glob('*.json.gz')))
        self.assertEqual(book.worksheet(FULL).get_all_values(),original)
        self.assertTrue(book.props[book.worksheet(FULL).id]['hidden'])
        self.assertEqual(sum(len(items) for items in book.charts.values()),1)
        self.assertEqual(sum(not book.props.get(sheet.id,{}).get('hidden',False) for sheet in book.worksheets()),7)
        self.assertEqual(book.worksheet('PR Excel').id,legacy.id)
        self.assertEqual(str(book.worksheet('PR Excel').get_all_values()[1][0]),'46296')
        clearing=[r['repeatCell'] for r in book.batches[-1]['requests']
                  if 'repeatCell' in r and r['repeatCell']['range']['sheetId']==legacy.id
                  and r['repeatCell']['fields']=='userEnteredFormat']
        self.assertTrue(any(r['range'].get('startColumnIndex')==9 and
                            r['range']['endColumnIndex']==13 and r['cell']=={} for r in clearing))
        self.assertTrue(any(r['range'].get('startRowIndex')==6 and
                            r['range'].get('endRowIndex')==500 and r['cell']=={} for r in clearing))

    def test_duplicate_details_survive_reload_and_identify_both_sources(self):
        imported=combine_sources(read_report_files([
            ('full.xlsx',excel([row('A1'),row('A2')])),
            ('remaining.xlsx',excel([row(' a1 ','Assign to ABS'),row('A3'),
                                     row('A1','Completed and Delivered')],title='Updates'))])['sources'])
        with tempfile.TemporaryDirectory() as path:
            store=PreviewStore(Path(path)/'previews')
            ReportWorkspace(store).save(mode='excel',imported=imported,files=imported['files'],
                                        row_count=len(imported['rows']))
            public=ReportWorkspace(store).public()
        self.assertEqual((public['input_row_count'],public['row_count'],public['duplicate_row_count']),(5,3,2))
        self.assertEqual(len(public['duplicate_orders']),2)
        first,last=public['duplicate_orders']
        self.assertEqual((first['Earlier file'],first['Replacement file'],first['Replacement worksheet']),
                         ('full.xlsx','remaining.xlsx','Updates'))
        self.assertEqual((last['Earlier file'],last['Status']),('remaining.xlsx','Completed and Delivered'))
        self.assertEqual(imported['rows'][0]['Status'],'Completed and Delivered')

    def test_single_tab_downloads_use_reference_colors_and_real_formatted_dates(self):
        imported=read_report_files([('full.xlsx',excel([row('A1'),row('A2',date='10/02/2026 09:00 AM')]))])
        with tempfile.TemporaryDirectory() as path:
            root=Path(path)/'previews'
            ReportWorkspace(PreviewStore(root)).save(mode='excel',imported=imported)
            with report_app(root) as app:
                client=app.test_client()
                response=client.get('/api/export/report?sheet=PR%20Excel')
                self.assertEqual(response.status_code,200)
                book=load_workbook(BytesIO(response.data))
                self.assertEqual(book.sheetnames,['PR Excel'])
                sheet=book.active
                self.assertEqual(sheet.max_column,9)
                self.assertEqual((sheet['A2'].number_format,sheet['A5'].number_format),('mmm-yy','dd-mm-yy'))
                self.assertEqual(sheet['A5'].value.strftime('%d-%m-%y'),'01-10-26')
                for cell,color in {'A1':'DCA683','A2':'C6E0B4','A3':'8EA9DB','A4':'DCA683',
                                   'A5':'E5E5E5','H5':'E5E5E5','B5':'FFFFFF','A7':'FFD966'}.items():
                    self.assertEqual(sheet[cell].fill.fgColor.rgb[-6:],color,cell)
                self.assertEqual(sheet['J1'].fill.patternType,None)
                self.assertEqual(len(sheet._charts),1)
                daily=load_workbook(BytesIO(client.get('/api/export/report?sheet=Daily%20Orders').data))
                self.assertEqual(daily.sheetnames,['Daily Status Report'])
                self.assertEqual(daily.active['A2'].number_format,'dd-mm-yy')
                self.assertEqual(daily.active['A2'].value.strftime('%d-%m-%y'),'01-10-26')
                self.assertEqual(daily.active['A1'].fill.fgColor.rgb[-6:],'DCA683')
                self.assertEqual(daily.active['A2'].fill.fgColor.rgb[-6:],'E5E5E5')
                self.assertEqual(daily.active['A4'].fill.fgColor.rgb[-6:],'FFD966')

    def test_api_failed_batch_keeps_prior_source_and_local_report_survives_cloud_failure(self):
        with tempfile.TemporaryDirectory() as path, patch('server.target_worksheet',side_effect=RuntimeError('Offline')), report_app(Path(path)/'previews') as app:
            client=app.test_client()
            response=client.post('/api/report-import',data={'files':[(BytesIO(excel([row('001')])),'full.xlsx')]})
            self.assertEqual(response.status_code,201)
            self.assertFalse(response.json['publication']['synced'])
            self.assertEqual(client.get('/api/live-sheets').json['source'],'Imported Excel')
            bad=client.post('/api/report-import',data={'files':[(BytesIO(excel([row('002')])),'full.xlsx'),(BytesIO(b'broken'),'bad.xlsx')]})
            self.assertEqual(bad.status_code,422)
            self.assertEqual(client.get('/api/live-sheets').json['sheets']['Overview']['rows'][0]['Order Number'],'001')
            self.assertEqual(PreviewStore(Path(path)/'previews').list(),[])
            changed=client.post('/api/report-source',json={'mode':'tracker'})
            self.assertEqual(changed.json['mode'],'tracker')
            client.post('/api/report-source',json={'mode':'excel'})
            selected=client.post('/api/report-date',json={'date':'2026-10-01'})
            self.assertEqual(selected.json['date'],'2026-10-01')
            self.assertEqual(client.post('/api/sla-comments',json={}).status_code,409)
            export=client.get('/api/export/report')
            self.assertEqual(export.status_code,200)
            book=load_workbook(BytesIO(export.data))
            self.assertEqual(len(book.sheetnames),7)
            state=ReportWorkspace(PreviewStore(Path(path)/'previews')).read()
            self.assertEqual(state['mode'],'excel')
            self.assertEqual(state['selected_date'],'2026-10-01')

    def test_queue_capture_does_not_replace_imported_public_reports(self):
        import pandas as pd
        from monthly_sync import sync_monthly
        from monthly_production import tab_name
        from order_reporting import automatic_sync_frame
        book=ReportBook()
        for base in (FULL,REMAINING):
            book.add(tab_name(base,'2026-10'),[])
        imported=read_report_files([('full.xlsx',excel([row('EXCEL-001')]))])
        with tempfile.TemporaryDirectory() as path:
            store=PreviewStore(Path(path)/'previews')
            ReportWorkspace(store).save(mode='excel',imported=imported)
            publish_reports(book,imported_snapshot(imported),{},Path(path)/'backups')
            before={name:book.worksheet(name).get_all_values() for name in ('All Products','Full_Search_OCT_2026','Daily Status Report')}
            saved=store.save(pd.DataFrame([row('CAPTURE-001')]),source='Queue')
            frame,_=automatic_sync_frame(store.get(saved['id']),store=store)
            sync_monthly(book,frame,store)
            for name,values in before.items():
                self.assertEqual(book.worksheet(name).get_all_values(),values)
            self.assertTrue(any('CAPTURE-001' in cells for cells in book.worksheet('Sheet1').get_all_values()))


if __name__=='__main__':
    unittest.main()

"""Imported SLA controls edit one date/workbook and reconcile every report."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, imported_snapshot, read_report_files, register_report_routes, report_frames
from test_pdf_report_workspace import ReportBook, row, excel
from tracker_sync import key
from workspace_backup import make_backup, restore_backup


class ImportedSlaManualTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.store = PreviewStore(Path(folder.name) / 'previews')
        self.workspace = ReportWorkspace(self.store)
        data = read_report_files([
            ('full.xlsx', excel([row('001', date='10/13/2026 09:00 AM', Status='Typing in Progress', Borrower='Original', **{'SLA Expiration': '10/14/2026 05:00 PM'}),
                                 row('old', date='10/12/2026 09:00 AM', **{'SLA Expiration': '2026-10-14'}),
                                 row('done', date='10/13/2026 09:00 AM', **{'Out Time': '10/13/2026 10:00 AM', 'SLA Expiration': '2026-10-14'})])),
            ('remaining.xlsx', excel([row('001', date='10/13/2026 09:00 AM', **{'SLA Expiration': '2026-10-14'})]))])
        self.workspace.save(mode='excel', imported=data, cloud_reset=True)
        self.app = Flask(__name__)
        self.book = ReportBook()
        self.start = patch('threading.Thread.start')
        self.start.start()
        self.addCleanup(self.start.stop)
        register_report_routes(self.app, self.workspace, threading.Lock(),
            lambda **kwargs: imported_snapshot(self.workspace.read()['imported']), lambda: (self.book, None), lambda: None)
        self.client = self.app.test_client()

    def report(self, day='2026-10-13'):
        result = self.client.get('/api/sla-status', query_string={'date': day})
        self.assertEqual(result.status_code, 200, result.json)
        return result.json

    def post(self, route, **values):
        return self.client.post(route, json={'mode': 'excel', 'report_date': '2026-10-13', **values})

    def lookup(self, number='001', source='full.xlsx'):
        return self.client.get('/api/sla-order', query_string={'mode': 'excel', 'date': '2026-10-13', 'order_number': number, 'source_workbook': source})

    def save(self, number='001', source='full.xlsx', status='Manual review', fields=None, lookup=None, **extra):
        return self.post('/api/sla-order', order_number=number, source_workbook=source, status=status,
                         fields=fields or {'SLA Expiration': '2026-10-14T17:00'},
                         expected_revision=(lookup or {}).get('revision'), **extra)

    def test_create_rename_and_export_empty_status_is_scoped_to_report_date(self):
        before = self.report()
        self.assertEqual(self.post('/api/sla-statuses', status='Manual review').status_code, 200)
        self.assertEqual(self.post('/api/sla-statuses', status='Manual QA', previous='Manual review').status_code, 200)
        after = self.report()
        self.assertEqual(after['totals'], before['totals'])
        self.assertEqual(after['custom_statuses'], ['Manual QA'])
        self.assertEqual(self.report('2026-10-12')['custom_statuses'], [])
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&date=2026-10-13').data))
        self.assertIn('Manual QA', [cells[0].value for cells in book['Sheet2'].iter_rows()])
        settings = self.workspace.read()
        frames, _ = report_frames(imported_snapshot(settings['imported']), settings)
        self.assertTrue(any(cells[0] == 'Manual QA' for cells in frames['SLA Status']))

    def test_lookup_and_move_edit_exact_workbook_without_duplicates_or_capture_changes(self):
        ambiguous = self.lookup(source='')
        self.assertEqual(ambiguous.status_code, 400)
        self.assertIn('multiple workbooks', ambiguous.json['error'])
        lookup = self.lookup().json
        self.assertEqual(lookup['fields']['SLA Expiration'], '2026-10-14 17:00:00')
        result = self.save(lookup=lookup, fields={'Borrower': 'Corrected', 'SLA Expiration': '2026-10-14T17:00'})
        self.assertEqual(result.status_code, 200, result.json)
        report = self.report()
        self.assertEqual(report['totals']['Grand Total'], 3)
        owners = {r['Source workbook']: r for r in report['rows'] if r['Order Number'] == '001'}
        self.assertEqual((owners['full.xlsx']['Status'], owners['full.xlsx']['Borrower']), ('Manual review', 'Corrected'))
        self.assertNotEqual(owners['remaining.xlsx']['Status'], 'Manual review')
        self.assertEqual(self.store.list(), [])
        self.assertEqual(self.save(lookup=lookup).status_code, 409)
        self.assertEqual(self.lookup().json['fields']['Borrower'], 'Corrected')

    def test_new_order_and_copy_are_persisted_once_and_visible_in_capacity_search_and_export(self):
        self.assertEqual(self.save(number='new', source='', fields={'Product': 'Full Title'}).status_code, 409)
        result = self.save(number='new', source='', fields={'Product': 'Full Title', 'Borrower': 'New borrower', 'SLA Expiration': '2026-10-13T17:00'})
        self.assertEqual(result.status_code, 200, result.json)
        self.assertTrue(result.json['created'])
        report = self.report()
        self.assertEqual((report['totals']['Grand Total'], report['manual_order_count']), (4, 1))
        self.assertEqual(self.save(number='NEW', source='Manual SLA orders', fields={'Product': 'Full Title', 'SLA Expiration': '2026-10-14T17:00'}).status_code, 409)
        copied = self.save(number='copy', source='full.xlsx', fields={'Product': 'Full Title', 'Borrower': 'Copy', 'SLA Expiration': '2026-10-14T17:00'})
        self.assertEqual(copied.status_code, 200, copied.json)
        snapshot = imported_snapshot(self.workspace.read()['imported'])
        self.assertEqual(len(snapshot['sheets']['Overview']['rows']), 6)
        self.assertEqual(self.report()['totals']['Grand Total'], 5)
        frames, capacity = report_frames(snapshot, self.workspace.read())
        self.assertTrue(any('new' in cells for cells in frames['All Products']))
        self.assertEqual(next(day for day in capacity['daily'] if day['Date'] == '2026-10-13')['Received'], 5)
        found = self.client.get('/api/report-order-search?search=new')
        self.assertEqual(found.json['total'], 1)
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&date=2026-10-13').data))
        self.assertEqual(book['Sheet1'].max_row - 1, 5)

    def test_delete_status_moves_exact_orders_and_bulk_editor_preserves_other_workbooks(self):
        self.post('/api/sla-statuses', status='Manual review')
        self.assertEqual(self.save(lookup=self.lookup().json).status_code, 200)
        report = self.report()
        identities = [key(r) for r in report['rows'] if r['Status'] == 'Manual review']
        stale = self.post('/api/sla-statuses', action='delete', status='Manual review', replacement='Typing in Progress', order_numbers=[])
        self.assertEqual(stale.status_code, 409)
        moved = self.post('/api/order-status', status='Manual review', replacement='QC in Progress', order_numbers=identities,
                          edit_scope='orders', sla_report_date='2026-10-13')
        self.assertEqual(moved.status_code, 200, moved.json)
        self.assertEqual(moved.json['updated'], 1)
        deleted = self.post('/api/sla-statuses', action='delete', status='Manual review', order_numbers=[])
        self.assertEqual(deleted.status_code, 200, deleted.json)
        self.assertNotIn('Manual review', self.report()['custom_statuses'])
        self.assertEqual(self.report()['totals']['Grand Total'], 3)

    def test_completion_protection_date_validation_and_backup_recovery(self):
        lookup = self.lookup('done').json
        blocked = self.save(number='done', lookup=lookup)
        self.assertEqual(blocked.status_code, 409)
        self.assertIn('Clear Out Time', blocked.json['error'])
        opened = self.save(number='done', lookup=lookup, fields={'SLA Expiration': '2026-10-14T17:00', 'Out Time': ''})
        self.assertEqual(opened.status_code, 200, opened.json)
        self.post('/api/sla-statuses', status='Keep after restore')
        backup = make_backup(self.store).getvalue()
        restore_backup(self.store, backup)
        self.assertIn('Keep after restore', self.report()['custom_statuses'])
        self.assertEqual(self.lookup('done').json['row']['Status'], 'Manual review')
        wrong_date = self.save(report_date='2026-10-12')
        self.assertEqual(wrong_date.status_code, 409)
        self.assertIn('another report date', wrong_date.json['error'])

    def test_local_manual_edits_survive_adding_another_import_before_sheet_publication(self):
        from import_excel_changes import preserve_sheet_edits
        from report_workspace import import_sources, combine_sources
        self.assertEqual(self.save(lookup=self.lookup().json, fields={'Borrower': 'Keep local', 'SLA Expiration': '2026-10-14T17:00'}).status_code, 200)
        previous = self.workspace.read()['imported']
        additional = read_report_files([('new.xlsx', excel([row('added', date='10/13/2026 09:00 AM')]))])
        combined = combine_sources(import_sources(previous) + import_sources(additional))
        retained = preserve_sheet_edits(previous, combined)
        actual = next(r for r in retained['rows'] if r['Order Number'] == '001' and r['Source workbook'] == 'full.xlsx')
        self.assertEqual((actual['Borrower'], actual['Status']), ('Keep local', 'Manual review'))


if __name__ == '__main__':
    unittest.main()

"""Publish derived reporting views without replacing canonical trackers or receipts."""
from collections import Counter
import hashlib
import json
import uuid
import pandas as pd

from report_metrics import report_matrices
from sheets_writer import coordinated

RECEIPT = '__TvTracker_ReportSource'
CHART_TITLE = 'Tv Tracker · Received and completed against capacity'


def capacity_chart_requests(book, sheet_id, capacity):
    """Replace only this app's chart; preserve any user-created charts."""
    metadata = book.fetch_sheet_metadata(params={'fields': 'sheets(charts(chartId,spec(title)))'})
    requests = [{'deleteEmbeddedObject': {'objectId': chart['chartId']}}
                for sheet in metadata.get('sheets', []) for chart in sheet.get('charts', [])
                if chart.get('spec', {}).get('title') == CHART_TITLE]
    if not capacity['daily']:
        return requests
    start = len(capacity['monthly']) + 3
    end = start + 1 + len(capacity['daily'])
    def source(column):
        return {'sourceRange': {'sources': [
            {'sheetId': sheet_id, 'startRowIndex': start, 'endRowIndex': end, 'startColumnIndex': column, 'endColumnIndex': column + 1}]}}
    requests.append({'addChart': {'chart': {'spec': {'title': CHART_TITLE,
        'basicChart': {'chartType': 'COMBO', 'legendPosition': 'BOTTOM_LEGEND', 'headerCount': 1,
            'domains': [{'domain': source(0)}],
            'series': [{'series': source(column), 'type': kind, 'targetAxis': 'LEFT_AXIS'}
                       for column, kind in ((1, 'COLUMN'), (2, 'COLUMN'), (7, 'LINE'), (8, 'LINE'))],
            'axis': [{'position': 'BOTTOM_AXIS', 'title': 'Received date'}, {'position': 'LEFT_AXIS', 'title': 'Orders'}]}},
        'position': {'overlayPosition': {'anchorCell': {'sheetId': sheet_id, 'rowIndex': end + 2, 'columnIndex': 0},
                                         'widthPixels': 1080, 'heightPixels': 400}}}}})
    return requests


def view_matrices(snapshot, capacity):
    from monthly_views import view_name
    tables = report_matrices(snapshot, capacity)
    combined = snapshot['sheets']['All Products']
    headers = combined['columns']
    tables['All Products'] = [headers] + [[r.get(c, '') for c in headers] for r in combined['rows']]
    for month in snapshot['reports']['monthly']:
        if month['Month'] == 'Undated':
            continue
        for full in (True, False):
            rows = [r for r in month['rows'] if (' '.join(str(r.get('Product', '')).casefold().split()) in ('full title', 'full search')) == full]
            tables[view_name(full, month['Month'])] = [headers] + [[r.get(c, '') for c in headers] for r in rows]
    status = Counter(r.get('Status', '') or '(Blank)' for r in combined['rows'])
    total = sum(status.values())
    tables['Status Report'] = [['Status', 'Orders', 'Share']] + [[s, n, n / total if total else 0] for s, n in status.items()]
    return tables


@coordinated('Publish reports')
def publish_reports(book, snapshot, capacity):
    from sheets_repository import Batch
    from datatrace_sync import write_sheet_batch, sheet_color
    from monthly_production import tab_identity
    from monthly_views import view_identity
    from tracker_sync import TRACKERS, OLD_FULL, OLD_REMAINING
    tables = view_matrices(snapshot, capacity)
    digest = hashlib.sha256(json.dumps(tables, sort_keys=True, default=str).encode()).hexdigest()
    publication = uuid.uuid4().hex
    batch = Batch(book)
    existing = set(batch.sheets)
    # Preserve the previous raw All Products archive before its first conversion.
    if 'All Products' in existing and '__TvTracker_OriginalAllProducts' not in existing:
        source = batch.sheets['All Products']
        backup_id = batch.next_id
        batch.next_id += 1
        batch.requests.append({'duplicateSheet': {'sourceSheetId': source.id, 'newSheetId': backup_id, 'newSheetName': '__TvTracker_OriginalAllProducts'}})
        batch.requests.append({'updateSheetProperties': {'properties': {'sheetId': backup_id, 'hidden': True}, 'fields': 'hidden'}})
    for title, values in tables.items():
        batch.replace_view(title, pd.DataFrame(values[1:], columns=values[0]))
        sheet = batch.sheets[title]
        # Sheets does not allow freezing every row of a one-row empty view.
        batch.capacity(sheet, max(2, len(values)), len(values[0]))
        if title == 'Capacity Report':
            # Keep the chart beneath the table and inside the allocated grid.
            batch.capacity(sheet, len(values) + 25, len(values[0]))
        batch.requests.extend([
            {'updateSheetProperties': {'properties': {'sheetId': sheet.id, 'hidden': False, 'gridProperties': {'frozenRowCount': 1}}, 'fields': 'hidden,gridProperties.frozenRowCount'}},
            {'repeatCell': {'range': {'sheetId': sheet.id, 'startRowIndex': 0, 'endRowIndex': 1, 'endColumnIndex': len(values[0])}, 'cell': {'userEnteredFormat': {'backgroundColor': sheet_color('#DDEBE7'), 'textFormat': {'bold': True, 'foregroundColor': sheet_color('#173D35')}}}, 'fields': 'userEnteredFormat'}}])
        if title in ('Daily Orders', 'Monthly Orders', 'Capacity Report'):
            batch.requests.extend([
                {'updateDimensionProperties': {'range': {'sheetId': sheet.id, 'dimension': 'COLUMNS', 'startIndex': 0, 'endIndex': 9}, 'properties': {'pixelSize': 126}, 'fields': 'pixelSize'}},
                {'updateDimensionProperties': {'range': {'sheetId': sheet.id, 'dimension': 'ROWS', 'startIndex': 0, 'endIndex': 1}, 'properties': {'pixelSize': 42}, 'fields': 'pixelSize'}},
                {'repeatCell': {'range': {'sheetId': sheet.id, 'endRowIndex': len(values), 'endColumnIndex': 9}, 'cell': {'userEnteredFormat': {'wrapStrategy': 'WRAP', 'verticalAlignment': 'MIDDLE'}}, 'fields': 'userEnteredFormat.wrapStrategy,userEnteredFormat.verticalAlignment'}}])
            if len(values) > 1:
                batch.requests.append({'repeatCell': {'range': {'sheetId': sheet.id, 'startRowIndex': 1, 'endRowIndex': len(values), 'startColumnIndex': 1, 'endColumnIndex': 9}, 'cell': {'userEnteredFormat': {'horizontalAlignment': 'RIGHT'}}, 'fields': 'userEnteredFormat.horizontalAlignment'}})
    batch.requests.extend(capacity_chart_requests(book, batch.sheets['Capacity Report'].id, capacity))
    if capacity['daily']:
        start = len(capacity['monthly']) + 3
        identity = batch.sheets['Capacity Report'].id
        batch.requests.extend([
            {'updateDimensionProperties': {'range': {'sheetId': identity, 'dimension': 'ROWS', 'startIndex': start, 'endIndex': start + 1}, 'properties': {'pixelSize': 42}, 'fields': 'pixelSize'}},
            {'repeatCell': {'range': {'sheetId': identity, 'startRowIndex': start, 'endRowIndex': start + 1, 'endColumnIndex': 9},
                            'cell': {'userEnteredFormat': {'backgroundColor': sheet_color('#DDEBE7'), 'textFormat': {'bold': True, 'foregroundColor': sheet_color('#173D35')}, 'wrapStrategy': 'WRAP', 'verticalAlignment': 'MIDDLE', 'horizontalAlignment': 'LEFT'}}, 'fields': 'userEnteredFormat'}},
            {'setBasicFilter': {'filter': {'range': {'sheetId': identity, 'startRowIndex': start, 'endRowIndex': start + 1 + len(capacity['daily']), 'endColumnIndex': 9}}}}])
    # Owned implementation tabs are hidden, never deleted. Unknown user tabs stay visible.
    owned = set(TRACKERS) | {OLD_FULL, OLD_REMAINING, 'Sheet1', 'Full Title', 'Remaining Products', 'Overview', 'Monthly report', 'Changes', 'Needs review', 'Preview History', 'Default import details'}
    visible_ids = {batch.sheets[title].id for title in tables}
    for title in existing - set(tables):
        if batch.sheets[title].id in visible_ids:
            continue
        if title in owned or title.startswith(('__DataTrace_', '__TvTracker_')) or tab_identity(title) or view_identity(title):
            batch.requests.append({'updateSheetProperties': {'properties': {'sheetId': batch.sheets[title].id, 'hidden': True}, 'fields': 'hidden'}})
    receipt_sheet = batch.sheet(RECEIPT)
    batch.capacity(receipt_sheet, 2, 4)
    batch.cells(receipt_sheet, 0, 0, [['Digest', 'Source', 'Updated', 'Publication'], [digest, snapshot.get('source', ''), snapshot.get('updated_at', ''), publication]])
    batch.requests.append({'updateSheetProperties': {'properties': {'sheetId': receipt_sheet.id, 'hidden': True}, 'fields': 'hidden'}})

    def committed():
        from gspread.exceptions import WorksheetNotFound
        try:
            values = book.worksheet(RECEIPT).get_all_values()
            return len(values) > 1 and len(values[1]) > 3 and values[1][0] == digest and values[1][3] == publication
        except WorksheetNotFound:
            return False
        except Exception as error:
            raise ValueError('The publication receipt could not be read. Publishing is paused because the write may have completed. Restore the connection before retrying; saved captures are retained.') from error

    write_sheet_batch(book, receipt_sheet, batch.requests, verify_commit=committed)
    if not committed():
        raise RuntimeError('Report publication could not be verified. Retry publishing; the report source is retained locally.')
    return {'digest': digest, 'daily_gid': batch.sheets['Daily Orders'].id,
            'daily_dates': [r[0] for r in tables['Daily Orders'][1:]], 'tabs': list(tables)}


def export_reports(snapshot, capacity):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.chart import BarChart, LineChart, Reference
    from datatrace_sync import status_color
    from tracker_formatting import foreground
    stream = BytesIO()
    book = Workbook()
    book.remove(book.active)
    for title, values in view_matrices(snapshot, capacity).items():
        sheet = book.create_sheet(title[:31])
        for row in values:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='DDEBE7')
            cell.font = Font(bold=True, color='173D35')
        if 'Status' in values[0]:
            index = values[0].index('Status')
            for number, row in enumerate(values[1:], 2):
                color = status_color(row[index]).lstrip('#')
                for cell in sheet[number]:
                    cell.fill = PatternFill('solid', fgColor=color)
                    cell.font = Font(color=foreground('#' + color).lstrip('#'))
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(14, max(len(str(c.value or '')) for c in column[:100]) + 2))
        if title == 'Capacity Report' and capacity['daily']:
            chart = BarChart()
            chart.title = 'Received and completed against capacity'
            start = len(capacity['monthly']) + 5
            for cell in sheet[start - 1]:
                cell.fill = PatternFill('solid', fgColor='DDEBE7')
                cell.font = Font(bold=True, color='173D35')
            sheet.auto_filter.ref = f'A{start - 1}:I{sheet.max_row}'
            lines = LineChart()
            for column in (2, 3, 8, 9):
                target = chart if column < 8 else lines
                target.add_data(Reference(sheet, min_col=column, min_row=start, max_row=sheet.max_row))
                target.series[-1].title = __import__('openpyxl').chart.series.SeriesLabel(strRef=None, v=values[0][column-1])
            chart.set_categories(Reference(sheet, min_col=1, min_row=start, max_row=sheet.max_row))
            lines.set_categories(Reference(sheet, min_col=1, min_row=start, max_row=sheet.max_row))
            chart += lines
            sheet.add_chart(chart, f'A{sheet.max_row + 3}')
    book.save(stream)
    stream.seek(0)
    return stream

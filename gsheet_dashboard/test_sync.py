import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

import datatrace_sync as sync


class SyncTests(unittest.TestCase):
    def test_status_rules_replace_original_values_once(self):
        source = pd.DataFrame({'Task Status': [' Workflow Suspended ', 'Available', 'Other'],
                               'Is Available': [False, True, False]})
        result = sync.apply_status_rules(source, [
            {'source': 'Workflow Suspended', 'target': 'Available'},
            {'source': 'Available', 'target': 'Search in Progress'}])
        self.assertEqual(result['Task Status'].tolist(), ['Available', 'Search in Progress', 'Other'])
        self.assertEqual(result['Is Available'].tolist(), [True, False, False])
        self.assertEqual(source.iloc[1]['Task Status'], 'Available')
        with self.assertRaises(ValueError):
            sync.apply_status_rules(source, [{'source': 'Available', 'target': 'A'},
                                            {'source': 'available', 'target': 'B'}])

    def test_queue_status_cells_have_colors_and_filter(self):
        df = pd.DataFrame({'Status': ['Awaiting for Clarification', 'Search in Progress'], 'Order': ['001', '002']})
        book, sheet = Mock(), Mock(id=0, row_count=20, col_count=10, title='Queue')
        sheet.get_all_values.return_value = [list(df.columns)] + df.values.tolist()
        sync.sync_dataframe(df, (book, sheet), validate=False, color_status=True)
        requests = book.batch_update.call_args.args[0]['requests']
        color_update = requests[-2]['updateCells']
        self.assertEqual(color_update['range']['startColumnIndex'], 0)
        self.assertEqual(color_update['range']['endColumnIndex'], len(df.columns))
        self.assertEqual(color_update['range']['startRowIndex'], 1)
        self.assertEqual(color_update['fields'], 'userEnteredFormat.backgroundColor')
        self.assertEqual(color_update['rows'][0]['values'][0], color_update['rows'][0]['values'][1])
        self.assertEqual(color_update['rows'][0]['values'][0]['userEnteredFormat']['backgroundColor'], sync.sheet_color('#a66ad3'))
        self.assertEqual(requests[-1]['setBasicFilter']['filter']['range']['endRowIndex'], 3)

    def frame(self):
        return pd.DataFrame([{'Arrival Time': '09/22/2026 01:00 PM',
                              'Task Status': 'Available', 'Parcel ID': '00123',
                              'Comment': '=1+1', 'Time Since Arrival': '2d 3h 30m',
                              'SLA Expiration*': '-1h 0m'}])

    def test_powerbi_and_literal_exports(self):
        df = sync.powerbi_table(self.frame(), '2026-09-22T12:00:00+00:00')
        self.assertEqual(df.loc[0, 'Queue Age Hours'], 51.5)
        self.assertEqual(df.loc[0, 'Arrival Date'], '2026-09-22')
        self.assertEqual(df.loc[0, 'SLA Status'], 'Overdue')
        self.assertTrue(df.loc[0, 'Is Available'])
        with tempfile.TemporaryDirectory() as folder:
            excel, csv = sync.export_to_excel_and_csv(df, Path(folder) / 'queue')
            read = pd.read_csv(csv, dtype=str, keep_default_na=False)
            self.assertEqual(read.loc[0, 'Parcel ID'], '00123')
            ws = load_workbook(excel).active
            self.assertIn('DataTraceQueue', ws.tables)
            self.assertEqual(ws['D2'].value, '=1+1')
            self.assertEqual(ws['D2'].data_type, 's')

    def test_report_frames_follow_sample_headers_and_product_split(self):
        source = pd.DataFrame([
            {'Order Number': 'A-1', 'Product': 'Full Title', 'St': 'NJ', 'Client': 'ONE',
             'Online/ Ground': 'Ground', 'Task Status': 'Available', 'Arrival Date': '2026-09-22',
             'Queue Age Hours': 4.5, 'OPON': '101'},
            {'Order Number': 'A-2', 'Product': 'Current Owner', 'St': 'NY', 'Client': 'TWO',
             'Online/ Ground': 'Online', 'Task Status': 'Task Suspended', 'Arrival Date': '2026-09-23',
             'Queue Age Hours': 2.0, 'OPON': '102'},
        ])
        all_products, full_title, remaining = sync.report_frames(source)
        self.assertEqual(list(all_products.columns), [name for name, _ in sync.ALL_PRODUCT_FIELDS])
        self.assertEqual(list(full_title.columns), [name for name, _ in sync.FULL_TITLE_FIELDS])
        self.assertEqual(list(remaining.columns), [name for name, _ in sync.REMAINING_PRODUCT_FIELDS])
        self.assertEqual(list(remaining.columns), [name for name, _ in sync.FULL_TITLE_FIELDS])
        self.assertEqual(all_products['Order Number'].tolist(), ['A-1', 'A-2'])
        self.assertEqual(all_products['Originator Product Order Number'].tolist(), ['101', '102'])
        self.assertEqual(full_title['Order Number'].tolist(), ['A-1'])
        self.assertEqual(remaining['Order Number'].tolist(), ['A-2'])
        self.assertEqual(remaining['Product'].tolist(), ['Current Owner'])
        self.assertEqual(remaining['Status'].tolist(), ['Task Suspended'])

    def test_report_frames_with_selected_products(self):
        source = pd.DataFrame([
            {'Order Number': 'A-1', 'Product': 'Full Title', 'St': 'NJ', 'Client': 'ONE',
             'Online/ Ground': 'Ground', 'Task Status': 'Available'},
            {'Order Number': 'A-2', 'Product': 'Current Owner', 'St': 'NY', 'Client': 'TWO',
             'Online/ Ground': 'Online', 'Task Status': 'Task Suspended'},
            {'Order Number': 'A-3', 'Product': 'Tax Search', 'St': 'TX', 'Client': 'THREE',
             'Online/ Ground': 'Online', 'Task Status': 'Available'},
        ])
        _, _, remaining_filtered = sync.report_frames(source, selected_products=['Tax Search'])
        self.assertEqual(remaining_filtered['Order Number'].tolist(), ['A-3'])
        self.assertEqual(remaining_filtered['Product'].tolist(), ['Tax Search'])

        _, _, remaining_all = sync.report_frames(source, selected_products=None)
        self.assertEqual(remaining_all['Order Number'].tolist(), ['A-2', 'A-3'])

    def test_status_report_counts_and_colors(self):
        source = pd.DataFrame([
            {'Task Status': 'Completed and delivered'},
            {'Task Status': 'Completed and delivered'},
            {'Task Status': 'Typing in Progress'},
            {'Task Status': 'Awaiting for Clarification'},
        ])
        report = sync.status_report_frame(source)
        self.assertEqual(list(report.columns), ['Status', 'Orders', 'Share'])
        self.assertEqual(report.loc[0].to_dict(), {
            'Status': 'Completed and delivered', 'Orders': 2, 'Share': '50.0%'
        })
        self.assertEqual(sync.status_color('Typing in Progress'), '#00b050')
        self.assertEqual(sync.status_color('Awaiting for Clarification'), '#a66ad3')

    def test_atomic_replace_gid_zero_and_literals(self):
        df = self.frame()
        book, sheet = Mock(), Mock(id=0, row_count=1, col_count=1, title='Actual tab')
        sheet.get_all_values.return_value = [list(df.columns)] + df.values.tolist()
        self.assertEqual(sync.sync_dataframe(df, (book, sheet)), 'Actual tab')
        requests = book.batch_update.call_args.args[0]['requests']
        self.assertEqual(len(requests), 2)
        self.assertEqual(requests[0]['updateSheetProperties']['properties']['gridProperties'],
                         {'rowCount': 2, 'columnCount': len(df.columns)})
        update = requests[-1]['updateCells']
        self.assertEqual(update['range'], {'sheetId': 0, 'startRowIndex': 0, 'endRowIndex': 2,
                                          'startColumnIndex': 0, 'endColumnIndex': len(df.columns)})
        self.assertEqual(update['fields'], 'userEnteredValue')
        self.assertEqual(update['rows'][1]['values'][2]['userEnteredValue'], {'stringValue': '00123'})
        self.assertEqual(update['rows'][1]['values'][3]['userEnteredValue'], {'stringValue': '=1+1'})

    def test_status_report_sync_applies_row_backgrounds(self):
        df = pd.DataFrame([{'Status': 'Ready to send', 'Orders': 1, 'Share': '100.0%'}])
        book, sheet = Mock(), Mock(id=7, row_count=1, col_count=1, title='Status Report')
        sheet.get_all_values.return_value = [list(df.columns)] + df.values.tolist()
        sync.sync_dataframe(df, (book, sheet), validate=False, row_backgrounds=[None, '#ffff00'])
        update = book.batch_update.call_args.args[0]['requests'][-1]['updateCells']
        self.assertEqual(update['fields'], 'userEnteredValue,userEnteredFormat.backgroundColor')
        self.assertIn('userEnteredFormat', update['rows'][1]['values'][0])

    def test_empty_data_never_writes(self):
        book = Mock()
        with self.assertRaises(ValueError):
            sync.sync_dataframe(self.frame().iloc[:0], (book, Mock()))
        book.batch_update.assert_not_called()

    def test_failed_scrape_never_uploads_old_files(self):
        with tempfile.TemporaryDirectory() as folder, patch.object(sync, 'BASE_DIR', Path(folder)), \
                patch.dict(sync.os.environ, {'DATATRACE_USERNAME': 'test', 'DATATRACE_PASSWORD': 'test'}), \
                patch.object(sync.subprocess, 'run', return_value=Mock(returncode=1)), \
                patch.object(sync, 'sync_workbook') as upload, \
                patch.object(sync, 'export_to_excel_and_csv') as export:
            result = sync.run_sync()
            self.assertEqual(result['scrape'], 'failed')
            upload.assert_not_called()
            export.assert_not_called()

    def test_extraction_retains_export_without_google_sync(self):
        def scrape(*args, **kwargs):
            Path(kwargs['env']['DATATRACE_OUTPUT_JSON']).write_text(
                self.frame().to_json(orient='records'), encoding='utf-8')
            return Mock(returncode=0)
        with tempfile.TemporaryDirectory() as folder, patch.object(sync, 'BASE_DIR', Path(folder)), \
                patch.dict(sync.os.environ, {'DATATRACE_USERNAME': 'test', 'DATATRACE_PASSWORD': 'test'}), \
                patch.object(sync.subprocess, 'run', side_effect=scrape), \
                patch.object(sync, 'sync_workbook') as upload:
            (Path(folder) / 'sync_status.json').write_text(json.dumps({'last_success_at': 'previous'}))
            result = sync.run_sync(auto_sync=False)
            self.assertEqual(result['scrape'], 'success')
            self.assertEqual(result['google_sheet'], 'not_synced')
            upload.assert_not_called()
            self.assertTrue((Path(folder) / 'queue_data_sheet2.xlsx').exists())

    def test_verification_failure_is_not_success(self):
        book, sheet = Mock(), Mock(id=0, row_count=100, col_count=30)
        sheet.get_all_values.return_value = []
        with self.assertRaisesRegex(RuntimeError, 'verification failed'):
            sync.sync_dataframe(self.frame(), (book, sheet))


if __name__ == '__main__':
    unittest.main()

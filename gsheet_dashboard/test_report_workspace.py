from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from preview_store import PreviewStore
from report_metrics import counts
from report_workspace import ReportWorkspace, read_workbooks
from report_publishing import export_reports, publish_reports
from workspace_backup import make_backup, restore_backup
from test_monthly_production import AtomicBook


def workbook(rows, title='Orders'):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    headers = list(dict.fromkeys(c for row in rows for c in row))
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(c, '') for c in headers])
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def order(number, status='Search In Progress', **extra):
    return {'Order Number': number, 'Product': 'Full Title', 'Status': status,
            'In-Time': '10/01/2026 09:00 AM', **extra}


class ReportWorkspaceTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.store = PreviewStore(Path(directory.name) / 'previews')
        self.workspace = ReportWorkspace(self.store)

    def test_status_buckets_partition_received_and_sla_requires_evidence(self):
        rows = [order('1', 'Completed and Delivered', **{'Out Time': '10/01/2026 10:00 AM', 'SLA Expiration': '10/01/2026 10:00 AM'}),
                order('2', ' CANCELLED '), order('3', 'Awaiting for Clarification'), order('4', ' Assign to ABS '),
                order('5', 'Need to assign ABS'), order('6', 'Task Suspended'),
                order('7', 'Completed and Delivered', **{'Out Time': '10/01/2026 11:00 AM', 'SLA Expiration': '10/01/2026 10:00 AM', 'Completion Evidence': 'Inferred from queue absence'})]
        result = counts(rows)
        self.assertEqual(sum(result[c] for c in ('Completed', 'Clarification', 'Cancelled', 'Vendor Pending', 'In-House Pending')), 7)
        self.assertEqual(result['Vendor Pending'], 1)
        self.assertEqual(result['In-House Pending'], 2)
        self.assertEqual(result['On time SLA'], 1)
        self.assertEqual(result['Missed SLA'], 0)
        self.assertEqual(result['SLA Unclassified'], 1)

    def test_multiple_workbooks_deduplicate_and_conflicts_fail_before_activation(self):
        raw = workbook([order('001')])
        saved = self.workspace.save_import([('a.xlsx', raw), ('b.xlsx', raw)])
        self.assertEqual(saved['count'], 1)
        self.assertEqual(saved['duplicates_removed'], 1)
        self.assertEqual(self.workspace.preferences()['source'], 'tracker')
        with self.assertRaisesRegex(ValueError, 'conflicting duplicate'):
            self.workspace.save_import([('a.xlsx', raw), ('b.xlsx', workbook([order('001', 'Cancelled')]))])
        self.assertEqual(len(self.workspace.imports()), 1)
        self.assertEqual(self.store.list(), [])

    def test_invalid_schema_dates_and_formulas_are_rejected(self):
        for row in ({'Order Number': 'x'}, order('x', **{'In-Time': 'nonsense'}), order('x', Client='=HYPERLINK("bad")')):
            with self.assertRaises(ValueError):
                read_workbooks([('bad.xlsx', workbook([row]))])

    def test_duplicate_formatting_and_complementary_blank_fields_merge(self):
        first = order('001', Client='Northstar', Comments='')
        second = order('001', **{'In-Time': '2026-10-01 09:00:00', 'Client': ' northstar ', 'Comments': 'Verified'})
        value = read_workbooks([('a.xlsx', workbook([first])), ('b.xlsx', workbook([second]))])
        self.assertEqual(len(value['rows']), 1)
        self.assertEqual(value['rows'][0]['Comments'], 'Verified')

    def test_conflict_review_is_explicit_and_bound_to_the_whole_upload(self):
        from report_workspace import ImportConflictError
        uploads = [('a.xlsx', workbook([order('001')])), ('b.xlsx', workbook([order('001', 'Cancelled')]))]
        with self.assertRaises(ImportConflictError) as caught:
            self.workspace.save_import(uploads)
        conflict = caught.exception.conflicts[0]
        self.assertEqual(conflict['order'], '001')
        self.assertIn('Status', conflict['columns'])
        self.assertIn('row 2', conflict['options'][1]['source'])
        choice = {'001': conflict['options'][1]['token']}
        result = self.workspace.save_import(uploads, choice)
        self.assertEqual(result['conflicts_resolved'], 1)
        self.assertEqual(self.workspace.imported_snapshot(result['id'])['reports']['daily'][0]['Cancelled'], 1)
        uploads[1] = ('b.xlsx', workbook([order('001', 'Completed and Delivered')]))
        with self.assertRaises(ImportConflictError):
            self.workspace.save_import(uploads, choice)
        self.assertEqual(len(self.workspace.imports()), 1)

    def test_timing_review_explains_uncertainty_without_changing_source_data(self):
        from report_metrics import enrich_reports
        rows = [order('1', 'Completed'), order('2', 'Completed', **{'Out Time': '10/01/2026'}),
                order('3', 'Completed', **{'Out Time': '10/01/2026 11:00 AM'}), order('4', 'Cancelled')]
        reports = {'daily': [{'rows': rows, 'columns': list(rows[0])}], 'monthly': []}
        enrich_reports(reports)
        day = reports['daily'][0]
        self.assertEqual(day['SLA Unclassified'], 3)
        self.assertEqual(sum(day['SLA Review Reasons'].values()), 3)
        self.assertEqual(day['rows'][0]['SLA Review Reason'], 'Out Time is missing')
        self.assertIn('date alone', day['rows'][1]['SLA Review Reason'])
        self.assertEqual(day['rows'][2]['SLA Review Reason'], 'SLA Expiration is missing')
        self.assertNotIn('SLA Review Reason', rows[0])

    def test_import_reports_are_independent_and_capacity_is_editable(self):
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1'), order('2', 'Cancelled')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        self.assertEqual(value['reports']['daily'][0]['Received'], 2)
        self.assertEqual(value['reports']['daily'][0]['Cancelled'], 1)
        self.workspace.set_targets('2026-10-01', 10, 15)
        report = self.workspace.capacity(value)
        self.assertEqual(report['monthly'][0]['Capacity'], 10)
        self.assertEqual(report['ytd']['Ext capacity'], 15)
        with self.assertRaises(ValueError):
            self.workspace.set_targets('default', 20, 10)

    def test_export_uses_same_counts_and_safe_literal_cells(self):
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1'), order('2', 'Assign to ABS')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        report = load_workbook(export_reports(value, self.workspace.capacity(value)))
        self.assertEqual(report['Daily Orders']['B2'].value, 2)
        self.assertEqual(report['Daily Orders']['F2'].value, 1)
        self.assertIn('Capacity Report', report.sheetnames)
        self.assertIn('Full_search_OCT_2026', report.sheetnames)
        self.assertEqual(len(report['Capacity Report']._charts), 1)

    def test_publish_is_atomic_and_preserves_canonical_trackers(self):
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        book = AtomicBook()
        book.add('Sheet1', [order('old')])
        original = deepcopy(book.worksheet('Sheet1').values)
        result = publish_reports(book, value, self.workspace.capacity(value))
        self.assertEqual(book.worksheet('Sheet1').values, original)
        self.assertEqual(result['daily_dates'], ['2026-10-01'])
        self.assertEqual(book.worksheet('Daily Orders').values[1][1], 1)

    def test_backup_includes_imports_and_detects_modified_contents(self):
        self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        raw = make_backup(self.store).getvalue()
        restore_backup(self.store, raw)
        self.assertEqual(len(self.workspace.imports()), 1)
        stream = BytesIO()
        with ZipFile(BytesIO(raw)) as source, ZipFile(stream, 'w', ZIP_DEFLATED) as output:
            for name in source.namelist():
                value = source.read(name)
                if name == 'manifest.json':
                    data = json.loads(value)
                    data['sha256']['previews.sqlite'] = 'wrong'
                    value = json.dumps(data).encode()
                output.writestr(name, value)
        with self.assertRaisesRegex(ValueError, 'integrity'):
            restore_backup(self.store, stream.getvalue())

    def test_old_identical_receipt_cannot_confirm_a_failed_new_publication(self):
        from requests.exceptions import Timeout
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        book = AtomicBook()
        publish_reports(book, value, self.workspace.capacity(value))
        with patch.object(book, 'batch_update', side_effect=Timeout('request never committed')) as write, patch('datatrace_sync.time.sleep'):
            with self.assertRaisesRegex(RuntimeError, 'three attempts'):
                publish_reports(book, value, self.workspace.capacity(value))
        self.assertEqual(write.call_count, 3)

    def test_new_report_sheets_allocate_chart_anchor_and_complete_receipt(self):
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        book = AtomicBook()
        publish_reports(book, value, self.workspace.capacity(value))
        requests = book.batches[-1]['requests']
        capacity_id = book.worksheet('Capacity Report').id
        receipt_id = book.worksheet('__TvTracker_ReportSource').id
        properties = [item['updateSheetProperties']['properties'] for item in requests if 'updateSheetProperties' in item]
        anchor = next(item['addChart']['chart']['position']['overlayPosition']['anchorCell'] for item in requests if 'addChart' in item)
        self.assertEqual(anchor['columnIndex'], 0)
        self.assertTrue(any(p['sheetId'] == capacity_id and p.get('gridProperties', {}).get('rowCount', 0) > anchor['rowIndex'] for p in properties))
        self.assertTrue(any(p['sheetId'] == receipt_id and p.get('gridProperties', {}).get('columnCount', 0) >= 4 and p['gridProperties']['rowCount'] >= 2 for p in properties))
        chart = next(item['addChart']['chart']['spec']['basicChart'] for item in requests if 'addChart' in item)
        for data in [chart['domains'][0]['domain']] + [series['series'] for series in chart['series']]:
            ranges = data['sourceRange']['sources']
            self.assertEqual(len(ranges), 1)
            self.assertEqual((ranges[0]['startRowIndex'], ranges[0]['endRowIndex']), (4, 6))
        matrix = book.worksheet('Capacity Report').values
        self.assertEqual(matrix[4][0], 'Date')
        self.assertEqual(matrix[5][0], '2026-10-01')
        exported = load_workbook(export_reports(value, self.workspace.capacity(value)))['Capacity Report']
        self.assertEqual(exported['A5'].value, 'Date')
        self.assertEqual(exported['A6'].value, '2026-10-01')
        self.assertIn('$B$6', exported._charts[0].series[0].val.numRef.f)

    def test_monthly_case_aliases_publish_without_duplicate_tabs(self):
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        book = AtomicBook()
        full = book.add('Full_Search_OCT_2026', [order('old')])
        remaining = book.add('Remaining_Search_OCT_2026')
        publish_reports(book, value, self.workspace.capacity(value))
        self.assertEqual(book.worksheet(full.title).id, full.id)
        self.assertEqual(book.worksheet(remaining.title).id, remaining.id)
        names = [s.title for s in book.sheets]
        self.assertNotIn('Full_search_OCT_2026', names)
        self.assertNotIn('Remaining_OCT_2026', names)
        matrix = book.worksheet(full.title).values
        index = matrix[0].index('Order Number')
        self.assertIn('1', [row[index] for row in matrix[1:] if len(row) > index])

    def test_unreadable_receipt_after_lost_reply_does_not_replay_write(self):
        from requests.exceptions import Timeout, ConnectionError
        saved = self.workspace.save_import([('one.xlsx', workbook([order('1')]))])
        value = self.workspace.imported_snapshot(saved['id'])
        book = AtomicBook()
        with patch.object(book, 'batch_update', side_effect=Timeout('lost reply')) as write, \
                patch.object(book, 'worksheet', side_effect=ConnectionError('receipt unavailable')):
            with self.assertRaisesRegex(ValueError, 'receipt could not be read'):
                publish_reports(book, value, self.workspace.capacity(value))
        self.assertEqual(write.call_count, 1)

    def test_backup_does_not_create_an_unrestorable_oversize_archive(self):
        with patch('workspace_backup.MAX_EXPANDED_BYTES', 1):
            with self.assertRaisesRegex(ValueError, 'recovery size'):
                make_backup(self.store)


if __name__ == '__main__':
    unittest.main()

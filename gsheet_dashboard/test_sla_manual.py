"""Manual SLA status/order workflows do not mutate captures or double count orders."""
from copy import deepcopy
from io import BytesIO
import unittest
from openpyxl import load_workbook
from sla_status import queue_rule
from sla_manual import custom_statuses
import test_sla_queue_rules as queue_tests


class SlaManualTests(unittest.TestCase):
    setUp = queue_tests.SlaQueueRulesTests.setUp
    capture = queue_tests.SlaQueueRulesTests.capture

    def initial(self):
        self.capture(1, [{'Order Number': '001', 'Borrower': 'Original', 'SLA Expiration*': '1d',
                          'Task Name': 'TypingModule', 'Task Status': 'In Progress'},
                         {'Order Number': '002', 'SLA Expiration*': '-2h', 'St': 'LA'}])
        return self.client.get('/api/sla-status?preview_id=1').json

    def lookup(self, number='001', preview=1):
        return self.client.get('/api/sla-order', query_string={'preview_id': preview, 'order_number': number}).json

    def save(self, number='001', status='Completed and Delivered', fields=None, lookup=None, preview=1):
        return self.client.post('/api/sla-order', json={'preview_id': preview, 'order_number': number,
            'status': status, 'fields': fields or {}, 'expected_revision': (lookup or {}).get('revision')})

    def test_typing_rule_is_exact_and_preserves_priority(self):
        self.assertEqual(queue_rule({'TaskName': 'typingModule', 'TaskStatus': 'In\nProgress'})[1], 'Typing in Progress')
        self.assertNotEqual(queue_rule({'TaskName': 'TypingModule', 'TaskStatus': 'Available'})[1], 'Typing in Progress')
        self.assertNotEqual(queue_rule({'TaskName': 'Search', 'TaskStatus': 'In Progress'})[1], 'Typing in Progress')
        self.assertEqual(queue_rule({'Task Name': 'TypingModule', 'Task Status': 'In Progress',
            'ETA Comments': 'ON HOLD'})[1], 'Waiting for Effective Date')
        self.assertEqual(queue_rule({'Task Name': 'TypingModule', 'Task Status': 'In Progress',
            'Product': 'ClientReview'})[1], 'SearchFix')
        self.assertEqual(self.initial()['summary'][-1]['Status'], 'Typing in Progress')

    def test_empty_custom_status_persists_without_changing_totals_and_can_rename(self):
        before = self.initial()
        for name in ('Need funds', ' need FUNDS '):
            response = self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': name})
            self.assertEqual(response.status_code, 200, response.json)
        report = self.client.get('/api/sla-status').json
        self.assertEqual(report['totals'], before['totals'])
        self.assertEqual([row['Grand Total'] for row in report['summary'] if row['Status'] == 'Need funds'], [0])
        renamed = self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'Need logins', 'previous': 'Need funds'})
        self.assertEqual(renamed.status_code, 200, renamed.json)
        self.assertEqual(custom_statuses(self.workspace.read(), 1), ['Need logins'])
        self.assertEqual(self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'Grand Total'}).status_code, 409)
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
        self.assertIn('Need logins', [row[0].value for row in book['Sheet2'].iter_rows()])

    def test_existing_order_moves_once_and_details_override_only_this_preview(self):
        self.initial()
        before = deepcopy(self.store.history())
        lookup = self.lookup(' 001 ')
        self.assertTrue(lookup['exists'])
        self.assertEqual(lookup['row']['Status'], 'Typing in Progress')
        response = self.save(fields={'Borrower': 'Corrected', 'State': 'NY', 'SLA Expiration': '2026-10-20T17:00'}, lookup=lookup)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertFalse(response.json['created'])
        report = self.client.get('/api/sla-status').json
        self.assertEqual(report['totals']['Grand Total'], 2)
        self.assertEqual(report['manual_order_count'], 0)
        row = next(row for row in report['rows'] if row['Order Number'] == '001')
        self.assertEqual((row['Status'], row['Borrower'], row['State']), ('Completed and Delivered', 'Corrected', 'NY'))
        self.assertEqual(row['SLA Expiration*'], row['SLA Expiration'])
        self.assertEqual(self.store.history(), before)
        self.capture(2, self.store.get(1)['rows'])
        other = self.client.get('/api/sla-status?preview_id=2').json
        self.assertEqual(other['rows'][0]['Status'], 'Typing in Progress')
        self.assertEqual(other['rows'][0]['Borrower'], 'Original')
        self.assertEqual(self.save(fields={'Borrower': 'Stale'}, lookup=lookup).status_code, 409)
        self.assertNotIn('sla_preview_order_edits', self.workspace.public())

    def test_new_orders_need_valid_deadline_are_deduplicated_and_exported(self):
        self.initial()
        self.assertFalse(self.lookup('new-003')['exists'])
        self.assertEqual(self.save(number='new-003').status_code, 409)
        self.assertEqual(self.save(number='new-003', fields={'SLA Expiration': 'garbage'}).status_code, 409)
        response = self.save(number='new-003', status='Manual review', fields={'Borrower': 'New', 'SLA Expiration': '2026-10-19T17:00'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue(response.json['created'])
        report = self.client.get('/api/sla-status').json
        self.assertEqual((report['captured_order_count'], report['manual_order_count'], report['totals']['Grand Total']), (2, 1, 3))
        self.assertEqual(sum(report['totals'][category] for category in ('Future SLA', 'Negative SLA', 'Today SLA')), 3)
        self.assertEqual(self.save(number='NEW-003', fields={'SLA Expiration': '2026-10-20T17:00'}).status_code, 409)
        lookup = self.lookup('NEW-003')
        self.assertEqual(lookup['row']['Queue Presence'], 'Manual addition')
        self.assertEqual(self.save(number='NEW-003', lookup=lookup).status_code, 200)
        self.assertEqual(self.client.get('/api/sla-status').json['totals']['Grand Total'], 3)
        self.assertEqual(len(self.store.get(1)['rows']), 2)
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
        self.assertEqual(book['Sheet1'].max_row - 1, 3)
        self.assertEqual(book['Sheet2']._pivots[0].cache.recordCount, 3)
        self.assertEqual(self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')[-1][-1], 3)

    def test_invalid_preview_and_import_mode_cannot_change_manual_orders(self):
        self.initial()
        self.assertEqual(self.save(preview=999).status_code, 404)
        self.assertEqual(self.save(preview=True).status_code, 409)
        self.workspace.save(mode='excel')
        self.assertEqual(self.save().status_code, 409)
        self.assertEqual(self.lookup().get('error'), 'Choose Tracker report before editing preview orders.')

    def test_empty_custom_status_can_be_deleted_and_does_not_return_after_refresh(self):
        before = self.initial()
        history = deepcopy(self.store.history())
        created = self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'In House Pending'})
        self.assertEqual(created.status_code, 200, created.json)
        response = self.client.post('/api/sla-statuses', json={'action': 'delete', 'preview_id': 1,
            'status': 'In House Pending', 'order_numbers': []})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['moved'], 0)
        report = self.client.get('/api/sla-status?preview_id=1').json
        self.assertNotIn('In House Pending', report['custom_statuses'])
        self.assertNotIn('In House Pending', [row['Status'] for row in report['summary']])
        self.assertEqual(report['totals'], before['totals'])
        self.assertEqual(self.store.history(), history)
        self.assertEqual(self.client.post('/api/sla-statuses', json={'action': 'delete', 'preview_id': 1,
            'status': 'Typing in Progress', 'order_numbers': ['001'], 'replacement': 'Search In Progress'}).status_code, 409)

    def test_deleting_populated_status_moves_every_order_once_and_keeps_manual_details(self):
        self.initial()
        history = deepcopy(self.store.history())
        self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'In House Pending'})
        moved = self.save(status='In House Pending', lookup=self.lookup())
        self.assertEqual(moved.status_code, 200, moved.json)
        added = self.save(number='003', status='In House Pending',
            fields={'Borrower': 'Manual borrower', 'SLA Expiration': '2026-10-20T17:00'})
        self.assertEqual(added.status_code, 200, added.json)
        before = self.client.get('/api/sla-status?preview_id=1').json
        body = {'action': 'delete', 'preview_id': 1, 'status': 'In House Pending', 'order_numbers': ['001', '003']}
        for changes in ({}, {'replacement': 'In House Pending'}, {'replacement': 'Unknown target'},
                        {'replacement': 'QC in Progress', 'order_numbers': ['001']}):
            rejected = self.client.post('/api/sla-statuses', json=dict(body, **changes))
            self.assertEqual(rejected.status_code, 409, rejected.json)
        response = self.client.post('/api/sla-statuses', json=dict(body, replacement='QC in Progress'))
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['moved'], 2)
        after = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(after['totals'], before['totals'])
        self.assertEqual(after['manual_order_count'], 1)
        self.assertNotIn('In House Pending', [row['Status'] for row in after['summary']])
        self.assertEqual({row['Order Number']: row['Status'] for row in after['rows']},
            {'001': 'QC in Progress', '002': 'Search In Progress', '003': 'QC in Progress'})
        self.assertEqual(next(row['Borrower'] for row in after['rows'] if row['Order Number'] == '003'), 'Manual borrower')
        self.assertEqual(len(after['rows']), len({row['Order Number'] for row in after['rows']}))
        self.assertEqual(self.store.history(), history)
        self.assertEqual(self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')[-1][-1], 3)

    def test_legacy_custom_names_are_preserved_and_deletion_is_isolated_to_selected_preview(self):
        self.initial()
        self.capture(2, self.store.get(1)['rows'])
        self.workspace.save(sla_custom_statuses=['In House Pending'],
            sla_preview_statuses={'1': {'001': 'In House Pending'}, '2': {'001': 'In House Pending'}})
        response = self.client.post('/api/sla-statuses', json={'action': 'delete', 'preview_id': 1,
            'status': 'In House Pending', 'order_numbers': ['001'], 'replacement': 'Search In Progress'})
        self.assertEqual(response.status_code, 200, response.json)
        first = self.client.get('/api/sla-status?preview_id=1').json
        second = self.client.get('/api/sla-status?preview_id=2').json
        self.assertNotIn('In House Pending', first['custom_statuses'])
        self.assertEqual(second['custom_statuses'], ['In House Pending'])
        self.assertEqual(next(row['Status'] for row in second['rows'] if row['Order Number'] == '001'), 'In House Pending')
        recreated = self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'In House Pending'})
        self.assertEqual(recreated.status_code, 200, recreated.json)
        self.assertIn('In House Pending', self.client.get('/api/sla-status?preview_id=1').json['custom_statuses'])
        self.assertNotIn('sla_preview_custom_statuses', self.workspace.public())

    def test_populated_custom_status_rename_updates_only_its_preview_registry(self):
        self.initial()
        self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'In House Pending'})
        moved = self.save(status='In House Pending', lookup=self.lookup())
        self.assertEqual(moved.status_code, 200, moved.json)
        self.capture(2, self.store.get(1)['rows'])
        renamed = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'In House Pending',
            'replacement': 'In House Review', 'order_numbers': ['001'], 'edit_scope': 'group'})
        self.assertEqual(renamed.status_code, 200, renamed.json)
        self.assertEqual(custom_statuses(self.workspace.read(), 1), ['In House Review'])
        self.assertEqual(custom_statuses(self.workspace.read(), 2), [])
        self.assertEqual(self.lookup()['row']['Status'], 'In House Review')

    def test_delete_includes_orders_with_invalid_deadlines(self):
        self.capture(1, [{'Order Number': 'bad', 'Status': 'In House Pending', 'SLA Expiration*': 'invalid'}])
        self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': 'In House Pending'})
        response = self.client.post('/api/sla-statuses', json={'action': 'delete', 'preview_id': 1,
            'status': 'In House Pending', 'order_numbers': ['bad'], 'replacement': 'Search In Progress'})
        self.assertEqual(response.status_code, 200, response.json)
        after = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(after['excluded_rows'][0]['Status'], 'Search In Progress')
        self.assertEqual(after['excluded_count'], 1)
        self.assertEqual(after['totals']['Grand Total'], 0)


if __name__ == '__main__':
    unittest.main()

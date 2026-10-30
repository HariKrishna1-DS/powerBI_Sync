"""Current deadline categories shared by the UI, Sheets and Excel reports."""
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import re
from portal_time import PORTAL_TIMEZONE

IST = timezone(timedelta(hours=5, minutes=30))
CATEGORIES = ['Future SLA', 'Negative SLA', 'Today SLA']
SUMMARY_COLUMNS = ['Status', *CATEGORIES, 'Grand Total']
DETAIL_COLUMNS = ['Order Number', 'SLA Expiration', 'SLA', 'Online/ Ground',
                  'Client', 'Product', 'St', 'County', 'Municipality',
                  'Task Name', 'Task Status', 'Status', 'Source workbook',
                  'ETA Comments', 'Last User', 'Status Reason', 'Source Preview', 'Queue Presence',
                  'Arrival Date', 'Arrival Time']
COLORS = ['#81cbb5', '#ed8796', '#e5bf68']


def tracker_status_rows(book):
    """Read only the tracker order-number/status columns for queue labels."""
    from gspread.utils import rowcol_to_a1
    from tracker_sync import canonical_headers
    headers = book.values_get("'All Products'!1:1", params={'valueRenderOption': 'UNFORMATTED_VALUE'}).get('values', [[]])[0]
    headers = canonical_headers(headers)
    ranges = []
    for name in ('Order Number', 'Status'):
        column = rowcol_to_a1(1, headers.index(name)+1)[:-1]
        ranges.append(f"'All Products'!{column}2:{column}")
    values = book.values_batch_get(ranges, params={'valueRenderOption': 'UNFORMATTED_VALUE'})['valueRanges']
    numbers, statuses = (item.get('values', []) for item in values)
    return [{'Order Number': cells[0], 'Status': statuses[index][0] if index < len(statuses) and statuses[index] else ''}
            for index, cells in enumerate(numbers) if cells and str(cells[0]).strip()]


def queue_field(row, name):
    """Portal headers may contain spaces, line breaks, or no spaces."""
    compact = lambda value: re.sub(r'[^a-z0-9]', '', str(value).casefold())
    return next((str(value or '').strip() for column, value in row.items()
                 if compact(column) == compact(name)), '')


def special_product_status(raw):
    """Recognize full names and the portal's displayed, ellipsized product labels."""
    product = re.sub(r'[^a-z0-9]', '', raw.casefold())
    truncated = bool(re.search(r'(?:\.{2,}|\u2026)\s*$', raw))
    for status, names, prefix in (
            ('Attorney Approval', ('attorneyapproval', 'attorneyapproach'), 'attorneyapp'),
            ('Home Builder', ('homebuilder', 'homebuilders'), 'homebuil')):
        if product in names or (truncated and product.startswith(prefix)
                                and any(name.startswith(product) for name in names)):
            return status
    return ''


def queue_rule(row, as_of=None):
    """Return priority, status and evidence shared by extraction and SLA reports."""
    from monthly_production import local_datetime
    from production_timing import completed, precise_timestamp
    if local_datetime(queue_field(row, 'Out Time')):
        return -1, 'Completed and Delivered', 'Valid Out Time'
    normalized = {name: queue_field(row, name) for name in ('Status', 'Task Status', 'Completed Time')}
    normalized['Status'] = normalized['Status'] or normalized['Task Status']
    if completed(normalized) and precise_timestamp(normalized['Completed Time']):
        return -1, 'Completed and Delivered', 'Completed status with a valid Completed Time'
    compact = lambda value: re.sub(r'[^a-z0-9]', '', value.casefold())
    product, task = compact(queue_field(row, 'Product')), compact(queue_field(row, 'Task Name'))
    comments = queue_field(row, 'ETA Comments')
    task_status = compact(queue_field(row, 'Task Status'))
    workflow_suspended = task_status == 'workflowsuspended' or task == 'workflowsuspended'
    portal_comments = ' '.join(queue_field(row, name) for name in ('Comment', 'Comments'))
    if workflow_suspended:
        if task_status == 'workflowsuspended' and re.search(r'\bon\s+hold\b', comments, re.I):
            return -0.5, 'Awaiting for Clarification', 'Workflow suspended; ETA Comments contains ON HOLD'
        if re.search(r'\bcannot\b', portal_comments, re.I):
            return -0.5, 'Awaiting for Clarification', 'Workflow suspended; Comment contains Cannot'
        if task in ('searchfix', 'crsp2'):
            return -0.5, 'Awaiting for Clarification', 'Workflow suspended SearchFix task'
        if re.search(r'\babst(?:raction|ractor|ractors|ract|racts)?\b', comments, re.I):
            return -0.5, 'Awaiting for Clarification', 'Workflow suspended; ETA Comments references abstraction'
    if task in ('crsp2', 'searchfix'):
        return 0, 'SearchFix', 'Task Name is CRSP2' if task == 'crsp2' else 'Task Name is SearchFix'
    if product in ('clientreview', 'qualitycontrol', 'quailtycontrol'):
        return 0, 'SearchFix', 'Product is ClientReview / Quality Control'
    if re.search(r'\bon\s+hold\b', comments, re.I):
        return 1, 'Waiting for Effective Date', 'ETA Comments contains ON HOLD'
    special = special_product_status(queue_field(row, 'Product'))
    if special:
        return 2, special, f"Product is {queue_field(row, 'Product')}"
    apology = re.search(r'\b(?:apoloize|apologize|apologise)\b', comments, re.I)
    if apology:
        return 3, 'Assign to ABS', f'ETA Comments contains {apology.group(0)}'
    assigned = re.search(r'\bassigned\s+to\s+the\s+(?:abstractor|abst)\b', comments, re.I)
    abstractor = re.search(r'\babstract(?:ors?)?\b', comments, re.I)
    commitment = re.search(r'\b(?:eod|sooner)\b', comments, re.I)
    if assigned or abstractor or commitment:
        evidence = assigned or abstractor or commitment
        reason = f'ETA Comments contains {" ".join(evidence.group(0).split())}'
        return 4, 'Assign to ABS', reason
    if task == 'typingmodule' and task_status == 'inprogress':
        return 5, 'Typing in Progress', 'Task Name is TypingModule and Task Status is In Progress'
    from tracker_sync import suspended
    if suspended(dict(row, **{'Task Status': queue_field(row, 'Task Status')})):
        return 6, 'Awaiting for Clarification', 'Workflow is suspended'
    last_user = queue_field(row, 'Last User').casefold()
    if (last_user in ('', '(blank)', 'null', 'none', 'n/a', '-')
            and task_status == 'available'):
        return 7, 'Search In Progress', 'Last User is blank and Task Status is Available'
    if task_status == 'tasksuspended':
        return 8, 'Typing in Progress', 'Task Status is Task Suspended; no higher-priority queue rule matches'
    return 9, '', ''


def completion_status(value):
    return re.sub(r'[^a-z]', '', str(value or '').casefold()) in (
        'complete', 'completed', 'completeanddelivered', 'completedanddelivered')


def sla_queue_rule(row, as_of=None):
    """The active SLA queue never inherits production completion evidence."""
    compact = lambda value: re.sub(r'[^a-z0-9]', '', str(value).casefold())
    active = {column: value for column, value in row.items()
              if compact(column) not in ('outtime', 'completedtime')}
    for column in active:
        if compact(column) in ('status', 'taskstatus') and completion_status(active[column]):
            active[column] = ''
    return queue_rule(active, as_of)


def queue_status_reason(row):
    task, state = queue_field(row, 'Task Name'), queue_field(row, 'Task Status')
    evidence = '; '.join(f'{column} is {value}' for column, value in
                         (('Task Name', task), ('Task Status', state)) if value and not completion_status(value))
    return evidence or 'Default active search queue status'


def queue_orders(preview, as_of=None, *, rule=None):
    """Count each Order Number once even when the queue contains several tasks."""
    from import_identity import number_key
    rule = rule or queue_rule
    orders = {}
    for original in preview['rows'] if preview else []:
        identity = number_key(original)
        if identity and (identity not in orders or rule(original, as_of)[0] < rule(orders[identity], as_of)[0]):
            orders[identity] = original
    return orders


def preview_snapshot(store, settings=None, preview_id=None, production=None, as_of=None):
    """Only orders present in the selected queue, with SLA-specific status rules."""
    settings = settings or {}
    as_of = as_of or datetime.now(IST)
    previews = store.list()
    if preview_id is not None:
        preview = store.get(preview_id)
    else:
        saved = settings.get('sla_preview_id')
        selected = next((p for p in previews if p['id'] == saved), None) or next(iter(previews), None)
        preview = store.get(selected['id']) if selected else None
    from reporting_dates import latest_arrival
    from status_updates import status_name
    reference = latest_arrival(preview['rows']) if preview else None
    overrides = (settings.get('sla_preview_statuses') or {}).get(str(preview['id']), {}) if preview else {}
    from tracker_sync import SOURCE, value, text, deadline
    from import_identity import number_key
    from sla_manual import FIELDS, apply_fields, revision, custom_statuses
    current_orders = queue_orders(preview, as_of, rule=sla_queue_rule)
    captured_order_count = len(current_orders)
    raw_rows = preview['rows'] if preview else []
    occurrences = Counter(number_key(row) for row in raw_rows)
    repeated = {identity for identity, count in occurrences.items() if identity and count > 1}
    duplicate_rows = [dict(row, **{'Capture row': index, 'SLA selection':
                      'Counted once' if row is current_orders[number_key(row)] else 'Repeated row'})
                      for index, row in enumerate(raw_rows, 1) if number_key(row) in repeated]
    repeated_sources, repeated_details = {}, {}
    for row in raw_rows:
        if number_key(row) in repeated:
            repeated_sources.setdefault(number_key(row), []).append(row)
    for row in duplicate_rows:
        repeated_details.setdefault(number_key(row), []).append(row)
    capture_positions = {id(row): index for index, row in enumerate(raw_rows, 1)}
    duplicate_groups = []
    for identity in occurrences:
        if identity not in repeated:
            continue
        selected = current_orders[identity]
        priority, _, condition = sla_queue_rule(selected, as_of)
        same_priority = sum(sla_queue_rule(row, as_of)[0] == priority for row in repeated_sources[identity])
        selection = f'Highest-priority matching condition: {condition}.' if condition else 'No automatic queue condition distinguishes these entries.'
        if same_priority > 1:
            selection += ' When conditions have equal priority, the first matching captured row is retained.'
        position = capture_positions[id(selected)]
        duplicate_groups.append({'order_number': str(selected['Order Number']).strip(),
                                 'row_count': occurrences[identity], 'selected_capture_row': position,
                                 'selection_reason': selection})
        for candidate in repeated_details[identity]:
            _, entry_status, entry_condition = sla_queue_rule(candidate, as_of)
            candidate.update({'Entry status': entry_status or 'Search In Progress',
                              'Entry condition': entry_condition or queue_status_reason(candidate),
                              'Counted capture row': position, 'Selection reason': selection})
    preview_key = str(preview['id']) if preview else ''
    manual_orders = (settings.get('sla_preview_manual_orders') or {}).get(preview_key, {})
    detail_edits = (settings.get('sla_preview_order_edits') or {}).get(preview_key, {})
    added = set(manual_orders) - set(current_orders)
    current_orders.update({identity: row for identity, row in manual_orders.items() if identity in added})
    statuses = {}
    for row in (production or {}).get('sheets', {}).get('Overview', {}).get('rows', []):
        identity = text(row.get('Order Number')).casefold()
        status = text(row.get('Status'))
        if identity and status:
            statuses.setdefault(identity, set()).add(status)
    rows = []
    for identity, original in current_orders.items():
        captured = preview
        anchor = local_time(captured['created'])
        row = dict(original)
        apply_fields(row, detail_edits.get(identity, {}))
        original = dict(row)
        for column in ('Product', 'Task Name', 'Task Status', 'Last User', 'ETA Comments'):
            if queue_field(original, column):
                row[column] = queue_field(original, column)
        for column in SOURCE:
            if not text(row.get(column)):
                row[column] = value(original, column)
        _, automatic, reason = sla_queue_rule(original, as_of)
        known = {status for status in statuses.get(identity, set()) if not completion_status(status)}
        saved_status = text(row.get('Status'))
        if automatic:
            row['Status'] = automatic
        elif saved_status and not completion_status(saved_status):
            reason = 'Saved preview Status'
        elif len(known) == 1:
            row['Status'], reason = next(iter(known)), 'Matched All Products Status by Order Number'
        else:
            row['Status'], reason = 'Search In Progress', queue_status_reason(original)
        expiry = deadline(row.get('SLA Expiration'), anchor)
        if expiry:
            row['SLA Expiration'] = expiry.isoformat(sep=' ')
        if identity in overrides and not completion_status(overrides[identity]):
            row['Status'], reason = overrides[identity], 'Manual status edit for this preview'
        elif (identity in (settings.get('sheet_status_overrides') or {})
              and not completion_status(settings['sheet_status_overrides'][identity])
              and str((settings.get('sheet_status_override_previews') or {}).get(identity, preview['id'])) == str(preview['id'])):
            row['Status'], reason = settings['sheet_status_overrides'][identity], 'Status edited in Google Sheets'
        row['Status'] = status_name(row['Status'])
        row.update({'Status Reason': reason, 'Source Preview': captured['name'],
                    'Queue Presence': 'Manual addition' if identity in added else 'In selected preview',
                    '_record_key': number_key(row)})
        row['_sla_revision'] = revision(row)
        rows.append(row)
    columns = list(dict.fromkeys((preview['columns'] if preview else []) + list(SOURCE) + FIELDS +
        ['Status', 'Status Reason', 'Source Preview', 'Queue Presence']))
    return {'mode': 'tracker', 'report_scope': 'preview', 'offline': False,
            'source': 'Saved preview', 'preview_id': preview['id'] if preview else None,
            'preview_name': preview['name'] if preview else None,
            'preview_row_count': len(preview['rows']) if preview else 0,
            'captured_order_count': captured_order_count, 'manual_order_count': len(added),
            'duplicate_order_count': len(repeated), 'duplicate_row_count': sum(occurrences[key]-1 for key in repeated),
            'missing_order_number_count': occurrences.get('', 0), 'duplicate_rows': duplicate_rows,
            'duplicate_groups': duplicate_groups,
            'custom_statuses': [status for status in custom_statuses(settings, preview['id']) if not completion_status(status)] if preview else [], 'editable_columns': FIELDS,
            'uses_queue_status': any((identity not in overrides or completion_status(overrides[identity]))
                                     and (not text(row.get('Status')) or completion_status(row.get('Status')))
                                     and not sla_queue_rule(row, as_of)[1]
                                     for identity, row in current_orders.items()),
            'previews': previews, 'updated_at': preview['created'] if preview else None,
            'sla_reference_at': reference.isoformat() if reference else None,
            'sheets': {'Overview': {'rows': rows, 'columns': columns}}}


def rename_preview_status(workspace, snapshot, old, new, selected, individual=False):
    from copy import deepcopy
    from tracker_sync import key
    from status_updates import status_name
    from sla_manual import checked_preview_status
    new = checked_preview_status(new)
    matching = {key(row) for row in snapshot['sheets']['Overview']['rows']
                if (status_name(row.get('Status') or row.get('Task Status')) or '(Blank)') == old}
    chosen = set(selected)
    if not chosen or not chosen.issubset(matching) or (not individual and matching != chosen):
        raise ValueError('Orders in this preview changed. Refresh SLA Status and try again.')
    from workspace_backup import save_safety_backup
    save_safety_backup(workspace.store, 'before-update')
    with workspace.lock:
        edits = deepcopy(workspace.read().get('sla_preview_statuses') or {})
        edits.setdefault(str(snapshot['preview_id']), {}).update(dict.fromkeys(chosen, new))
        changes = {'sla_preview_statuses': edits}
        from sla_manual import custom_statuses, custom_status_changes
        settings = workspace.read()
        custom = custom_statuses(settings, snapshot['preview_id'])
        if not individual and old in custom:
            changes.update(custom_status_changes(settings, snapshot['preview_id'],
                [new if status == old else status for status in custom]))
        workspace.save(**changes)
    return len(chosen)


def local_time(value):
    from portal_time import capture_time
    return capture_time(value)


def category(value, as_of, anchor=None):
    from monthly_production import local_datetime
    raw = str(value if value is not None else '').strip()
    if not raw:
        return ''
    # Queue exports express overdue deadlines as signed elapsed durations.
    relative = re.fullmatch(r'(-)?(?:(\d+)d\s*)?(?:(\d+)h\s*)?(?:(\d+)m\s*)?', raw, re.I)
    if relative and any(relative.groups()[1:]):
        if relative[1]:
            return 'Negative SLA'
        days, hours, minutes = (int(part or 0) for part in relative.groups()[1:])
        due = (anchor or as_of) + timedelta(days=days, hours=hours, minutes=minutes)
    else:
        due = local_datetime(value)
    if due is None:
        return ''
    return 'Today SLA' if due.date() == as_of.date() else 'Negative SLA' if due.date() < as_of.date() else 'Future SLA'


def imported_date_snapshot(snapshot, selected=None, strict=True):
    """Select one imported arrival date, keeping workbook order identities intact."""
    from reporting_dates import arrival_details
    source = snapshot.get('sheets', {}).get('Overview') or snapshot.get('sheets', {}).get('All Products') or {'rows': [], 'columns': []}
    rows = [dict(row, **arrival_details(row)) for row in source['rows'] if str(row.get('Order Number') or '').strip()]
    counts = Counter(row['Arrival Date'] or 'Undated' for row in rows)
    dates = sorted((day for day in counts if day != 'Undated'), reverse=True)
    dates += ['Undated'] if counts.get('Undated') else []
    if selected and selected not in dates and strict:
        raise ValueError('That date is no longer in the imported report. Choose an available report date.')
    chosen = selected if selected in dates else next(iter(dates), '')
    selected_rows = [row for row in rows if (row['Arrival Date'] or 'Undated') == chosen]
    return dict(snapshot, report_scope='imported_date', report_date=chosen,
        report_dates=[{'date': day, 'orders': counts[day]} for day in dates],
        imported_order_count=len(rows), date_order_count=len(selected_rows),
        undated_order_count=counts.get('Undated', 0),
        sheets={**snapshot.get('sheets', {}), 'Overview': {**source, 'rows': selected_rows}})


def build_sla_status(snapshot, as_of=None):
    from monthly_production import local_datetime
    from status_updates import status_name
    from reporting_dates import arrival_details
    selected_day = snapshot.get('report_date')
    reference = local_datetime(snapshot.get('sla_reference_at'))
    as_of = as_of or (datetime.fromisoformat(selected_day) if selected_day and selected_day != 'Undated'
                      else reference or local_time(snapshot.get('updated_at')) or datetime.now(PORTAL_TIMEZONE))
    if isinstance(as_of, date) and not isinstance(as_of, datetime):
        as_of = datetime.combine(as_of, datetime.min.time())
    if as_of.tzinfo:
        as_of = as_of.astimezone(PORTAL_TIMEZONE).replace(tzinfo=None)
    sheets = snapshot.get('sheets', {})
    source = sheets.get('Overview') or sheets.get('All Products') or {'rows': [], 'columns': []}
    anchor = as_of if snapshot.get('report_scope') == 'imported_date' else local_time(snapshot.get('updated_at'))
    rows, excluded, counts = [], [], Counter()
    for original in source['rows']:
        if not str(original.get('Order Number') or '').strip():
            continue
        expiry = original.get('SLA Expiration', original.get('SLA Expiration*'))
        result = category(expiry, as_of, anchor)
        row = dict(original, **arrival_details(original), **{'SLA Category': result})
        row['Status'] = status_name(original.get('Status') or original.get('Task Status') or '(Blank)')
        if not result and snapshot.get('report_scope') == 'preview':
            result = row['SLA Category'] = 'Unknown SLA'
            row['SLA issue'] = 'Missing SLA expiration' if not str(expiry or '').strip() else 'Invalid SLA expiration'
        if not result:
            row['SLA issue'] = 'Missing SLA expiration' if not str(original.get('SLA Expiration') or '').strip() else 'Invalid SLA expiration'
            excluded.append(row)
            continue
        counts[row['Status'], result] += 1
        rows.append(row)
    custom = list(dict.fromkeys(status_name(status) for status in snapshot.get('custom_statuses') or []))
    statuses = sorted({status for status, _ in counts} | set(custom), key=str.casefold)
    categories = CATEGORIES + (['Unknown SLA'] if any(row['SLA Category'] == 'Unknown SLA' for row in rows) else [])
    summary_columns = ['Status', *categories, 'Grand Total']
    summary = [dict(Status=status, **{key: counts[status, key] for key in categories},
                    **{'Grand Total': sum(counts[status, key] for key in categories)}) for status in statuses]
    totals = dict(Status='Grand Total', **{key: sum(row[key] for row in summary) for key in categories},
                  **{'Grand Total': len(rows)})
    from import_identity import number_key
    counted = {number_key(row): row for row in rows + excluded}
    duplicate_groups = []
    for group in snapshot.get('duplicate_groups') or []:
        order = counted[number_key({'Order Number': group['order_number']})]
        duplicate_groups.append(dict(group, status=order['Status'], status_reason=order.get('Status Reason', ''),
                                     sla_category=order['SLA Category']))
    duplicate_rows = [dict(row, **{'Counted status': counted[number_key(row)]['Status'],
                                  'Counted SLA': counted[number_key(row)]['SLA Category'],
                                  'Applied status condition': counted[number_key(row)].get('Status Reason', '')})
                      for row in snapshot.get('duplicate_rows') or []]
    return {'as_of': as_of.date().isoformat(), 'date_basis': 'arrival_time' if reference or selected_day and selected_day != 'Undated' else 'capture_time',
            'sla_reference_at': reference.isoformat() if reference else None, 'rows': rows, 'excluded_rows': excluded,
            'excluded_count': len(excluded), 'summary': summary, 'totals': totals,
            'columns': list(dict.fromkeys(source.get('columns', []) + ['Arrival Date', 'Arrival Time', 'SLA Category'])),
            'summary_columns': summary_columns, **{key: snapshot.get(key) for key in (
                'source', 'mode', 'offline', 'updated_at', 'report_scope', 'preview_id',
                'preview_name', 'preview_row_count', 'previews', 'captured_order_count',
                'duplicate_order_count', 'duplicate_row_count', 'missing_order_number_count', 'duplicate_rows',
                'manual_order_count', 'custom_statuses', 'editable_columns', 'report_date',
                'report_dates', 'imported_order_count', 'date_order_count', 'undated_order_count', 'source_workbooks', 'data_revision')},
            'custom_statuses': custom, 'duplicate_groups': duplicate_groups, 'duplicate_rows': duplicate_rows}


def summary_values(report):
    columns = report['summary_columns']
    return [columns] + [[row[column] for column in columns]
        for row in report['summary'] + [report['totals']]]


def sheet_summary_values(report):
    """Reference layout, with empty category counts displayed as blank cells."""
    return [['Count of Order Number', 'Column Labels'] + [''] * (len(report['summary_columns'])-2)] + [
        [value if index in (0, len(row)-1) or value else '' for index, value in enumerate(row)]
        for row in summary_values(report)]


def sheet_format_requests(sheet_id, values, existing):
    """Keep the SLA table plain, including when replacing an older colored report."""
    table = {'sheetId': sheet_id, 'endRowIndex': len(values), 'endColumnIndex': len(values[1])}
    requests = [{'deleteConditionalFormatRule': {'sheetId': sheet_id, 'index': index}}
                for index in reversed(range(len(existing.get('conditionalFormats', []))))]
    requests.append({'repeatCell': {'range': {'sheetId': sheet_id},
        'cell': {}, 'fields': 'userEnteredFormat'}})
    requests.append({'repeatCell': {'range': table,
        'cell': {'userEnteredFormat': {'horizontalAlignment': 'CENTER',
            'verticalAlignment': 'MIDDLE', 'textFormat': {'fontFamily': 'Calibri', 'fontSize': 11}}},
        'fields': 'userEnteredFormat'}})
    for start, end in ((0, 2), (len(values)-1, len(values))):
        requests.append({'repeatCell': {'range': dict(table, startRowIndex=start, endRowIndex=end),
            'cell': {'userEnteredFormat': {'textFormat': {'bold': True}}},
            'fields': 'userEnteredFormat.textFormat.bold'}})
    requests.append({'repeatCell': {'range': dict(table, startRowIndex=2, startColumnIndex=1),
        'cell': {'userEnteredFormat': {'numberFormat': {'type': 'NUMBER', 'pattern': '#,##0'}}},
        'fields': 'userEnteredFormat.numberFormat'}})
    border = {'style': 'SOLID', 'color': {'red': 0, 'green': 0, 'blue': 0}}
    requests.append({'updateBorders': {'range': dict(table, startRowIndex=1),
        **{side: border for side in ('top', 'bottom', 'left', 'right', 'innerHorizontal', 'innerVertical')}}})
    for start, end, width in ((0, 1, 220), (1, 4, 115), (4, 5, 120)):
        requests.append({'updateDimensionProperties': {
            'range': {'sheetId': sheet_id, 'dimension': 'COLUMNS', 'startIndex': start, 'endIndex': end},
            'properties': {'pixelSize': width}, 'fields': 'pixelSize'}})
    return requests


def sample_rows(report):
    aliases = {'SLA': 'SLA Category', 'Online/ Ground': 'Online/Ground', 'St': 'State'}
    return [{column: row.get(column, row.get(aliases.get(column, column), ''))
             if column != 'SLA' else row['SLA Category'] for column in DETAIL_COLUMNS} for row in report['rows']]


def add_excel_pivot(summary, detail, report):
    """Serialize the reference's native Status/SLA/count pivot and its cache.

    Materialized cells and cache records both use the shared report, so the
    downloaded workbook opens with its totals and can refresh from Sheet1.
    """
    from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
    from openpyxl.pivot.fields import Text, Index, DateTimeField
    from openpyxl.pivot.record import Record, RecordList
    from openpyxl.pivot.table import (TableDefinition, Location, PivotField, FieldItem,
                                     RowColField, RowColItem, DataField, PivotTableStyle)
    statuses = [row['Status'] for row in report['summary']]
    dimensions = {2: report['summary_columns'][1:-1], 11: statuses}
    lookups = {column: {value: index for index, value in enumerate(values)} for column, values in dimensions.items()}
    fields = [CacheField(name=name, sharedItems=SharedItems(
        _fields=[Text(v=value) for value in dimensions.get(column, [])]))
        for column, name in enumerate(DETAIL_COLUMNS)]
    cache = CacheDefinition(cacheSource=CacheSource(type='worksheet',
        worksheetSource=WorksheetSource(ref=detail.dimensions, sheet=detail.title)),
        cacheFields=fields, recordCount=len(report['rows']), saveData=True,
        refreshOnLoad=False, enableRefresh=True, createdVersion=6, refreshedVersion=6,
        minRefreshableVersion=3)
    cache.records = RecordList(r=[Record(_fields=[
        Index(v=lookups[column][str(value)]) if column in dimensions else
        DateTimeField(v=value) if isinstance(value, datetime) else Text(v=str(value if value is not None else ''))
        for column, value in enumerate(values)]) for values in detail.iter_rows(min_row=2, values_only=True)])
    pivot_fields = [PivotField(showAll=False) for _ in DETAIL_COLUMNS]
    pivot_fields[0].dataField = True
    for column, axis in ((2, 'axisCol'), (11, 'axisRow')):
        pivot_fields[column] = PivotField(axis=axis, showAll=True, compact=False,
            defaultSubtotal=False, items=[FieldItem(x=index) for index in range(len(dimensions[column]))] + [FieldItem(t='default')])
    def items(length):
        return [RowColItem(x=[Index(v=index)]) for index in range(length)] + [RowColItem(t='grand', x=[Index(v=0)])]
    pivot = TableDefinition(name='SLAStatusPivot', cacheId=1, dataCaption='Values',
        rowHeaderCaption='Status', compact=False, compactData=False, outline=False, outlineData=False,
        rowGrandTotals=True, colGrandTotals=True,
        location=Location(ref=f'A3:{chr(64+len(report["summary_columns"]))}{len(statuses)+5}', firstHeaderRow=1, firstDataRow=2, firstDataCol=1),
        pivotFields=pivot_fields, rowFields=[RowColField(x=11)], rowItems=items(len(statuses)),
        colFields=[RowColField(x=2)], colItems=items(len(dimensions[2])),
        dataFields=[DataField(name='Count of Order Number', fld=0, subtotal='count')],
        pivotTableStyleInfo=PivotTableStyle(name='PivotStyleLight16', showRowHeaders=True,
            showColHeaders=True, showRowStripes=False, showColStripes=False, showLastColumn=True))
    pivot.cache = cache
    summary.add_pivot(pivot)


def excel_report(report):
    """A report snapshot in the supplied Sheet2/Sheet1 layout."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter
    book = Workbook()
    summary = book.active
    summary.title = 'Sheet2'
    summary.append([])
    summary.append([])
    summary['A1'] = (f"Arrival date: {report['report_date']} | SLA as of: {report['as_of']}"
                     if report.get('report_date') else f"SLA as of: {report['as_of']}")
    if report.get('excluded_count'):
        summary['A2'] = f"{report['excluded_count']} orders excluded: missing or invalid SLA expiration."
    summary.append(['Count of Order Number', 'Column Labels'])
    for values in summary_values(report):
        # The reference displays zero category cells as blanks.
        summary.append([value if index == 0 or value or index == len(values)-1 else None
                        for index, value in enumerate(values)])
    detail = book.create_sheet('Sheet1')
    detail.append(DETAIL_COLUMNS)
    for row in sample_rows(report):
        values = [row[column] for column in DETAIL_COLUMNS]
        values[1] = local_time(values[1]) or values[1]
        detail.append(values)
    border = Border(**{edge: Side(style='thin', color='D4D4D4') for edge in ('left', 'right', 'top', 'bottom')})
    for sheet in (summary, detail):
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.font = Font(name='Calibri', size=11)
                if cell.row >= (4 if sheet == summary else 1):
                    cell.border = border
                cell.alignment = Alignment(vertical='center')
        header = 4 if sheet == summary else 1
        for cell in sheet[header]:
            cell.font = Font(name='Calibri', size=11, bold=True)
            cell.fill = PatternFill('solid', fgColor='F2F2F2')
        sheet.freeze_panes = f'A{header+1}'
        for index, width in enumerate(([30, 17, 17, 15, 17] if sheet == summary else [29, 24, 17, 18, 22, 27, 8, 22, 24, 22, 22, 30, 56, 48, 26, 64, 20, 32, 18, 28]), 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    for cells in detail.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(cells[0].value, datetime):
            cells[0].number_format = 'mm/dd/yyyy hh:mm'
    for cell in summary[summary.max_row]:
        cell.font = Font(name='Calibri', size=11, bold=True)
    detail.auto_filter.ref = detail.dimensions
    if report['summary']:
        add_excel_pivot(summary, detail, report)
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = 'bar', 'stacked', 100
        chart.title = 'SLA Status'
        chart.x_axis.title = 'Orders'
        chart.y_axis.title = 'Status'
        chart.add_data(Reference(summary, min_col=2, max_col=len(report['summary_columns'])-1, min_row=4, max_row=summary.max_row-1), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=5, max_row=summary.max_row-1))
        chart.width, chart.height = 27, max(12, len(report['summary']) * 0.65)
        for series, color in zip(chart.series, COLORS + ['#a8c5ed']):
            series.graphicalProperties.solidFill = color.lstrip('#')
        summary.add_chart(chart, 'G4')
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def sheet_chart_requests(sheet_id, values, existing, preserve_format=False):
    """Replace the report's chart in place, excluding Grand Total from series."""
    chart = next((item for item in existing.get('charts', []) if item.get('spec', {}).get('title') == 'SLA Status'), None)
    if preserve_format and chart:
        if len(values) <= 3:
            return []
        from copy import deepcopy
        spec = deepcopy(chart['spec'])
        basic = spec['basicChart']
        target_count = len(values[1])-2
        series = basic.setdefault('series', [])
        while len(series) < target_count:
            index = len(series)+1
            series.append({'series': {'sourceRange': {'sources': [{'sheetId': sheet_id,
                'startRowIndex': 1, 'endRowIndex': len(values)-1,
                'startColumnIndex': index, 'endColumnIndex': index+1}]}},
                'targetAxis': 'BOTTOM_AXIS'})
        del series[target_count:]
        for item in basic.get('domains', []):
            for source in item['domain']['sourceRange']['sources']:
                source.update(startRowIndex=1, endRowIndex=len(values)-1)
        for item in basic.get('series', []):
            for source in item['series']['sourceRange']['sources']:
                source.update(startRowIndex=1, endRowIndex=len(values)-1)
        return [{'updateChartSpec': {'chartId': chart['chartId'], 'spec': spec}}]
    if len(values) <= 3:
        return [{'deleteEmbeddedObject': {'objectId': chart['chartId']}}] if chart else []
    def source(column):
        return {'sourceRange': {'sources': [{'sheetId': sheet_id, 'startRowIndex': 1,
            'endRowIndex': len(values)-1, 'startColumnIndex': column, 'endColumnIndex': column+1}]}}
    from datatrace_sync import sheet_color
    spec = {'title': 'SLA Status', 'basicChart': {'chartType': 'BAR', 'stackedType': 'STACKED',
        'legendPosition': 'BOTTOM_LEGEND', 'headerCount': 1,
        'axis': [{'position': 'BOTTOM_AXIS', 'title': 'Orders'}, {'position': 'LEFT_AXIS', 'title': 'Status'}],
        'domains': [{'domain': source(0)}], 'series': [
            {'series': source(index), 'targetAxis': 'BOTTOM_AXIS', 'color': sheet_color(color)}
            for index, color in enumerate(('#555555', '#999999', '#cccccc', '#a8c5ed')[:len(values[1])-2], 1)]}}
    position = {'overlayPosition': {'anchorCell': {'sheetId': sheet_id, 'rowIndex': 1, 'columnIndex': 6},
                                   'widthPixels': 950, 'heightPixels': max(400, (len(values)-3)*32)}}
    if chart:
        return [{'updateChartSpec': {'chartId': chart['chartId'], 'spec': spec}},
                {'updateEmbeddedObjectPosition': {'objectId': chart['chartId'], 'newPosition': position,
                                                 'fields': 'anchorCell,widthPixels,heightPixels'}}]
    return [{'addChart': {'chart': {'spec': spec, 'position': position}}}]


def preserve_table_formats(sheet_id, previous, values):
    """Move the existing total style and extend its existing body style."""
    old_total = next((i for i, row in enumerate(previous) if row and row[0] == 'Grand Total'), None)
    new_total = len(values)-1
    if old_total is None or old_total == new_total:
        return []
    def copy(start, end, row):
        return {'copyPaste': {'source': {'sheetId': sheet_id, 'startRowIndex': row,
                    'endRowIndex': row+1, 'endColumnIndex': len(values[1])},
                'destination': {'sheetId': sheet_id, 'startRowIndex': start,
                    'endRowIndex': end, 'endColumnIndex': len(values[1])}, 'pasteType': 'PASTE_FORMAT'}}
    requests = [copy(new_total, new_total+1, old_total)]
    if old_total >= 3 and new_total > old_total:
        requests.append(copy(old_total, new_total, 2))
    return requests


def publish_preview_sla(book, report, backup_root):
    """Replace only SLA values and chart ranges; preserve existing Sheet styling."""
    from pathlib import Path
    from datetime import timezone
    import gzip
    import json
    import secrets
    from datatrace_sync import sheet_cell
    metadata = book.fetch_sheet_metadata(params={'fields': 'sheets(properties,conditionalFormats,charts)'})['sheets']
    existing = next((s for s in metadata if s['properties']['title'] == 'SLA Status'), None)
    previous = book.worksheet('SLA Status').get_all_values(value_render_option='FORMULA') if existing else []
    values = sheet_summary_values(report)
    number = existing['properties']['sheetId'] if existing else secrets.randbelow(2**30)
    requests = []
    if not existing:
        requests.append({'addSheet': {'properties': {'sheetId': number, 'title': 'SLA Status',
            'gridProperties': {'rowCount': len(values)+25, 'columnCount': 18, 'frozenRowCount': 2}}}})
    elif existing['properties'].get('gridProperties', {}).get('rowCount', 0) < len(values):
        requests.append({'updateSheetProperties': {'properties': {'sheetId': number,
            'gridProperties': {'rowCount': len(values)+25}}, 'fields': 'gridProperties.rowCount'}})
    requests.extend(preserve_table_formats(number, previous, values))
    requests.append({'updateCells': {'range': {'sheetId': number,
        'endRowIndex': max(len(values), len(previous)), 'endColumnIndex': max(len(values[1]), max((len(r) for r in previous), default=0))},
        'rows': [{'values': [sheet_cell(value) for value in row]} for row in values],
        'fields': 'userEnteredValue'}})
    if not existing:
        requests.extend(sheet_format_requests(number, values, {}))
    requests.extend(sheet_chart_requests(number, values, existing or {}, preserve_format=True))
    backup_root = Path(backup_root)
    backup_root.mkdir(parents=True, exist_ok=True)
    backup = backup_root / ('sla-preview-' + datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S-%f') + '.json.gz')
    with gzip.open(backup, 'wt', encoding='utf-8') as stream:
        json.dump({'metadata': existing, 'values': previous, 'preview_id': report['preview_id']}, stream)
    book.batch_update({'requests': requests})
    def normalized(matrix):
        result = []
        for row in matrix:
            row = [str(value) if value is not None else '' for value in row[:len(values[1])]]
            while row and not row[-1]:
                row.pop()
            result.append(row)
        while result and not result[-1]:
            result.pop()
        return result
    actual = book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')
    if normalized(actual) != normalized(values):
        raise RuntimeError('SLA Status could not be verified. Refresh SLA Status to retry.')
    return {'synced': True, 'tabs': ['SLA Status'], 'preview_id': report['preview_id']}

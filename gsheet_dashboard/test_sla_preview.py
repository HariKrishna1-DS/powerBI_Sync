"""Selected-preview SLA isolation across API, Excel, edits and Sheets publication."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import Mock, patch

from flask import Flask
from openpyxl import load_workbook
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, register_report_routes, imported_snapshot, read_report_files, report_frames
from sla_status import build_sla_status, preview_snapshot, publish_preview_sla, tracker_status_rows
from test_pdf_report_workspace import ReportBook, row, excel


class SlaPreviewTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.store = PreviewStore(Path(self.folder.name)/'previews')
        self.columns = ['Order Number', 'Status', 'Product', 'In-Time', 'SLA Expiration']
        self.first = [row('shared', 'Cancelled', **{'SLA Expiration': '2099-10-16'}),
                      row('old-only', 'Assign to ABS', **{'SLA Expiration': '-1d'})]
        self.latest = [row('shared', 'Completed and Delivered', **{'SLA Expiration': '-2d'}),
                       row('new-only', 'QC in Progress', **{'SLA Expiration': '2099-10-17'}),
                       row('missing', 'Search In Progress')]
        self.store.restore(1, self.columns, self.first, created='2026-10-14T12:00:00Z')
        self.store.restore(2, self.columns, self.latest, created='2026-10-15T12:00:00Z')
        self.workspace = ReportWorkspace(self.store)
        self.imported = read_report_files([('full.xlsx', excel([
            row('production-only', **{'SLA Expiration': '-90d'})]))])
        self.production = imported_snapshot(self.imported)
        self.workspace.save(mode='tracker', imported=self.imported)
        self.book = ReportBook()
        self.book.add('Other report', [{'Keep': 'preserve', 'unchanged': 99}])
        self.gate = threading.Lock()
        self.load = Mock(return_value=dict(self.production, mode='tracker', source='Google Sheets'))
        self.get_book = Mock(return_value=(self.book, None))
        self.app = Flask(__name__)
        self.publish = register_report_routes(self.app, self.workspace, self.gate,
                         self.load, self.get_book, lambda: None)
        self.client = self.app.test_client()

    def test_capture_rows_reconcile_to_unique_orders_with_repeat_details(self):
        rows = [
            {'Order Number': ' A ', 'Product': 'Document', 'Task Status': 'Available', 'SLA Expiration*': '2h'},
            {'Order Number': 'a', 'Product': 'Full Title', 'Task Status': 'Workflow Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': 'B', 'Task Name': 'Search', 'Task Status': 'Available', 'SLA Expiration*': '1d'},
            {'Order Number': 'b ', 'Task Name': 'SearchFix', 'Task Status': 'Available', 'SLA Expiration*': '1d'},
            {'Order Number': ' ', 'Task Status': 'Available', 'SLA Expiration*': '1d'},
        ]
        self.store.restore(3, list(dict.fromkeys(column for row in rows for column in row)), rows,
                           created='2026-10-16T04:00:00Z')
        before = deepcopy(self.store.get(3))
        report = self.client.get('/api/sla-status?preview_id=3').json
        self.assertEqual(report['preview_row_count'], 5)
        self.assertEqual(report['captured_order_count'], 2)
        self.assertEqual(report['duplicate_row_count'], 2)
        self.assertEqual(report['duplicate_order_count'], 2)
        self.assertEqual(report['missing_order_number_count'], 1)
        self.assertEqual(report['totals']['Grand Total'], 2)
        self.assertEqual(len(report['duplicate_rows']), 4)
        counted = [row for row in report['duplicate_rows'] if row['SLA selection'] == 'Counted once']
        self.assertEqual([row['Capture row'] for row in counted], [2, 4])
        groups = {group['order_number'].casefold(): group for group in report['duplicate_groups']}
        self.assertEqual(groups['a']['status'], 'Awaiting for Clarification')
        self.assertEqual(groups['a']['status_reason'], 'Workflow is suspended')
        self.assertEqual(groups['a']['selected_capture_row'], 2)
        self.assertIn('Highest-priority matching condition', groups['a']['selection_reason'])
        self.assertEqual(groups['b']['status_reason'], 'Task Name is SearchFix')
        self.assertTrue(all(row['Entry condition'] and row['Counted status'] for row in report['duplicate_rows']))
        self.assertEqual({row['Status'] for row in report['rows']}, {'Awaiting for Clarification', 'SearchFix'})
        self.assertEqual(self.store.get(3), before)

    def test_repeat_conditions_explain_equal_priority_and_manual_override(self):
        rows = [{'Order Number': 'A', 'Product': product, 'Task Status': 'Available', 'SLA Expiration*': '2h'}
                for product in ('Document', 'Full Title')]
        self.store.restore(3, list(rows[0]), rows, created='2026-10-16T04:00:00Z')
        self.workspace.save(sla_preview_statuses={'3': {'a': 'QC in Progress'}})
        report = self.client.get('/api/sla-status?preview_id=3').json
        group = report['duplicate_groups'][0]
        self.assertEqual(group['selected_capture_row'], 1)
        self.assertIn('equal priority', group['selection_reason'])
        self.assertEqual(group['status'], 'QC in Progress')
        self.assertEqual(group['status_reason'], 'Manual status edit for this preview')
        self.assertEqual(report['totals']['Grand Total'], 1)
        self.assertTrue(all(row['Counted status'] == 'QC in Progress' and row['Entry status'] == 'Search In Progress'
                            for row in report['duplicate_rows']))

    def test_selected_preview_does_not_add_orders_missing_since_its_predecessor(self):
        latest = self.client.get('/api/sla-status').json
        self.assertEqual(latest['preview_id'], 2)
        self.assertEqual(latest['preview_row_count'], 3)
        self.assertEqual(latest['totals']['Grand Total'], 2)
        self.assertEqual(latest['excluded_count'], 1)
        self.assertEqual({r['Order Number'] for r in latest['rows']}, {'shared', 'new-only'})
        self.assertEqual(next(r['Status'] for r in latest['rows'] if r['Order Number'] == 'shared'),
                         'Search In Progress')
        earlier = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual({r['Status'] for r in earlier['summary']}, {'Cancelled', 'Assign to ABS'})
        self.assertEqual([p['id'] for p in earlier['previews']], [2, 1])
        self.assertEqual(earlier['summary_columns'], ['Status', 'Future SLA', 'Negative SLA', 'Today SLA', 'Grand Total'])
        self.load.assert_not_called()
        self.get_book.assert_called_once_with()

    def test_download_contains_only_the_requested_preview_and_reconciles(self):
        response = self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1')
        self.assertEqual(response.status_code, 200)
        self.assertIn('preview1', response.headers['Content-Disposition'])
        book = load_workbook(BytesIO(response.data), data_only=True)
        self.assertEqual(book.sheetnames, ['Sheet2', 'Sheet1'])
        self.assertEqual({r[0] for r in book['Sheet1'].iter_rows(min_row=2, values_only=True)}, {'shared', 'old-only'})
        self.assertEqual(book['Sheet2']['E7'].value, 2)
        self.load.assert_not_called()

    def test_imported_mode_ignores_preview_selection_and_uses_only_imported_orders(self):
        self.workspace.save(mode='excel')
        self.load.return_value = self.production
        response = self.client.get('/api/sla-status?preview_id=1')
        self.assertEqual(response.json['source'], 'Imported Excel')
        self.assertEqual([r['Order Number'] for r in response.json['rows']], ['production-only'])
        self.assertEqual(self.client.post('/api/sla-preview', json={'preview_id': 1}).status_code, 409)
        self.assertEqual(self.workspace.read()['mode'], 'excel')

    def test_selection_persists_and_only_sla_sheet_changes(self):
        before = deepcopy(self.book.worksheet('Other report').get_all_values())
        response = self.client.post('/api/sla-preview', json={'preview_id': 1})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json['publication']['synced'])
        self.assertEqual(self.workspace.read()['sla_preview_id'], 1)
        self.assertEqual(self.client.get('/api/sla-status').json['preview_id'], 1)
        self.assertEqual(self.book.worksheet('Other report').get_all_values(), before)
        self.assertEqual(int(self.book.worksheet('SLA Status').get_all_values()[-1][-1]), 2)
        self.assertEqual(self.workspace.read()['mode'], 'tracker')
        self.assertEqual(self.workspace.read()['imported'], self.imported)
        self.load.assert_not_called()

    def test_existing_sla_styles_chart_colors_position_and_other_cells_are_preserved(self):
        first = build_sla_status(preview_snapshot(self.store, preview_id=1))
        publish_preview_sla(self.book, first, Path(self.folder.name)/'backups')
        sheet = self.book.worksheet('SLA Status')
        sheet.conditional_formats = [{'ranges': [{'sheetId': sheet.id}], 'booleanRule': {'condition': {'type': 'TEXT_EQ'}}}]
        chart = self.book.charts[sheet.id][0]
        chart['spec']['basicChart']['series'][0]['color'] = {'red': 1, 'green': .5, 'blue': .2}
        chart['position']['overlayPosition']['widthPixels'] = 777
        before = deepcopy(chart)
        publish_preview_sla(self.book, build_sla_status(preview_snapshot(self.store, preview_id=2)), Path(self.folder.name)/'backups')
        self.assertEqual(chart['position'], before['position'])
        self.assertEqual(chart['spec']['basicChart']['series'][0]['color'], before['spec']['basicChart']['series'][0]['color'])
        requests = self.book.batches[-1]['requests']
        self.assertFalse(any('deleteConditionalFormatRule' in r or 'updateEmbeddedObjectPosition' in r for r in requests))
        self.assertFalse(any('repeatCell' in r and 'userEnteredFormat' in r['repeatCell']['fields'] for r in requests))
        self.assertEqual(len(sheet.conditional_formats), 1)

    def test_status_edit_is_limited_to_the_selected_preview_and_preserves_raw_captures(self):
        before = deepcopy(self.store.history())
        response = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'Cancelled',
                    'replacement': 'Assign to ABS', 'order_numbers': ['shared']})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['updated'], 1)
        older = self.client.get('/api/sla-status?preview_id=1').json
        latest = self.client.get('/api/sla-status?preview_id=2').json
        self.assertEqual(next(r['Status'] for r in older['rows'] if r['Order Number']=='shared'), 'Assign to ABS')
        self.assertEqual(next(r['Status'] for r in latest['rows'] if r['Order Number']=='shared'), 'Search In Progress')
        self.assertEqual(self.store.history(), before)
        self.assertEqual(self.workspace.read()['imported'], self.imported)

    def test_missing_invalid_and_empty_previews_never_fall_back_to_production(self):
        for value in ('0', '-1', 'abc'):
            self.assertEqual(self.client.get('/api/sla-status?preview_id='+value).status_code, 400)
        self.assertEqual(self.client.get('/api/sla-status?preview_id=999').status_code, 404)
        self.assertEqual(self.client.post('/api/sla-preview', json={'preview_id': True}).status_code, 400)
        for number in (1, 2):
            self.store.delete(number)
        response = self.client.get('/api/sla-status')
        self.assertEqual(response.json['totals']['Grand Total'], 0)
        self.assertEqual(response.json['previews'], [])
        self.load.assert_not_called()

    def test_full_report_publication_keeps_production_reports_and_uses_preview_for_sla_only(self):
        result = self.publish()
        self.assertTrue(result['synced'])
        all_products = self.book.worksheet('All Products').get_all_values()
        self.assertEqual(all_products[1][all_products[0].index('Order Number')], 'production-only')
        self.assertEqual(int(self.book.worksheet('SLA Status').get_all_values()[-1][-1]), 2)
        self.assertEqual(int(self.book.worksheet('Status Report').get_all_values()[1][1]), 1)

    def test_raw_queue_aliases_business_statuses_and_yearless_deadlines_use_capture_time(self):
        self.store.restore(3, ['Order Number','Task Name','Task Status','SLA Expiration*'], [
            {'Order Number':'queue-1','Task Name':'Search','Task Status':'Available','SLA Expiration*':'10/16 05:00 PM'},
            {'Order Number':'queue-2','Task Name':'Search','Task Status':'Workflow Suspended','SLA Expiration*':'10/15 05:00 PM'},
            {'Order Number':'queue-3','Status':'Assign to ABS','SLA Expiration*':'2h'},
        ], created='2026-10-16T10:30:00Z')
        snapshot = preview_snapshot(self.store, preview_id=3)
        report = build_sla_status(snapshot, datetime(2026,10,16,16,15))
        self.assertEqual(report['totals']['Grand Total'], 3)
        self.assertEqual(report['excluded_count'], 0)
        current = [row for row in report['rows'] if row['Queue Presence']=='In selected preview']
        self.assertEqual([row['Status'] for row in current], ['Search In Progress','Awaiting for Clarification','Assign to ABS'])
        self.assertEqual([row['SLA Category'] for row in current], ['Today SLA','Negative SLA','Today SLA'])
        self.assertIn('2026-10-16', report['rows'][0]['SLA Expiration'])

    def test_cloud_failure_keeps_the_selection_and_local_preview_available(self):
        self.get_book.side_effect = RuntimeError('Google Sheets is offline')
        response = self.client.post('/api/sla-preview', json={'preview_id': 1})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json['publication']['synced'])
        local = self.client.get('/api/sla-status').json
        self.assertEqual(local['preview_id'], 1)
        self.assertEqual(local['totals']['Grand Total'], 2)
        self.assertIn('offline', local['sync_error'])

    def test_raw_queue_uses_tracker_business_status_for_its_orders_only(self):
        self.store.restore(3, ['Order Number','Task Name','Task Status','SLA Expiration*'], [
            {'Order Number':'production-only','Task Name':'Search','Task Status':'Available','Last User':'Named user','SLA Expiration*':'10/19 05:00 PM'},
            {'Order Number':'new-order','Task Name':'Search','Task Status':'Available','SLA Expiration*':'-1d'},
        ], created='2026-10-16T10:30:00Z')
        with patch('sla_status.tracker_status_rows', return_value=[{'Order Number': 'production-only', 'Status': 'Assign to ABS'}]):
            report = self.client.get('/api/sla-status?preview_id=3').json
        self.assertEqual(report['totals']['Grand Total'], 2)
        self.assertEqual([r['Status'] for r in report['rows'] if r['Queue Presence']=='In selected preview'], ['Assign to ABS','Search In Progress'])
        self.load.assert_not_called()
        self.assertEqual(self.workspace.read()['mode'], 'tracker')

    def test_status_lookup_reads_only_the_order_and_status_columns(self):
        book = Mock()
        book.values_get.return_value = {'values': [['No','Date','Order Number','Product','Status']]}
        book.values_batch_get.return_value = {'valueRanges': [
            {'values': [['001'], [], ['002']]}, {'values': [['Assign to ABS'], [], ['Cancelled']]}]}
        self.assertEqual(tracker_status_rows(book), [
            {'Order Number':'001','Status':'Assign to ABS'}, {'Order Number':'002','Status':'Cancelled'}])
        self.assertEqual(book.values_batch_get.call_args.args[0], ["'All Products'!C2:C", "'All Products'!E2:E"])

    def test_latest_selection_follows_new_captures_and_historical_selection_stays_selected(self):
        self.client.post('/api/sla-preview', json={'preview_id': 2})
        self.store.restore(3, self.columns, self.first, created='2026-10-16T12:00:00Z')
        self.assertEqual(self.client.get('/api/sla-status').json['preview_id'], 3)
        self.client.post('/api/sla-preview', json={'preview_id': 1})
        self.store.restore(4, self.columns, self.first, created='2026-10-17T12:00:00Z')
        self.assertEqual(self.client.get('/api/sla-status').json['preview_id'], 1)


if __name__ == '__main__':
    unittest.main()

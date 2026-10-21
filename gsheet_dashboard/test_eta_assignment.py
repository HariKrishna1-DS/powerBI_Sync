"""Business queue rules, order-level comparison and selected-order corrections."""
from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, register_report_routes
from sla_status import preview_snapshot, build_sla_status, excel_report, queue_rule, special_product_status
from test_pdf_report_workspace import ReportBook


class EtaAssignmentTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.store = PreviewStore(Path(folder.name) / 'previews')
        self.workspace = ReportWorkspace(self.store)
        self.book = ReportBook()
        self.app = Flask(__name__)
        register_report_routes(self.app, self.workspace, threading.Lock(), lambda: None,
                               lambda: (self.book, None), lambda: None)
        self.client = self.app.test_client()

    def capture(self, number, rows, created='2026-10-19T04:30:00Z'):
        columns = list(dict.fromkeys(['Order Number', 'SLA Expiration*'] +
                                    [column for row in rows for column in row]))
        self.store.restore(number, columns, rows, created=created)

    def test_abstract_and_dated_eod_or_sooner_comments(self):
        today = datetime(2026, 10, 21, 10)
        screenshot = ('ETA: Hello, This order was received on 10/08, and the fee was approved on 10/15, '
                      'and as per them, the ETA will be {} EOD or sooner.')
        cases = [
            ('Update from Abstract.', 'Assign to ABS'),
            ('ABSTRACT is working on the order.', 'Assign to ABS'),
            (screenshot.format('10/21'), 'Need to Assign ABS'),
            (screenshot.format('10/22'), 'Assign to ABS'),
            (screenshot.format('10/23'), 'Assign to ABS'),
            (screenshot.format('10/20'), 'Assign to ABS'),
            ('Received on 10/21; fee approved 10/21; ETA 10/22 EOD or sooner.', 'Assign to ABS'),
            ('ETA 10/21 sooner.', 'Need to Assign ABS'),
            ('ETA 10/21 or SOONER.', 'Need to Assign ABS'),
            ('ETA 10/21/2026\nEoD\nor sooner.', 'Need to Assign ABS'),
            ('ETA 10/21/26 EOD.', 'Need to Assign ABS'),
            ('ETA 2026-10-21 EOD.', 'Need to Assign ABS'),
            ('ETA 10/21/2025 EOD.', 'Assign to ABS'),
            ('As per the abstractor, ETA 10/21 EOD or sooner.', 'Need to Assign ABS'),
            ('ETA was 10/21 EOD; updated ETA 10/22 EOD.', 'Assign to ABS'),
            ('ETA was 10/22 EOD; updated ETA 10/21 EOD.', 'Need to Assign ABS'),
            ('Received 10/21; ETA EOD or sooner.', 'Assign to ABS'),
            ('Will finish sooner.', 'Assign to ABS'),
            ('ETA EOD.', 'Assign to ABS'),
            ('ETA 02/30 EOD or sooner.', 'Assign to ABS'),
            ('ETA 13/21 EOD or sooner.', 'Assign to ABS'),
            ('ETA 110/21 EOD.', 'Assign to ABS'),
        ]
        for comment, expected in cases:
            with self.subTest(comment=comment):
                self.assertEqual(queue_rule({'ETAComments': comment, 'TaskStatus': 'Available'}, today)[1], expected)
        for comment in ('abstractorElse', 'nonabstractors', 'abstraction', 'EODish', 'soonerish'):
            self.assertEqual(queue_rule({'ETA Comments': comment}, today)[1], '')
        self.assertEqual(queue_rule({'Comment': 'Abstract; ETA 10/21 EOD or sooner'}, today)[1], '')
        self.assertIn('2026-10-21 matches today 2026-10-21',
                      queue_rule({'ETA Comments': screenshot.format('10/21')}, today)[2])

    def test_eta_date_uses_india_today_and_keeps_specific_rules(self):
        fields = {'ETA Comments': 'ETA 10/21 EOD or sooner.'}
        self.assertEqual(queue_rule(fields, datetime(2026, 10, 20, 20, tzinfo=timezone.utc))[1], 'Need to Assign ABS')
        self.assertEqual(queue_rule(fields, datetime(2026, 10, 21, 19, tzinfo=timezone.utc))[1], 'Assign to ABS')
        self.assertEqual(queue_rule(fields, date(2026, 10, 21))[1], 'Need to Assign ABS')
        self.assertEqual(queue_rule({'ETA Comments': 'ETA 02/29 EOD'}, date(2028, 2, 29))[1], 'Need to Assign ABS')
        self.assertEqual(queue_rule({'ETA Comments': 'ETA 12/31 EOD'}, date(2027, 1, 1))[1], 'Assign to ABS')
        cases = [({'Out Time': '10/21/2026 09:15 AM'}, 'Completed and Delivered'),
                 ({'TaskName': 'CRSP2'}, 'SearchFix'),
                 ({'Product': 'Quality Control'}, 'SearchFix'),
                 ({'Product': 'Attorney Approval'}, 'Attorney Approval'),
                 ({'Product': 'Home Builders'}, 'Home Builder'),
                 ({'ETA Comments': 'ON HOLD; ETA 10/21 EOD'}, 'Waiting for Effective Date'),
                 ({'ETA Comments': 'We apologize; ETA 10/22 EOD'}, 'Need to Assign ABS')]
        for changed, expected in cases:
            with self.subTest(changed=changed):
                self.assertEqual(queue_rule(dict(fields, **changed), date(2026, 10, 21))[1], expected)

    def test_eta_date_refreshes_old_previews_deduplicates_and_keeps_manual_status(self):
        originals = [
            {'Order Number': 'same', 'ETA Comments': 'ETA 10/22 EOD'},
            {'Order Number': 'SAME ', 'ETA Comments': 'ETA 10/21 EOD'},
            {'Order Number': 'tomorrow', 'ETA Comments': 'ETA 10/22 EOD or sooner'},
            {'Order Number': 'manual', 'ETA Comments': 'ETA 10/21 sooner'},
        ]
        self.capture(1, [dict(row, **{'SLA Expiration*': '2099-10-22'}) for row in originals])
        before = deepcopy(self.store.history())
        settings = {'sla_preview_statuses': {'1': {'manual': 'QC in Progress'}}}
        for today in (datetime(2026, 10, 21, 10), datetime(2026, 10, 22, 10)):
            snapshot = preview_snapshot(self.store, settings, as_of=today)
            report = build_sla_status(snapshot, today)
            statuses = {row['Order Number'].strip().casefold(): row['Status'] for row in report['rows']}
            self.assertEqual(statuses, {'same': 'Need to Assign ABS', 'manual': 'QC in Progress',
                'tomorrow': 'Need to Assign ABS' if today.day == 22 else 'Assign to ABS'})
            self.assertEqual(report['totals']['Grand Total'], 3)
        self.assertEqual(self.store.history(), before)

    def test_eta_statuses_reconcile_across_api_excel_and_sheets(self):
        today = datetime(2026, 10, 21, 10)
        self.capture(1, [
            {'Order Number': 'today', 'ETA Comments': 'ETA 10/21 EOD or sooner', 'SLA Expiration*': '1d'},
            {'Order Number': 'future', 'ETA Comments': 'Received 10/21; ETA 10/22 EOD or sooner', 'SLA Expiration*': '1d'},
            {'Order Number': 'plain', 'ETA Comments': 'Update from Abstract', 'SLA Expiration*': '1d'},
        ])
        before = deepcopy(self.store.history())
        with patch('sla_status.queue_rule', side_effect=lambda row, as_of=None: queue_rule(row, today)):
            report = self.client.get('/api/sla-status?preview_id=1').json
            expected = {'today': 'Need to Assign ABS', 'future': 'Assign to ABS', 'plain': 'Assign to ABS'}
            self.assertEqual({row['Order Number']: row['Status'] for row in report['rows']}, expected)
            self.assertEqual({row['Status']: row['Grand Total'] for row in report['summary']},
                             {'Assign to ABS': 2, 'Need to Assign ABS': 1})
            book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
            self.assertEqual({row[0]: row[11] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)}, expected)
            edited = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'Need to Assign ABS',
                'replacement': 'QC in Progress', 'order_numbers': ['today'], 'edit_scope': 'orders'})
            self.assertEqual(edited.status_code, 200, edited.json)
            self.assertTrue(edited.json['publication']['synced'])
            updated = self.client.get('/api/sla-status?preview_id=1').json
            self.assertEqual({row['Order Number']: row['Status'] for row in updated['rows']},
                             dict(expected, today='QC in Progress'))
            sheet = self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')
            self.assertEqual({row[0]: row[-1] for row in sheet[2:-1]}, {'Assign to ABS': 2, 'QC in Progress': 1})
        self.assertEqual(self.store.history(), before)

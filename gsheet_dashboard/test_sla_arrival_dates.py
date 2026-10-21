"""Arrival preservation and date-scoped Excel SLA reports share one dataset."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
import pandas as pd

from datatrace_sync import powerbi_table
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, imported_snapshot, read_report_files, register_report_routes, report_frames
from reporting_dates import arrival_details
from sla_status import build_sla_status, imported_date_snapshot, preview_snapshot, sheet_summary_values
from test_pdf_report_workspace import ReportBook, row, excel
from tracker_sync import key


class SlaArrivalDatesTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.store = PreviewStore(Path(self.folder.name) / 'previews')
        self.workspace = ReportWorkspace(self.store)
        self.book = ReportBook()
        self.book.add('Other report', [{'Keep': 'unchanged'}])
        self.rows = [
            row('old', date='10/12/2026 09:15 AM', **{'SLA Expiration': '2026-10-13'}),
            row('future', date='10/13/2026 01:02:03 PM', **{'SLA Expiration': '2026-10-14'}),
            row('today', date='10/13/2026 10:00 AM', **{'SLA Expiration': '2026-10-13 23:59:00'}),
            row('late', date='10/13/2026 11:00 AM', **{'SLA Expiration': '2026-10-12 23:59:00'}),
            row('no-deadline', date='10/13/2026 12:00 PM'),
            row('date-only', date='', **{'Date': '2026-10-13', 'SLA Expiration': '2026-10-13'}),
            row('undated', date='', **{'SLA Expiration': '2099-10-14'}),
        ]
        self.imported = read_report_files([('full.xlsx', excel(self.rows))])
        self.workspace.save(mode='excel', imported=self.imported, cloud_reset=True)
        self.gate = threading.Lock()
        self.app = Flask(__name__)
        def load(**kwargs):
            return imported_snapshot(self.workspace.read()['imported'])
        register_report_routes(self.app, self.workspace, self.gate, load, lambda: (self.book, None), lambda: None)
        self.client = self.app.test_client()

    def test_arrival_date_uses_source_calendar_and_preserves_seconds(self):
        for raw, day in [('10/13/2026 01:02:03 PM', '2026-10-13'),
                         ('2026-10-13T23:59:00-07:00', '2026-10-13'),
                         ('46308.5', '2026-10-13')]:
            with self.subTest(raw=raw):
                self.assertEqual(arrival_details({'Arrival Time': raw}), {'Arrival Date': day, 'Arrival Time': raw})
        self.assertEqual(arrival_details({'Date': '2026-10-13'}), {'Arrival Date': '2026-10-13', 'Arrival Time': ''})
        self.assertEqual(arrival_details({'Date': 'invalid', 'In-Time': '10/13/2026 09:00 AM'})['Arrival Date'], '2026-10-13')
        self.assertEqual(arrival_details({'Date': '2026-10-13', 'In-Time': '10/12/2026 09:00 AM'})['Arrival Date'], '2026-10-13')
        self.assertEqual(arrival_details({}), {'Arrival Date': '', 'Arrival Time': ''})

    def test_capture_saves_arrival_fields_and_sla_preview_exposes_them(self):
        frame = pd.DataFrame([{'Order Number': 'capture', 'Arrival Time': '10/13/2026 01:02:03 PM',
                               'Task Status': 'Available', 'Task Name': 'Search', 'SLA Expiration*': '1d'}])
        frame = powerbi_table(frame, '2026-10-19T12:00:00Z')
        captured = self.store.save(frame)
        self.assertEqual(captured['rows'][0]['Arrival Date'], '2026-10-13')
        self.assertEqual(captured['rows'][0]['Arrival Time'], '10/13/2026 01:02:03 PM')
        report = build_sla_status(preview_snapshot(self.store), datetime(2026, 10, 19))
        self.assertEqual(report['rows'][0]['Arrival Time'], captured['rows'][0]['Arrival Time'])
        self.assertIn('Arrival Date', report['columns'])
        self.assertIn('Arrival Time', report['columns'])

    def test_default_is_latest_arrival_date_and_categories_use_that_date(self):
        report = self.client.get('/api/sla-status').json
        self.assertEqual(report['report_date'], '2026-10-13')
        self.assertEqual(report['as_of'], '2026-10-13')
        self.assertEqual(report['report_dates'], [{'date': '2026-10-13', 'orders': 5},
            {'date': '2026-10-12', 'orders': 1}, {'date': 'Undated', 'orders': 1}])
        self.assertEqual(report['date_order_count'], 5)
        self.assertEqual(report['imported_order_count'], 7)
        self.assertEqual(report['totals'], {'Status': 'Grand Total', 'Future SLA': 1,
            'Negative SLA': 1, 'Today SLA': 2, 'Grand Total': 4})
        self.assertEqual(report['excluded_count'], 1)
        self.assertEqual(len({key(row) for row in report['rows'] + report['excluded_rows']}), 5)
        self.assertEqual(next(r for r in report['rows'] if r['Order Number'] == 'date-only')['Arrival Time'], '')

    def test_historical_selection_persists_and_publishes_only_matching_sla(self):
        before = deepcopy(self.book.worksheet('Other report').get_all_values())
        response = self.client.post('/api/sla-report-date', json={'date': '2026-10-12'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json['publication']['synced'])
        report = self.client.get('/api/sla-status').json
        self.assertEqual(report['report_date'], '2026-10-12')
        self.assertEqual(report['totals']['Future SLA'], 1)
        self.assertEqual(self.workspace.read()['sla_import_date'], '2026-10-12')
        self.assertEqual(self.book.worksheet('Other report').get_all_values(), before)
        expected = sheet_summary_values(report)
        self.assertEqual(self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE'),
                         [values[:2] if index == 0 else values for index, values in enumerate(expected)])
        frames, _ = report_frames(imported_snapshot(self.imported), self.workspace.read())
        self.assertEqual(frames['SLA Status'], expected)
        # Picking the latest date follows newer imports instead of pinning a day.
        self.client.post('/api/sla-report-date', json={'date': '2026-10-13'})
        self.assertEqual(self.workspace.read()['sla_import_date'], '')

    def test_download_and_detail_rows_share_date_arrival_fields_and_totals(self):
        report = self.client.get('/api/sla-status?date=2026-10-13').json
        download = self.client.get('/api/export/report?sheet=SLA%20Status&date=2026-10-13')
        self.assertEqual(download.status_code, 200)
        self.assertIn('2026-10-13.xlsx', download.headers['Content-Disposition'])
        book = load_workbook(BytesIO(download.data), data_only=True)
        sheet = book['Sheet1']
        headers = [cell.value for cell in sheet[1]]
        records = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(records), report['totals']['Grand Total'])
        for original in report['rows']:
            record = next(r for r in records if r['Order Number'] == original['Order Number'])
            self.assertEqual(record['Arrival Date'], original['Arrival Date'])
            self.assertEqual(record['Arrival Time'] or '', original['Arrival Time'])
        self.assertIn('Arrival date: 2026-10-13', book['Sheet2']['A1'].value)
        pivot = book['Sheet2']._pivots[0]
        self.assertEqual(pivot.cache.recordCount, len(report['rows']))
        # A date in a download URL must not change the selected UI date.
        self.assertNotIn('sla_import_date', self.workspace.read())

    def test_invalid_date_is_rejected_and_undated_orders_remain_accessible(self):
        for url in ('/api/sla-status?date=2026-02-30', '/api/export/report?sheet=SLA%20Status&date=2026-10-11'):
            self.assertEqual(self.client.get(url).status_code, 400)
        self.assertEqual(self.client.post('/api/sla-report-date', json={'date': []}).status_code, 400)
        undated = self.client.get('/api/sla-status?date=Undated').json
        self.assertEqual([r['Order Number'] for r in undated['rows']], ['undated'])
        self.workspace.save(mode='tracker')
        self.assertEqual(self.client.post('/api/sla-report-date', json={'date': '2026-10-13'}).status_code, 409)

    def test_status_edit_is_scoped_to_selected_date_and_keeps_other_days(self):
        report = self.client.get('/api/sla-status').json
        identities = [key(r) for r in report['rows'] + report['excluded_rows']]
        with patch('threading.Thread.start'):
            response = self.client.post('/api/order-status', json={'status': 'Search In Progress',
                'replacement': 'QC in Progress', 'order_numbers': identities,
                'sla_report_date': '2026-10-13'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['updated'], 5)
        statuses = {r['Order Number']: r['Status'] for r in self.workspace.read()['imported']['rows']}
        self.assertEqual(statuses['old'], 'Search In Progress')
        self.assertEqual(statuses['undated'], 'Search In Progress')
        self.assertEqual(statuses['future'], 'QC in Progress')

    def test_import_and_source_switch_reset_historical_selection_to_latest(self):
        self.workspace.save(sla_import_date='2026-10-12')
        with patch('threading.Thread.start'):
            self.client.post('/api/report-source', json={'mode': 'excel'})
        self.assertEqual(self.client.get('/api/sla-status').json['report_date'], '2026-10-13')
        self.workspace.save(sla_import_date='2026-10-12')
        with patch('threading.Thread.start'):
            response = self.client.post('/api/report-import', data={'files': (BytesIO(excel([
                row('newest', date='10/14/2026 09:00 AM', **{'SLA Expiration': '2026-10-14'})])), 'new.xlsx')})
        self.assertEqual(response.status_code, 201)
        latest = self.client.get('/api/sla-status').json
        self.assertEqual(latest['report_date'], '2026-10-14')
        self.assertEqual([r['Order Number'] for r in latest['rows']], ['newest'])

    def test_workbook_identity_survives_date_filter_and_missing_dates_never_use_import_time(self):
        data = read_report_files([(name, excel([self.rows[1]])) for name in ('full.xlsx', 'remaining.xlsx')])
        report = build_sla_status(imported_date_snapshot(imported_snapshot(data)))
        self.assertEqual(report['totals']['Grand Total'], 2)
        self.assertEqual(len({key(r) for r in report['rows']}), 2)
        data = read_report_files([('empty-dates.xlsx', excel([self.rows[-1]]))])
        report = build_sla_status(imported_date_snapshot(imported_snapshot(data)))
        self.assertEqual(report['report_date'], 'Undated')
        self.assertEqual(report['rows'][0]['Arrival Date'], '')
        self.assertEqual(report['rows'][0]['Arrival Time'], '')
        empty = imported_date_snapshot({'sheets': {'Overview': {'rows': [], 'columns': []}}})
        self.assertEqual(build_sla_status(empty)['totals']['Grand Total'], 0)

    def test_relative_import_deadlines_use_selected_report_day(self):
        data = read_report_files([('relative.xlsx', excel([
            row('hours', date='10/13/2026 09:00 AM', **{'SLA Expiration': '2h'}),
            row('tomorrow', date='10/13/2026 09:00 AM', **{'SLA Expiration': '1d'}),
        ]))])
        data['created'] = '2026-10-19T12:00:00Z'
        report = build_sla_status(imported_date_snapshot(imported_snapshot(data)))
        self.assertEqual(report['totals']['Today SLA'], 1)
        self.assertEqual(report['totals']['Future SLA'], 1)


if __name__ == '__main__':
    unittest.main()

from datetime import datetime
from io import BytesIO
import unittest
from openpyxl import load_workbook
from sla_status import build_sla_status, sla_queue_rule, excel_report, sheet_summary_values


class SuspendedSlaTests(unittest.TestCase):
    def test_on_hold_requires_suspended_task_status_for_clarification(self):
        for task in ('Search', 'SearchFix', 'TypingModule'):
            for comment in ('ON HOLD - Awaiting response', 'on\n hold', 'On Hold'):
                with self.subTest(task=task, comment=comment):
                    self.assertEqual(sla_queue_rule({'Task Name': task, 'Task Status': 'Workflow Suspended',
                        'ETA Comments': comment})[1], 'Awaiting for Clarification')
        self.assertEqual(sla_queue_rule({'Task Name': 'Search', 'Task Status': 'Available',
            'ETA Comments': 'ON HOLD - Awaiting response'})[1], 'Waiting for Effective Date')

    def test_clarification_overrides_searchfix_and_vendor_assignment(self):
        for fields in [
            {'Task Name': 'Search', 'Comment': 'Cannot Complete Order', 'ETA Comments': 'assigned to the abstractor'},
            {'Task Name': 'Workflow Suspended', 'Comments': 'CANNOT complete'},
            {'Task Name': 'SearchFix', 'ETA Comments': 'ON HOLD'},
            {'Task Name': 'CRSP2'},
            {'Task Name': 'Search', 'ETA Comments': 'Abstraction pending'},
            {'Task Name': 'Search', 'ETA Comments': 'Order assigned to the abst..'},
        ]:
            with self.subTest(fields=fields):
                row = {'Task Status': 'Workflow Suspended', **fields}
                self.assertEqual(sla_queue_rule(row)[1], 'Awaiting for Clarification')
        self.assertEqual(sla_queue_rule({'Task Name': 'SearchFix', 'Task Status': 'Available'})[1], 'SearchFix')
        self.assertEqual(sla_queue_rule({'Task Name': 'Search', 'Task Status': 'Available',
                                       'ETA Comments': 'assigned to the abstractor'})[1], 'Assign to ABS')

    def test_unknown_deadlines_count_in_details_summary_and_exports(self):
        snapshot = {'report_scope': 'preview', 'updated_at': '2026-10-29T09:00:00Z',
                    'sheets': {'Overview': {'columns': [], 'rows': [
                        {'Order Number': 'missing', 'Status': 'Awaiting for Clarification'},
                        {'Order Number': 'invalid', 'Status': 'SearchFix', 'SLA Expiration': 'PAUSED'},
                        {'Order Number': 'dated', 'Status': 'SearchFix', 'SLA Expiration': '2026-10-30'},
                    ]}}}
        report = build_sla_status(snapshot, datetime(2026, 10, 29))
        self.assertEqual(report['totals']['Grand Total'], 3)
        self.assertEqual(report['totals']['Unknown SLA'], 2)
        self.assertEqual(report['excluded_count'], 0)
        self.assertEqual({r['Order Number'] for r in report['rows']}, {'missing', 'invalid', 'dated'})
        self.assertEqual(len(sheet_summary_values(report)[1]), 6)
        book = load_workbook(BytesIO(excel_report(report).getvalue()))
        self.assertEqual(book['Sheet1'].max_row, 4)
        self.assertEqual(book['Sheet2']['E4'].value, 'Unknown SLA')
        self.assertEqual(book['Sheet2']['F7'].value, 3)
        self.assertEqual(book['Sheet2']._pivots[0].cache.recordCount, 3)

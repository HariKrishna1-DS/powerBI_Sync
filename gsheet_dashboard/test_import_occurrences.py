"""Repeated Excel order numbers are independent source occurrences."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook
from import_identity import order_key, upgrade_workbook_orders
from import_excel_changes import editable_frames, reconcile, preserve_sheet_edits
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, read_report_files, imported_snapshot, report_frames, register_report_routes, combine_sources
from sla_status import build_sla_status, imported_date_snapshot
from test_pdf_report_workspace import excel


def occurrence(date, status='Completed and Delivered', out='', product='Legal & Vesting'):
    return {'Order Number':'SYNTHETIC-REPEAT-1','Date':date,'In-Time':date,'Product':product,
            'Status':status,'Out Time':out,'SLA Expiration':'2099-10-20 17:00'}


class ImportedOccurrenceTests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.workspace=ReportWorkspace(PreviewStore(Path(self.folder.name)/'previews'))
        self.raw=[occurrence('10/01/2026',out='10/01/2026'),
                  occurrence('10/14/2026','Waiting for Effective Date',product='Update')]
        self.imported=read_report_files([('full.xlsx',excel(self.raw))])
        self.workspace.save(mode='excel',cloud_reset=True,imported=self.imported,
                            files=self.imported['files'],row_count=2)

    def snapshot(self, **kwargs):
        return imported_snapshot(self.workspace.read()['imported'])

    def client(self):
        app=Flask(__name__)
        register_report_routes(app,self.workspace,threading.Lock(),self.snapshot,lambda:None,lambda:None)
        return app.test_client()

    def test_repeated_order_keeps_completed_earlier_row_and_pending_later_row(self):
        self.assertEqual(len(self.imported['rows']),2)
        self.assertEqual(len({order_key(row) for row in self.imported['rows']}),2)
        frames,report=report_frames(self.snapshot(),{'mode':'excel','capacity_month':'2026-10'})
        self.assertEqual([(row['Date'],row['Received'],row['Completed']) for row in report['daily']],
                         [('2026-10-01',1,1),('2026-10-14',1,0)])
        self.assertEqual([row['Status'] for row in self.imported['rows']],
                         ['Completed and Delivered','Waiting for Effective Date'])
        self.assertEqual(len(frames['All Products'])-1,2)
        public=self.workspace.public()
        self.assertEqual((public['input_row_count'],public['row_count'],public['duplicate_row_count']),(2,2,1))

    def test_every_completed_occurrence_counts_even_when_its_other_values_are_identical(self):
        same=occurrence('10/14/2026',out='10/15/2026')
        rows=[occurrence('10/01/2026',out='10/02/2026'),same,same,
              occurrence('10/14/2026','Waiting for Effective Date',product='Update')]
        imported=read_report_files([('full.xlsx',excel(rows))])
        self.assertEqual(len(imported['rows']),4)
        self.assertEqual(len({order_key(row) for row in imported['rows']}),4)
        snapshot=imported_snapshot(imported)
        _,report=report_frames(snapshot,{'mode':'excel','capacity_month':'2026-10'})
        self.assertEqual(sum(day['Received'] for day in report['daily']),4)
        self.assertEqual(sum(day['Completed'] for day in report['daily']),3)
        self.assertEqual(next(day['Completed'] for day in report['daily'] if day['Date']=='2026-10-15'),2)
        self.assertEqual(build_sla_status(imported_date_snapshot(snapshot,'2026-10-14'))['totals']['Grand Total'],3)

    def test_matching_row_numbers_in_two_worksheets_stay_separate(self):
        book=Workbook();book.active.title='First'
        for sheet in (book.active,book.create_sheet('Second')):
            sheet.append(list(self.raw[0]));sheet.append(list(self.raw[0].values()))
        stream=BytesIO();book.save(stream)
        imported=read_report_files([('full.xlsx',stream.getvalue())])
        self.assertEqual({row['Source row ID'] for row in imported['rows']},{'First!2','Second!2'})
        self.assertEqual(len({order_key(row) for row in imported['rows']}),2)

    def test_old_import_restores_earlier_occurrence_and_preserves_only_later_row_edits(self):
        legacy=deepcopy(self.imported)
        legacy['order_identity']='workbook'
        legacy['rows']=[deepcopy(legacy['rows'][-1])]
        legacy['rows'][0].pop('Source row ID')
        legacy['rows'][0]['Status']='Typing in Progress'
        saved=deepcopy(legacy)
        upgraded=upgrade_workbook_orders(legacy)
        self.assertEqual([row['Status'] for row in upgraded['rows']],
                         ['Completed and Delivered','Typing in Progress'])
        self.assertEqual([row['Out Time'] for row in upgraded['rows']],['10/01/2026',''])
        self.assertEqual(legacy,saved)
        self.assertIs(upgrade_workbook_orders(upgraded),upgraded)
        self.workspace.save(imported=legacy,row_count=1)
        reopened=ReportWorkspace(self.workspace.store)
        self.assertEqual(reopened.public()['row_count'],2)
        self.assertTrue(reopened.read()['cloud_reset'])
        self.assertTrue(list((Path(self.folder.name)/'backups').glob('before-update-*.zip')))

    def test_google_sheet_edit_and_removal_affect_only_the_selected_occurrence(self):
        baseline=editable_frames(self.snapshot(),{'mode':'excel'})
        unchanged,events=reconcile(self.imported,baseline,deepcopy(baseline))
        self.assertFalse(events)
        self.assertEqual(unchanged,self.imported)
        remote=deepcopy(baseline)
        values=remote['All Products'];headers=values[0]
        values[1][headers.index('Out Time')]='10/05/2026'
        updated,events=reconcile(self.imported,baseline,remote)
        self.assertEqual(len(events),1)
        self.assertEqual([row['Out Time'] for row in updated['rows']],['10/05/2026',''])
        self.assertEqual(updated['rows'][1]['Status'],'Waiting for Effective Date')
        remote=deepcopy(baseline);remote['All Products'].pop(1)
        removed,events=reconcile(self.imported,baseline,remote)
        self.assertEqual(len(removed['rows']),1)
        self.assertEqual(removed['rows'][0]['Status'],'Waiting for Effective Date')
        replacement=read_report_files([('other.xlsx',excel([dict(self.raw[1],**{'Order Number':'OTHER'})]))])
        rebuilt=combine_sources(self.imported['sources']+replacement['sources'])
        preserved=preserve_sheet_edits(removed,rebuilt,['other.xlsx'])
        self.assertEqual(len(preserved['rows']),2)
        self.assertNotIn('Completed and Delivered',[row['Status'] for row in preserved['rows']])
        bad=deepcopy(baseline);index=bad['All Products'][0].index('Source row ID')
        for cells in bad['All Products']:cells.pop(index)
        with self.assertRaisesRegex(ValueError,'Source row ID'):
            reconcile(self.imported,baseline,bad)

    def test_sla_editor_targets_one_of_two_occurrences_on_the_same_date(self):
        imported=read_report_files([('full.xlsx',excel([self.raw[0],dict(self.raw[1],Date='10/01/2026',**{'In-Time':'10/01/2026'})]))])
        self.workspace.save(imported=imported)
        client=self.client()
        record=order_key(imported['rows'][1])
        lookup=client.get('/api/sla-order',query_string={'mode':'excel','date':'2026-10-01',
            'order_number':'SYNTHETIC-REPEAT-1','source_workbook':'full.xlsx','record_key':record})
        self.assertEqual(lookup.status_code,200,lookup.json)
        self.assertEqual(lookup.json['row']['Status'],'Waiting for Effective Date')
        with patch('threading.Thread.start'):
            saved=client.post('/api/sla-order',json={'mode':'excel','report_date':'2026-10-01',
                'order_number':'SYNTHETIC-REPEAT-1','source_workbook':'full.xlsx','record_key':record,
                'status':'Typing in Progress','expected_revision':lookup.json['revision'],
                'fields':{'SLA Expiration':'2099-10-20 17:00'}})
        self.assertEqual(saved.status_code,200,saved.json)
        self.assertEqual([row['Status'] for row in self.workspace.read()['imported']['rows']],
                         ['Completed and Delivered','Typing in Progress'])

    def test_capacity_editor_and_export_target_the_specific_row(self):
        client=self.client()
        self.assertEqual(client.get('/api/order-details?mode=excel&order=SYNTHETIC-REPEAT-1&source_workbook=full.xlsx').status_code,409)
        record=order_key(self.imported['rows'][0])
        result=client.get('/api/order-details',query_string={'mode':'excel','order':'SYNTHETIC-REPEAT-1',
                          'source_workbook':'full.xlsx','record_key':record})
        self.assertEqual(result.status_code,200,result.json)
        editor=result.json
        self.assertNotIn('Source row ID',editor['editable_columns'])
        with patch('threading.Thread.start'):
            saved=client.post('/api/order-details',json={'mode':'excel','order_number':'SYNTHETIC-REPEAT-1',
                'source_workbook':'full.xlsx','record_key':record,'version':editor['version'],
                'report_date':'2026-10-01','changes':{'Out Time':'10/05/2026'}})
        self.assertEqual(saved.status_code,200,saved.json)
        rows=self.workspace.read()['imported']['rows']
        self.assertEqual([row['Out Time'] for row in rows],['10/05/2026',''])
        detail=client.get('/api/capacity-report/details?date=2026-10-05').json
        completed=next(group['rows'] for group in detail['groups'] if group['key']=='completed')
        self.assertEqual(len(completed),1)
        self.assertEqual(completed[0]['Source row ID'],self.imported['rows'][0]['Source row ID'])
        export=client.get('/api/export/report?sheet=All%20Products')
        self.assertEqual(export.status_code,200)
        book=load_workbook(BytesIO(export.data),data_only=True)
        sheet=book['All Products'];values=list(sheet.values)
        number=values[0].index('Order Number')
        self.assertEqual([cells[number] for cells in values[1:]],['SYNTHETIC-REPEAT-1']*2)
        self.assertIn('Source row ID',values[0])


if __name__=='__main__':unittest.main()

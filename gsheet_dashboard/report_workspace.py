"""Persist report source selection and publish the PDF's public reporting tabs.

Imported workbooks are a separate report dataset. They never become queue captures
and therefore cannot infer completion of missing tracker orders.
"""
from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import csv
import gzip
import hashlib
import json
import re
import secrets
import threading

DAILY_COLUMNS = ['Date', 'Received', 'Completed', 'Clarification', 'Cancelled',
                 'Vendor Pending', 'In-House Pending', 'SLA OnTime', 'Missing']
CAPACITY_COLUMNS = DAILY_COLUMNS[:7] + ['Capacity', 'Ext capacity']


class ReportWorkspace:
    def __init__(self, store):
        self.store = store
        if not hasattr(self, "lock"):
            self.lock = threading.RLock()
        with store.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS report_workspace (id INTEGER PRIMARY KEY, payload TEXT NOT NULL)')
        from import_excel_changes import initialize_history
        initialize_history(self)
        settings = self.read()
        imported = settings.get('imported') or {}
        if imported.get('sources') and imported.get('order_identity') != 'workbook':
            from import_identity import upgrade_workbook_orders
            from tracker_sync import sync_lock
            from workspace_backup import save_safety_backup
            with sync_lock(store.root), self.lock:
                latest = self.read()
                if latest.get('imported', {}).get('order_identity') != 'workbook':
                    updated = upgrade_workbook_orders(latest['imported'])
                    save_safety_backup(store, 'before-update')
                    self.save(imported=updated, files=updated['files'], row_count=len(updated['rows']),
                              cloud_baseline=None, cloud_reset=True, cloud_refresh_error=None,
                              **({'publication_state': 'pending'} if latest['mode'] == 'excel' else {}))

    def read(self):
        with self.store.connect() as db:
            row = db.execute('SELECT payload FROM report_workspace WHERE id=1').fetchone()
        return json.loads(row[0]) if row else {'mode': 'tracker', 'enabled': True,
            'selected_date': '', 'capacity': 700, 'extended_capacity': 750}

    def save(self, **changes):
        with self.lock:
            value = dict(self.read(), **changes)
            if {'mode', 'imported', 'selected_date', 'capacity', 'extended_capacity'} & changes.keys():
                value['data_revision'] = secrets.token_hex(12)
            with self.store.connect() as db:
                db.execute('INSERT OR REPLACE INTO report_workspace VALUES (1,?)', (json.dumps(value),))
            return value

    def public(self):
        value = self.read()
        public = {k: v for k, v in value.items() if k not in ('imported', 'cloud_baseline', 'tracker_status_baseline')}
        effective = Counter(row.get('Source workbook', '') for row in (value.get('imported') or {}).get('rows', []))
        public['imports'] = [{'name': s['name'], 'rows': len(s['rows']), 'orders': len((value.get('imported') or {}).get('rows', [])) if s.get('legacy') else effective[s['name']],
                              'worksheets': len(s['files']), 'legacy': s.get('legacy', False)}
                             for s in import_sources(value.get('imported'))]
        public['input_row_count'] = sum(file.get('rows', 0) for file in value.get('files', []))
        public['worksheet_data_rows'] = sum(file.get('rows_seen', file.get('rows', 0)) for file in value.get('files', []))
        public['missing_order_number_rows'] = sum(file.get('missing_order_number', 0) for file in value.get('files', []))
        public['row_diagnostics_available'] = bool(value.get('files')) and all(
            'rows_seen' in file for file in value.get('files', []))
        imported = value.get('imported') or {}
        diagnostics = (imported if imported.get('diagnostics_version') == 4 else
                       combine_sources(import_sources(imported))) if imported else {}
        repeated = diagnostics.get('duplicate_orders', [])
        public['duplicate_orders'] = [row for row in repeated if row['Occurrence type'] == 'Within workbook']
        public['source_overlap_orders'] = diagnostics.get('source_overlap_orders', [])
        public['duplicate_row_count'] = len(public['duplicate_orders'])
        public['source_overlap_count'] = len(public['source_overlap_orders'])
        public['repeated_source_row_count'] = len(repeated)
        public['duplicate_order_numbers'] = list(dict.fromkeys(row['Order Number'] for row in public['duplicate_orders']))
        if imported:
            public['files'] = diagnostics['files']
        return public


def import_sources(imported):
    """Older releases without provenance remain one removable batch."""
    if not imported:
        return []
    if 'sources' in imported:
        return imported['sources']
    # Old releases did not store file provenance. Keep ambiguous batches intact.
    return [{'name': 'Previous imported batch', 'legacy': True, 'rows': imported['rows'],
             'columns': imported['columns'], 'files': imported.get('files', []),
             'reviews': imported.get('reviews', [])}]


def combine_sources(sources):
    from import_identity import combine_sources as combine
    return combine(sources)


def read_report_files(uploads):
    """Read all order worksheets, preserving text IDs and positional headers."""
    from openpyxl import load_workbook
    from tracker_sync import canonical_headers, text, key, STATUS_NAMES
    from production_timing import sla_result
    if not 1 <= len(uploads) <= 30:
        raise ValueError('Choose between 1 and 30 Excel or CSV files.')
    saved, columns, counts, reviews, sources = {}, [], [], [], []
    total = 0
    for name, content in uploads:
        name = Path(name.replace('\\', '/')).name
        suffix = Path(name).suffix.lower()
        book = None
        if suffix == '.xlsx':
            try:
                with ZipFile(BytesIO(content)) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                        raise ValueError(f'{name}: expanded workbook exceeds 100 MB.')
                book = load_workbook(BytesIO(content), read_only=True, data_only=True)
                tabs = [(sheet.title, sheet.iter_rows()) for sheet in book.worksheets
                        if sheet.max_column <= 200 and sheet.max_row <= 100001]
                if len(tabs) != len(book.worksheets):
                    raise ValueError(f'{name}: a sheet exceeds 100,000 rows or 200 columns.')
            except ValueError:
                raise
            except Exception as exc:
                raise ValueError(f'{name}: cannot read this XLSX workbook.') from exc
        elif suffix == '.csv':
            try:
                tabs = [(name, iter(csv.reader(content.decode('utf-8-sig').splitlines())))]
            except UnicodeError as exc:
                raise ValueError(f'{name}: CSV must use UTF-8 encoding.') from exc
        else:
            raise ValueError(f'{name}: choose XLSX or CSV files.')
        found = False
        file_rows, file_columns, file_counts, file_reviews = [], [], [], []
        try:
            for title, iterator in tabs:
                headers = None
                for index, cells in enumerate(iterator):
                    raw = [cell.value if hasattr(cell, 'value') else cell for cell in cells]
                    while raw and raw[-1] is None:
                        raw.pop()
                    candidate = canonical_headers(raw)
                    if 'Order Number' in candidate:
                        headers = candidate
                        break
                    if index >= 19:
                        break
                if not headers:
                    continue  # Summary/chart worksheets are not order datasets.
                if not any(c in headers for c in ('Status', 'Task Status')):
                    raise ValueError(f'{name} / {title}: include a Status or Task Status column.')
                found = True
                columns = list(dict.fromkeys(columns + [c for c in headers if c]))
                file_columns = list(dict.fromkeys(file_columns + [c for c in headers if c]))
                count = {'file': name, 'sheet': title, 'rows': 0, 'rows_seen': 0,
                         'missing_order_number': 0, 'duplicates': 0}
                for source_row, cells in enumerate(iterator, index + 2):
                    values = []
                    for cell in cells:
                        value = cell.value if hasattr(cell, 'value') else cell
                        # Excel often stores order IDs as numbers with a 000000 format.
                        fmt = getattr(cell, 'number_format', '')
                        if isinstance(value, int) and re.fullmatch('0+', fmt or ''):
                            value = str(value).zfill(len(fmt))
                        values.append(text(value))
                    row = {h: v for h, v in zip(headers, values) if h}
                    if not any(values):
                        continue
                    count['rows_seen'] += 1
                    if not key(row):
                        count['missing_order_number'] += 1
                        continue
                    total += 1
                    if total > 100000:
                        raise ValueError('The combined import exceeds 100,000 orders.')
                    for target, aliases in {'In-Time': ('Arrival Time', 'RequestArrivalTime', 'In Time'),
                                             'Date': ('Arrival Date',), 'SLA Expiration': ('SLA Expiration*',)}.items():
                        if not row.get(target):
                            row[target] = next((row[a] for a in aliases if row.get(a)), '')
                    if not row.get('Product'):
                        hint = (name + ' ' + title).casefold().replace('_', ' ')
                        if 'full' in hint or title == 'TV Orders':
                            row['Product'] = 'Full Search'
                        elif any(term in hint for term in ('remaining', 'co update', 'c-o', 'current owner')):
                            row['Product'] = 'Remaining Products'
                        else:
                            raise ValueError(f'{name} / {title}: order {row["Order Number"]} needs a Product.')
                    status = ' '.join((row.get('Status') or row.get('Task Status', '')).casefold().split())
                    row['Status'] = STATUS_NAMES.get(status, {'canceled': 'Cancelled', 'completed': 'Completed and Delivered'}.get(status, row.get('Status') or row.get('Task Status', '')))
                    result, reason = sla_result(row)
                    if row.get('Free Site') not in ('On Time', 'Missing'):
                        row['Free Site'] = result
                    if reason:
                        reviews.append({'Order Number': row['Order Number'], 'Reason': reason})
                        file_reviews.append({'Order Number': row['Order Number'], 'Reason': reason})
                    if key(row) in saved:
                        count['duplicates'] += 1
                    row['_source_row'] = source_row
                    saved[key(row)] = row  # Last file / worksheet wins, explicitly reported.
                    file_rows.append(row)
                    count['rows'] += 1
                counts.append(count)
                file_counts.append(count)
            if not found:
                raise ValueError(f'{name}: no worksheet has an Order Number header in its first 20 rows.')
            sources.append({'name': name, 'rows': file_rows, 'columns': file_columns,
                            'files': file_counts, 'reviews': file_reviews})
        finally:
            if book:
                book.close()
    if not saved:
        raise ValueError('No order rows were found. The existing report was kept.')
    return combine_sources(sources)


def imported_snapshot(imported):
    from tracker_sync import FULL, REMAINING, HEADERS, sheet_reports, key
    from monthly_production import arrival_sort, local_datetime
    rows = [dict(row, _record_key=key(row)) for row in imported['rows']]
    for row in rows:
        received = local_datetime(row.get('Date')) or local_datetime(row.get('In-Time'))
        row['_received_at'] = received.isoformat() if received else ''
    # Completed imported orders move to the completion month only when Out Time
    # is later than their arrival month. Daily reports retain the arrival date.
    for row in rows:
        if str(row.get('Status') or row.get('Task Status') or '').strip().casefold() != 'completed and delivered':
            continue
        completed_at = local_datetime(row.get('Out Time'))
        arrived_at = local_datetime(row.get('_received_at'))
        if completed_at and (not arrived_at or completed_at.strftime('%Y-%m') > arrived_at.strftime('%Y-%m')):
            row['_month'] = completed_at.strftime('%Y-%m')
    rows = arrival_sort(rows)
    rows = [dict(row, No=index) for index, row in enumerate(rows, 1)]
    columns = list(dict.fromkeys(HEADERS + imported['columns']))
    members, owners = {True: set(), False: set()}, {}
    from import_identity import source_group
    source_groups = {source['name']: source_group(source['name']) for source in import_sources(imported)}
    for source in import_sources(imported):
        if sum(metadata['rows'] for metadata in source['files']) != len(source['rows']):
            continue  # Legacy batches cannot be safely split by worksheet offsets.
        offset = 0
        for metadata in source['files']:
            group = source_group(source['name'], metadata['sheet'])
            for row in source['rows'][offset:offset + metadata['rows']]:
                identity = key(dict(row, **{'Source workbook': source['name']}))
                full = group if group is not None else ' '.join(row.get('Product', '').casefold().split()) in ('full title', 'full search')
                members[full].add(identity)
                owners[identity] = full
            offset += metadata['rows']
    if 'view_members' in imported:
        members = {full: set(imported['view_members'].get(label, []))
                   for full, label in ((True, 'Full Title'), (False, 'Remaining Products'))}
        for row in rows:
            identity = key(row)
            if source_groups.get(row.get('Source workbook')) is None and (identity in members[True]) != (identity in members[False]):
                owners[identity] = identity in members[True]
    for row in rows:
        identity = key(row)
        group = source_groups.get(row.get('Source workbook'))
        if group is not None:
            owners[identity] = group
        if identity not in owners:
            full = ' '.join(row.get('Product', '').casefold().split()) in ('full title', 'full search')
            members[full].add(identity)
            owners[identity] = full
    # Keep each uploaded report's membership, including its other products and
    # carried orders. Matching numbers in different workbooks remain independent.
    full = arrival_sort([r for r in rows if owners[key(r)]])
    remaining = arrival_sort([r for r in rows if not owners[key(r)]])
    sheets = {name: {'columns': columns, 'rows': data} for name, data in (
        ('Full Title', full), ('Remaining Products', remaining), ('Overview', rows), ('All Products', rows))}
    reports = sheet_reports({FULL: [r for r in rows if owners[key(r)]],
                             REMAINING: [r for r in rows if not owners[key(r)]]})
    statuses = sorted({str(r.get('Status') or r.get('Task Status') or '(Blank)').strip() for r in rows}, key=str.casefold)
    for periods in reports.values():
        for period in periods:
            counts = Counter(str(r.get('Status') or r.get('Task Status') or '(Blank)').strip() for r in period['rows'])
            period['status_counts'] = {status: counts[status] for status in statuses}
            period.update({'Status: ' + status: counts[status] for status in statuses})
    return {'sheets': sheets, 'reports': reports, 'statuses': statuses,
            'source': 'Imported Excel', 'mode': 'excel', 'offline': False, 'updated_at': imported['created'],
            'revision': hashlib.sha256(json.dumps(imported, sort_keys=True).encode()).hexdigest()}


def current_capacity_month():
    return datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime('%Y-%m')


def capacity_report(reports, settings):
    from capacity_calculations import capacity_days, COUNTS
    daily_capacity = settings.get('capacity', 700)
    extended = settings.get('extended_capacity', 750)
    activity, missing_dates = capacity_days(reports)
    days = [dict(row, Capacity=daily_capacity, **{'Ext capacity': extended}) for row in activity]
    # A Daily Orders selection must not move the live capacity table back to an
    # earlier month. Only an explicitly requested monthly export overrides IST.
    month = settings.get('capacity_month') or current_capacity_month()
    monthly = []
    for key in sorted({day['Date'][:7] for day in days}):
        items = [day for day in days if day['Date'].startswith(key)]
        totals = {column: sum(int(day.get(column, 0) or 0) for day in items) for column in COUNTS}
        monthly.append(dict(Date=key, **totals, Capacity=len(items)*daily_capacity,
                            **{'Ext capacity': len(items)*extended}))
    return {'columns': CAPACITY_COLUMNS, 'monthly': monthly,
            'daily': sorted([r for r in days if r['Date'].startswith(month)], key=lambda r: r['Date']),
            'month': month, 'scope_label': datetime.strptime(month, '%Y-%m').strftime('%B %Y'),
            'capacity': daily_capacity, 'extended_capacity': extended,
            'missing_completion_date_count': missing_dates}


def report_frames(snapshot, settings):
    from tracker_sync import matrix
    from monthly_views import view_name, product_view_headers
    from sla_status import build_sla_status, sheet_summary_values
    reports = snapshot['reports']
    selected = settings.get('selected_date')
    day = next((r for r in reports['daily'] if r['Date'] == selected), None)
    day = day or next((r for r in reports['daily'] if r['Date'] != 'Undated'), None) or next(iter(reports['daily']), None)
    if snapshot.get('mode') == 'excel':
        day = next((r for r in reports['daily'] if r['Date'] != 'Undated'), None) or next(iter(reports['daily']), None)
    month = day['Date'][:7] if day and day['Date'] != 'Undated' else datetime.now().strftime('%Y-%m')
    detail_columns = product_view_headers(snapshot['sheets']['Full Title']['columns'])
    all_rows = snapshot['sheets']['Overview']['rows']
    frames = {'All Products': matrix(all_rows, snapshot['sheets']['Overview']['columns'])}
    for full, label in ((True, 'Full Title'), (False, 'Remaining Products')):
        from status_updates import full_product
        rows = snapshot['sheets'][label]['rows'] if snapshot.get('mode') == 'excel' else [r for r in all_rows if full_product(r) == full]
        frames[view_name(full, month)] = matrix([dict(row, No=index) for index, row in enumerate(rows, 1)], detail_columns)
    status_columns = ['Status: ' + status for status in snapshot.get('statuses', [])]
    monthly_columns = ['Month'] + DAILY_COLUMNS[1:7] + ['SLA OnTime', 'SLA on Missing']
    monthly_rows = [dict(row, **{'SLA on Missing': row.get('Missing', 0)}) for row in reports['monthly']]
    frames['Monthly Orders'] = matrix(monthly_rows, monthly_columns)
    daily_status_columns = DAILY_COLUMNS[:7] + ['SLA On Time', 'SLA Missing']
    daily_status = [dict(row, **{'SLA On Time': row.get('SLA OnTime', 0),
                                'SLA Missing': row.get('Missing', 0)})
                    for row in sorted(reports['daily'], key=lambda row: row['Date'])]
    daily_total = {'Date': 'Total', **{column: sum(row.get(column, 0) for row in daily_status)
                                      for column in daily_status_columns[1:]}}
    frames['Daily Status Report'] = matrix(daily_status + [daily_total], daily_status_columns)
    counts = Counter(str(r.get('Status', '')).strip() or '(Blank)' for r in all_rows)
    frames['Status Report'] = [['Status', 'Orders', 'Share']] + [[s, n, n / len(all_rows)] for s, n in counts.items()]
    frames['SLA Status'] = sheet_summary_values(build_sla_status(snapshot))
    capacity = capacity_report(reports, dict(settings, mode=settings.get('mode', snapshot.get('mode', 'tracker'))))
    monthly = capacity['monthly']
    year = capacity['month'][:4]
    total = {'Date': f'{year} YTD Total', **{c: sum(r.get(c, 0) for r in monthly if str(r['Date']).startswith(year)) for c in capacity['columns'][1:]}}
    daily_total = {'Date': 'Total', **{c: sum(r.get(c, 0) for r in capacity['daily']) for c in capacity['columns'][1:]}}
    frames['PR Excel'] = (matrix(monthly + [total], capacity['columns'])
                                 + matrix(capacity['daily'] + [daily_total], capacity['columns']))
    return frames, capacity


def all_products_date_filter(sheet_id, values, selected_date):
    """Show the selected received day in the shared All Products tab."""
    if sheet_id is None or not values:
        return None
    headers = values[0]
    order = ('Date', 'In-Time') if 'Source workbook' in headers else ('In-Time', 'Date')
    source = next((column for column in order if column in headers), None)
    if source is None:
        return None
    filter_ = {'range': {'sheetId': sheet_id, 'endRowIndex': len(values),
                         'endColumnIndex': len(headers)}}
    if selected_date:
        from openpyxl.utils import get_column_letter

        def date_serial(column):
            cell = f'${get_column_letter(headers.index(column) + 1)}2'
            raw = f'TO_TEXT({cell})'
            # Tracker dates are month/day/year strings. DATEVALUE interprets them
            # using the spreadsheet locale, which can silently select another day.
            iso = f'DATE(VALUE(LEFT({raw},4)),VALUE(MID({raw},6,2)),VALUE(MID({raw},9,2)))'
            month = f'VALUE(REGEXEXTRACT({raw},"^([0-9]{{1,2}})/"))'
            day = f'VALUE(REGEXEXTRACT({raw},"^[0-9]{{1,2}}/([0-9]{{1,2}})/"))'
            year = f'VALUE(REGEXEXTRACT({raw},"^[0-9]{{1,2}}/[0-9]{{1,2}}/([0-9]{{2,4}})"))'
            year = f'IF({year}<100,IF({year}>=69,1900+{year},2000+{year}),{year})'
            american = f'DATE({year},{month},{day})'
            return (f'IFERROR(IF(ISNUMBER({cell}),INT({cell}),'
                    f'IF(REGEXMATCH({raw},"^[0-9]{{4}}-[0-9]{{2}}-[0-9]{{2}}"),{iso},'
                    f'IF(REGEXMATCH({raw},"^[0-9]{{1,2}}/[0-9]{{1,2}}/[0-9]{{2,4}}"),{american},0))),0)')

        received = date_serial(source)
        fallback = next((column for column in order if column != source and column in headers), None)
        if fallback:
            received = f'IF({received}>0,{received},{date_serial(fallback)})'
        if selected_date == 'Undated':
            match = '0'
        else:
            date = datetime.strptime(selected_date, '%Y-%m-%d')
            match = f'DATE({date.year},{date.month},{date.day})'
        filter_['criteria'] = {str(headers.index(source)): {'condition': {'type': 'CUSTOM_FORMULA',
            'values': [{'userEnteredValue': f'={received}={match}'}]}}}
    return {'setBasicFilter': {'filter': filter_}}


def publish_reports(book, snapshot, settings, backup_root, daily_only=False, _collision_retry=0):
    """Back up replaced reports, publish atomically, and keep public report tabs visible.

    Internal tracker/history tabs are hidden because sync and recovery need them.
    The retired Daily Orders output is backed up locally before removal.
    """
    from datatrace_sync import sheet_cell, sheet_color
    from tracker_formatting import format_requests, capacity_sheet_values, REPORT_COLORS
    from monthly_views import view_identity
    frames, capacity = report_frames(snapshot, settings)
    all_products_frame = frames['All Products']
    if daily_only:
        # The tracker sync can refresh All Products between report publications.
        # Publish the rows and filter together so a date never filters stale data.
        frames = {'All Products': all_products_frame,
                  'Daily Status Report': frames['Daily Status Report']}
    metadata = book.fetch_sheet_metadata(params={'fields': 'sheets(properties,conditionalFormats,charts)'})['sheets']
    by_title = {s['properties']['title']: s for s in metadata}
    if daily_only and 'All Products' not in by_title:
        # An initial publication may still be pending when the date is chosen.
        frames, capacity = report_frames(snapshot, settings)
        daily_only = False
    all_products_id = by_title.get('All Products', {}).get('properties', {}).get('sheetId')
    ids = {s['properties']['sheetId'] for s in metadata}
    requests, originals, retired_value_ranges = [], {}, []
    retired = by_title.get('Daily Orders')
    retired_id = retired['properties']['sheetId'] if retired else None
    if retired:
        originals['Daily Orders'] = book.worksheet('Daily Orders').get_all_values(value_render_option='FORMULA')
        requests.append({'deleteSheet': {'sheetId': retired_id}})
    for title, values in frames.items():
        existing = by_title.get(title)
        if not existing:
            existing = next((s for name, s in by_title.items() if name.casefold() == title.casefold()), None)
        if not existing and title == 'PR Excel':
            existing = by_title.get('Capacity Report')
        if not existing and view_identity(title):
            existing = next((s for name, s in by_title.items() if view_identity(name) == view_identity(title)), None)
        if existing:
            props = existing['properties']
            number = props['sheetId']
            originals[props['title']] = book.worksheet(props['title']).get_all_values(value_render_option='FORMULA')
            if props['title'] != title:
                requests.append({'updateSheetProperties': {'properties': {'sheetId': number, 'title': title}, 'fields': 'title'}})
        else:
            number = secrets.randbelow(2**30)
            while number in ids:
                number = secrets.randbelow(2**30)
            ids.add(number)
            props = {}
            requests.append({'addSheet': {'properties': {'sheetId': number, 'title': title}}})
        if title == 'All Products':
            all_products_id = number
        cols = max((len(r) for r in values), default=1)
        old_grid = props.get('gridProperties', {})
        old_cols = max(old_grid.get('columnCount', cols),
                       max((len(row) for row in originals.get(props.get('title'), [])), default=cols))
        requests.append({'updateSheetProperties': {'properties': {'sheetId': number, 'hidden': False,
            'gridProperties': {'rowCount': max(len(values) + 25, old_grid.get('rowCount', 0)),
                               'columnCount': max(cols, 18 if title == 'SLA Status' else 0, old_cols), 'frozenRowCount': 2 if title == 'SLA Status' else 1}},
            'fields': 'hidden,gridProperties'}})
        sheet_values = capacity_sheet_values(values) if title in ('PR Excel', 'Daily Status Report') else values
        if existing:
            # Small status edits should not resend hundreds of thousands of
            # unchanged cells. Explicit clearing below handles retired ranges.
            prior = originals[props['title']]
            def comparable(value):
                if value is None:
                    return ''
                return str(float(value)) if isinstance(value, (int, float)) and not isinstance(value, bool) else str(value)
            for row_index, row in enumerate(sheet_values):
                saved = prior[row_index] if row_index < len(prior) else []
                changed = [column for column, value in enumerate(row)
                           if comparable(value) != comparable(saved[column] if column < len(saved) else '')]
                if changed:
                    left, right = min(changed), max(changed)+1
                    requests.append({'updateCells': {'start': {'sheetId': number, 'rowIndex': row_index, 'columnIndex': left},
                        'rows': [{'values': [sheet_cell(value) for value in row[left:right]]}],
                        'fields': 'userEnteredValue'}})
        else:
            output_rows = [{'values': [sheet_cell(value) for value in row]} for row in sheet_values]
            requests.append({'updateCells': {'range': {'sheetId': number},
                'rows': output_rows, 'fields': 'userEnteredValue'}})
        if title == 'SLA Status':
            from sla_status import sheet_format_requests
            requests.extend(sheet_format_requests(number, values, existing or {}))
        else:
            requests.extend(format_requests(number, values, (existing or {}).get('conditionalFormats', []),
                                            replace_rules=title == 'Status Report', grid_shape=(
                                                max(len(values) + 25, old_grid.get('rowCount', 0)),
                                                max(cols, old_grid.get('columnCount', 0)))))
        # Old report fills can extend far beyond the current data. Clear those
        # blank rows as well; status rules are bounded to populated rows.
        requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': len(values),
            'endColumnIndex': cols}, 'cell': {}, 'fields': 'userEnteredFormat'}})
        requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': len(values),
            'endRowIndex': max(len(values)+25, old_grid.get('rowCount', 0)),
            'endColumnIndex': cols}, 'cell': {}, 'fields': 'userEnteredValue'}})
        from gspread.utils import rowcol_to_a1
        quoted = "'" + title.replace("'", "''") + "'!"
        # Empty string values still contribute to Google Sheets' selected-cell
        # Count. Remove retired cells rather than filling thousands of blanks.
        retired_value_ranges.append(quoted + rowcol_to_a1(len(values)+1, 1) + ':' +
                                    rowcol_to_a1(max(len(values)+25, old_grid.get('rowCount', 0)), old_cols))
        if old_cols > cols:
            from gspread.utils import rowcol_to_a1
            quoted = "'" + title.replace("'", "''") + "'!"
            retired_value_ranges.append(quoted + rowcol_to_a1(1, cols+1) + ':' +
                                        rowcol_to_a1(max(len(values), old_grid.get('rowCount', len(values))), old_cols))
            requests.append({'repeatCell': {'range': {'sheetId': number,
                'endRowIndex': max(len(values), old_grid.get('rowCount', len(values))),
                'startColumnIndex': cols, 'endColumnIndex': old_cols},
                'cell': {}, 'fields': 'userEnteredValue'}})
            requests.append({'repeatCell': {'range': {'sheetId': number,
                'endRowIndex': max(len(values), old_grid.get('rowCount', len(values))),
                'startColumnIndex': cols, 'endColumnIndex': old_cols},
                'cell': {}, 'fields': 'userEnteredFormat'}})
        if title == 'Status Report' and len(values) > 1:
            requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': 1,
                'endRowIndex': len(values), 'startColumnIndex': 2, 'endColumnIndex': 3},
                'cell': {'userEnteredFormat': {'numberFormat': {'type': 'PERCENT', 'pattern': '0.00%'}}},
                'fields': 'userEnteredFormat.numberFormat'}})
        if title == 'SLA Status':
            from sla_status import sheet_chart_requests
            requests.extend(sheet_chart_requests(number, values, existing or {}))
        if title == 'Daily Status Report':
            if len(values) > 2:
                requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': 1,
                    'endRowIndex': len(values)-1, 'endColumnIndex': cols},
                    'cell': {'userEnteredFormat': {'backgroundColor': sheet_color('#ffffff'),
                        'textFormat': {'bold': False, 'foregroundColor': sheet_color('#111827')}}},
                    'fields': 'userEnteredFormat.backgroundColor,userEnteredFormat.textFormat.bold,userEnteredFormat.textFormat.foregroundColor'}})
                requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': 1,
                    'endRowIndex': len(values)-1, 'startColumnIndex': 0, 'endColumnIndex': 1},
                    'cell': {'userEnteredFormat': {'numberFormat': {'type': 'DATE', 'pattern': 'dd-MM-yy'},
                        'backgroundColor': sheet_color('#' + REPORT_COLORS['daily'])}},
                    'fields': 'userEnteredFormat.numberFormat,userEnteredFormat.backgroundColor'}})
            for row_index, row in enumerate(values[1:-1], 1):
                if str(row[0]) == settings.get('selected_date'):
                    requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': row_index,
                        'endRowIndex': row_index+1, 'endColumnIndex': cols},
                        'cell': {'userEnteredFormat': {'backgroundColor': sheet_color('#cfe2f3'),
                            'textFormat': {'bold': True, 'foregroundColor': sheet_color('#111827')}}},
                        'fields': 'userEnteredFormat.backgroundColor,userEnteredFormat.textFormat.bold,userEnteredFormat.textFormat.foregroundColor'}})
            for row_index, color in ((0, '#' + REPORT_COLORS['header']),
                                     (len(values)-1, '#' + REPORT_COLORS['total'])):
                requests.append({'repeatCell': {'range': {'sheetId': number, 'startRowIndex': row_index,
                    'endRowIndex': row_index+1, 'endColumnIndex': cols},
                    'cell': {'userEnteredFormat': {'backgroundColor': sheet_color(color), 'textFormat': {'bold': True}}},
                    'fields': 'userEnteredFormat.backgroundColor,userEnteredFormat.textFormat.bold'}})
        if title == 'PR Excel':
            last_month_row = len(capacity['monthly']) + 1
            daily_start = last_month_row + 1
            daily_end = daily_start + len(capacity['daily']) + 1
            def source(column):
                return {'sourceRange': {'sources': [
                    {'sheetId': number, 'startRowIndex': daily_start, 'endRowIndex': max(daily_start + 2, daily_end),
                     'startColumnIndex': column, 'endColumnIndex': column + 1}]}}
            spec = {'title': 'Overall — ' + capacity['scope_label'], 'basicChart': {'chartType': 'COMBO', 'legendPosition': 'BOTTOM_LEGEND', 'headerCount': 1,
                'axis': [{'position': 'BOTTOM_AXIS', 'title': 'Date'}, {'position': 'LEFT_AXIS', 'title': 'Orders'}],
                'domains': [{'domain': source(0)}], 'series': [
                    {'series': source(c), 'type': kind, 'targetAxis': 'LEFT_AXIS'} for c, kind in ((1, 'COLUMN'), (2, 'COLUMN'), (7, 'LINE'), (8, 'LINE'))]}}
            chart = next((c for c in (existing or {}).get('charts', []) if c.get('spec', {}).get('title', '').startswith('Overall')), None)
            if chart:
                requests.append({'updateChartSpec': {'chartId': chart['chartId'], 'spec': spec}})
                requests.append({'updateEmbeddedObjectPosition': {'objectId': chart['chartId'], 'newPosition': {'overlayPosition': {'anchorCell': {'sheetId': number, 'rowIndex': len(values)+2, 'columnIndex': 0}, 'widthPixels': 900, 'heightPixels': 350}}, 'fields': 'anchorCell,widthPixels,heightPixels'}})
            else:
                requests.append({'addChart': {'chart': {'spec': spec, 'position': {'overlayPosition': {'anchorCell': {'sheetId': number, 'rowIndex': len(values)+2, 'columnIndex': 0}, 'widthPixels': 900, 'heightPixels': 350}}}}})
    date_filter = all_products_date_filter(all_products_id, all_products_frame, settings.get('selected_date', ''))
    if date_filter:
        requests.append(date_filter)
    if not daily_only:
        shown_ids = {req['updateSheetProperties']['properties']['sheetId'] for req in requests if req.get('updateSheetProperties', {}).get('properties', {}).get('hidden') is False}
        for item in metadata:
            props = item['properties']
            if props['sheetId'] != retired_id and props['sheetId'] not in shown_ids and not props.get('hidden'):
                requests.append({'updateSheetProperties': {'properties': {'sheetId': props['sheetId'], 'hidden': True}, 'fields': 'hidden'}})
    backup_root = Path(backup_root)
    backup_root.mkdir(parents=True, exist_ok=True)
    backup = backup_root / f'reports-{datetime.now().strftime("%Y%m%d-%H%M%S-%f")}.json.gz'
    with gzip.open(backup, 'wt', encoding='utf-8') as stream:
        json.dump({'metadata': metadata, 'values': originals}, stream)
    try:
        book.batch_update({'requests': requests})
        if retired_value_ranges and hasattr(book, 'values_batch_clear'):
            # The values clear endpoint also removes legacy populated cells
            # beyond the report schema that a formatting-only reset retains.
            book.values_batch_clear(body={'ranges': retired_value_ranges})
    except Exception as exc:
        # Another installation may have created a tab after metadata was read.
        # Sheets batch updates are atomic, so rebuild once against fresh tabs.
        from gspread.exceptions import APIError
        message = str(exc).casefold()
        if isinstance(exc, APIError) and 'addsheet' in message and 'already exists' in message and not _collision_retry:
            return publish_reports(book, snapshot, settings, backup_root, daily_only, _collision_retry=1)
        raise
    # A failed/lost reply leaves the source saved and retryable. Each retry rereads
    # metadata, so no chart, tab or appended data can be duplicated.
    for title, wanted in frames.items():
        actual = book.worksheet(title).get_all_values(value_render_option='UNFORMATTED_VALUE')
        # Verify the report's declared columns. Unnamed cells beyond the
        # schema are not order fields and cannot invalidate matching orders.
        width = max((len(row) for row in wanted), default=0)
        actual = [row[:width] for row in actual]
        if title in ('PR Excel', 'Daily Status Report'):
            wanted = capacity_sheet_values(wanted)
        def normalized(matrix):
            result = []
            for row in matrix:
                row = [str(float(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v) for v in row]
                while row and not row[-1]:
                    row.pop()
                result.append(row)
            while result and not result[-1]:
                result.pop()
            return result
        if normalized(actual) != normalized(wanted):
            raise RuntimeError(f'{title} could not be verified. Retry publishing the saved report.')
    highlighted = settings.get('selected_date', '')
    highlighted = highlighted if any(str(row[0]) == highlighted for row in frames['Daily Status Report'][1:-1]) else ''
    return {'tabs': list(frames), 'backup': str(backup), 'highlighted_date': highlighted}


def register_report_routes(app, workspace, gate, load_snapshot, get_book, invalidate):
    from flask import request, jsonify, send_file
    from tracker_sync import sync_lock
    from import_excel_changes import sync_imported_tabs, editable_frames, change_history, preserve_sheet_edits, tabs_unchanged
    import time

    requested = threading.Event()
    full_requested = threading.Event()
    worker_lock = threading.Lock()
    worker = None
    stopped = app.extensions.get('stop_scheduler', threading.Event())

    def publish(daily_only=False):
        revision = workspace.read().get('data_revision')
        try:
            with sync_lock(workspace.store.root):
                settings = workspace.read()
                revision = settings.get('data_revision')
                with workspace.lock:
                    if workspace.read().get('data_revision') == revision:
                        workspace.save(publication_state='working', sync_error=None)
                book, _ = get_book()
                if sync_imported_tabs(workspace, book, settings):
                    daily_only = False
                settings = workspace.read()
                revision = settings.get('data_revision')
                snapshot = (imported_snapshot(settings['imported'])
                            if settings['mode'] == 'excel' and settings.get('imported')
                            else load_snapshot(force=True, tracker_only=True))
                if snapshot.get('offline'):
                    raise RuntimeError('Reconnect Google Sheets before publishing tracker reports.')
                if workspace.read().get('data_revision') != revision:
                    requested.set()
                    return {'synced': False, 'queued': True}
                if (settings.get('mode') == 'excel' and not settings.get('cloud_reset')
                        and settings.get('cloud_baseline') and not daily_only
                        and not tabs_unchanged(book, settings['cloud_baseline'])):
                    requested.set()
                    full_requested.set()
                    return {'synced': False, 'queued': True}
                result = publish_reports(book, snapshot, settings, workspace.store.root.parent / 'backups' / 'reports', daily_only)
            with workspace.lock:
                if workspace.read().get('data_revision') != revision:
                    requested.set()
                    return dict(result, synced=False, queued=True)
                workspace.save(sync_error=None, publication_state='synced', highlighted_date=result['highlighted_date'],
                               synced_at=datetime.now(timezone.utc).isoformat(), cloud_refresh_error=None,
                               **({'cloud_baseline': editable_frames(snapshot, settings), 'cloud_reset': False}
                                  if settings.get('mode') == 'excel' and not daily_only else {}))
            return dict(result, synced=True)
        except Exception as exc:
            app.logger.exception('Report publication failed; saved source retained')
            with workspace.lock:
                if workspace.read().get('data_revision') == revision:
                    workspace.save(sync_error=str(exc), publication_state='error')
                else:
                    requested.set()
                    return {'synced': False, 'queued': True}
            return {'synced': False, 'sync_error': str(exc)}

    def publication_worker():
        while not stopped.is_set():
            if not requested.wait(1):
                continue
            if not gate.acquire(timeout=1):
                continue
            try:
                with workspace.lock:
                    requested.clear()
                    full = full_requested.is_set()
                    full_requested.clear()
                publish(daily_only=not full)
            finally:
                gate.release()

    def queue_publication(daily_only=False):
        nonlocal worker
        with workspace.lock:
            workspace.save(publication_state='pending', sync_error=None, enabled=True)
            if not daily_only:
                full_requested.set()
            requested.set()
        with worker_lock:
            if worker is None or not worker.is_alive():
                worker = threading.Thread(target=publication_worker, daemon=True, name='report-publication')
                worker.start()
        return {'synced': False, 'queued': True}

    app.extensions['queue_report_publication'] = queue_publication

    last_cloud_check = float('-inf')
    def refresh_imported_reports(force=False):
        nonlocal last_cloud_check
        settings = workspace.read()
        if settings.get('mode') == 'excel' and not settings.get('imported'):
            return {'changed': False, 'mode': 'excel'}
        reset = settings.get('mode') == 'excel' and settings.get('cloud_reset')
        if reset or (not force and time.monotonic() - last_cloud_check < 30):
            return {'changed': False, 'busy': bool(reset)}
        if not gate.acquire(blocking=False):
            return {'changed': False, 'busy': True}
        try:
            last_cloud_check = time.monotonic()
            with sync_lock(workspace.store.root):
                book, _ = get_book()
                if settings.get('mode') == 'excel':
                    changed = sync_imported_tabs(workspace, book, workspace.read())
                else:
                    from status_updates import sync_tracker_status_tabs
                    changed = sync_tracker_status_tabs(workspace, book)
            if changed:
                invalidate()
            publication = queue_publication() if changed else None
            return {'changed': changed, 'publication': publication, **workspace.public()}
        except Exception as exc:
            workspace.save(cloud_refresh_error=str(exc))
            last_cloud_check = time.monotonic() + 30
            if force:
                raise
            return {'changed': False, 'sync_error': str(exc)}
        finally:
            gate.release()

    app.extensions['refresh_imported_reports'] = refresh_imported_reports

    @app.post('/api/report-refresh')
    def refresh_imported_source():
        result = refresh_imported_reports(force=True)
        if result.get('busy'):
            return jsonify(error='Wait for the current import, capture or Sheets update to finish, then refresh.'), 409
        return jsonify(result)

    @app.get('/api/import-excel-changes')
    def imported_change_history():
        offset = max(0, request.args.get('offset', default=0, type=int) or 0)
        result = change_history(workspace, request.args.get('action', ''), request.args.get('search', '')[:200], offset)
        settings = workspace.read()
        return jsonify(**result, mode=settings['mode'], checked_at=settings.get('cloud_checked_at'),
                       sync_error=settings.get('cloud_refresh_error'), publication_state=settings.get('publication_state'))

    @app.get('/api/report-workspace')
    def get_report_workspace():
        return jsonify(workspace.public())

    @app.post('/api/report-source')
    def choose_report_source():
        mode = (request.get_json(silent=True) or {}).get('mode')
        if mode not in ('tracker', 'excel'):
            raise ValueError('Choose Tracker report or Import Excel report.')
        with workspace.lock:
            if mode == 'excel' and not workspace.read().get('imported'):
                raise ValueError('Import Excel files before selecting the Excel report.')
            workspace.save(mode=mode, enabled=True, selected_date='', cloud_baseline=None,
                           cloud_reset=mode == 'excel', cloud_refresh_error=None)
        invalidate()
        return jsonify(**workspace.public(), publication=queue_publication())

    @app.post('/api/report-import')
    def import_reports():
        files = request.files.getlist('files')
        # Parse before any mutation. Local report edits do not wait for cloud sync.
        uploaded = read_report_files([(file.filename, file.read()) for file in files if file.filename])
        with workspace.lock:
            previous_imported = workspace.read().get('imported')
            replace_all = request.form.get('replace_all') == 'true'
            sources = [] if replace_all else import_sources(previous_imported)
            uploaded_names = {s['name'].casefold() for s in uploaded['sources']}
            sources = [s for s in sources if not (s.get('legacy') and
                       {f['file'].casefold() for f in s['files']}.issubset(uploaded_names))]
            replaced = []
            for source in uploaded['sources']:
                name = source['name'].casefold()
                replaced.extend(s['name'] for s in sources if s['name'].casefold() == name)
                sources = [s for s in sources if s['name'].casefold() != name]
                sources.append(source)
            imported = combine_sources(sources)
            imported = preserve_sheet_edits(previous_imported, imported, uploaded_names, replace_all)
            workspace.save(imported=imported, files=imported['files'], row_count=len(imported['rows']),
                           imported_at=imported['created'], mode='excel', enabled=True, selected_date='',
                           cloud_baseline=None, cloud_reset=True, cloud_refresh_error=None)
        invalidate()
        publication = queue_publication()
        return jsonify(**workspace.public(), reviews=imported['reviews'], replaced_files=replaced, publication=publication), 201

    @app.post('/api/report-remove')
    def remove_imported_file():
        body = request.get_json(silent=True) or {}
        with workspace.lock:
            previous_imported = workspace.read().get('imported')
            sources = import_sources(previous_imported)
            name = body.get('file')
            if body.get('all') is True:
                sources = []
            elif isinstance(name, str) and any(s['name'] == name for s in sources):
                sources = [s for s in sources if s['name'] != name]
            else:
                raise ValueError('Choose an imported file to remove.')
            imported = combine_sources(sources)
            imported = preserve_sheet_edits(previous_imported, imported, [name] if isinstance(name, str) else [], body.get('all') is True)
            # An empty Excel report stays empty until the user selects Tracker.
            workspace.save(imported=imported, files=imported['files'], row_count=len(imported['rows']),
                           imported_at=imported['created'], selected_date='', cloud_baseline=None,
                           cloud_reset=True, cloud_refresh_error=None)
        invalidate()
        return jsonify(**workspace.public(), publication=queue_publication())

    @app.post('/api/report-publish')
    def publish_current_reports():
        return jsonify(**queue_publication())

    @app.post('/api/report-date')
    def choose_report_date():
        date = (request.get_json(silent=True) or {}).get('date', '')
        if not isinstance(date, str) or (date and not re.fullmatch(r'\d{4}-\d{2}-\d{2}|Undated', date)):
            raise ValueError('Choose a report date.')
        if date and not any(row['Date'] == date for row in load_snapshot()['reports']['daily']):
            raise ValueError('That date is no longer in the selected report source.')
        workspace.save(selected_date=date, enabled=True)
        return jsonify(date=date, **queue_publication(daily_only=True))

    @app.get('/api/capacity-report')
    def get_capacity_report():
        snapshot = load_snapshot()
        return jsonify(**capacity_report(snapshot['reports'], workspace.read()),
                       source=snapshot.get('source'), offline=snapshot.get('offline', False))

    @app.get('/api/capacity-report/details')
    def get_capacity_details():
        from capacity_calculations import capacity_details
        from order_details import recorded_changes
        date = request.args.get('date', '')
        month = request.args.get('month', '')
        if month:
            try:
                if not re.fullmatch(r'\d{4}-\d{2}', month):
                    raise ValueError()
                datetime.strptime(month, '%Y-%m')
            except ValueError:
                return jsonify(error='Choose a valid capacity report month.'), 400
        if (not date and not month) or (date and not re.fullmatch(r'\d{4}-\d{2}-\d{2}', date)):
            return jsonify(error='Choose a capacity report date.'), 400
        if date and month and not date.startswith(month + '-'):
            return jsonify(error='Choose a date within the selected month.'), 400
        snapshot = load_snapshot()
        settings = workspace.read()
        mode = settings['mode']
        month_report = capacity_report(snapshot['reports'], dict(settings, capacity_month=month)) if month else None
        if month_report and not date:
            date = next((row['Date'] for row in month_report['daily']), '')
        month_fields = {'month': month, 'month_label': month_report['scope_label'],
                        'daily': month_report['daily'], 'capacity_columns': month_report['columns']} if month_report else {}
        if not date:
            return jsonify(date='', columns=[], groups=[], raw_received=0, rollover_received=0,
                           **month_fields, mode=mode, source=snapshot.get('source'), offline=snapshot.get('offline', False))
        changes = recorded_changes(workspace, mode, date)
        try:
            details = capacity_details(snapshot['reports'], date)
        except ValueError as exc:
            if not changes:
                return jsonify(error=str(exc)), 404
            details = {'date': date, 'previous_date': None, 'next_date': None, 'raw_received': 0,
                       'rollover_received': 0, 'columns': [], 'groups': [{'key': 'received',
                       'label': 'Received workload', 'rows': [],
                       'description': 'No orders remain assigned to this reporting date. Its recorded changes are retained.'}]}
        details['groups'].append({'key': 'changes', 'label': 'Recorded changes',
            'description': 'Edits saved from this report date or detected on this date in India time. History starts when changes were first recorded.',
            'rows': changes})
        return jsonify(**details, **month_fields, mode=mode, source=snapshot.get('source'), offline=snapshot.get('offline', False))

    @app.get('/api/daily-orders/details')
    def get_daily_details():
        from order_details import recorded_changes
        date = request.args.get('date', '')
        snapshot = load_snapshot()
        mode = workspace.read()['mode']
        changes = recorded_changes(workspace, mode, date)
        day = next((row for row in snapshot['reports']['daily'] if row['Date'] == date), None)
        if day is None:
            if not changes:
                return jsonify(error='That date is no longer in Daily Orders. Refresh the report and try again.'), 404
            day = {'rows': [], 'columns': snapshot['sheets']['Overview']['columns']}
        groups = [{'key': 'received', 'label': 'Received', 'rows': day['rows']},
                  {'key': 'changes', 'label': 'Recorded changes', 'rows': changes,
                   'description': 'Edits saved from this report date or detected on this date in India time. History starts when changes were first recorded.'}]
        return jsonify(date=date, columns=day['columns'], groups=groups, mode=mode,
                       source=snapshot.get('source'), offline=snapshot.get('offline', False))

    @app.get('/api/order-details')
    def get_order_details():
        from order_details import editor_details
        settings = workspace.read()
        if request.args.get('mode') != settings['mode']:
            return jsonify(error='The report source changed. Reopen the date before editing.'), 409
        try:
            return jsonify(editor_details(load_snapshot(), settings,
                           request.args.get('order', ''), request.args.get('month', ''), request.args.get('source_workbook', '')))
        except ValueError as exc:
            return jsonify(error=str(exc)), 409

    @app.post('/api/order-details')
    def save_order_details():
        from order_details import save_order
        body = request.get_json(silent=True)
        if (not isinstance(body, dict) or not isinstance(body.get('order_number'), str)
                or not body['order_number'].strip() or not isinstance(body.get('month', ''), str)
                or not isinstance(body.get('source_workbook', ''), str)):
            return jsonify(error='Choose an order from the report date table.'), 400
        if not gate.acquire(blocking=False):
            return jsonify(error='Wait for the current capture or Sheets update to finish, then save.'), 409
        try:
            with sync_lock(workspace.store.root), workspace.lock:
                settings = workspace.read()
                if body.get('mode') != settings['mode']:
                    return jsonify(error='The report source changed. Reopen the order before saving.'), 409
                book = None
                if settings['mode'] == 'excel':
                    if settings.get('cloud_baseline') and not settings.get('cloud_reset'):
                        book, _ = get_book()
                        sync_imported_tabs(workspace, book, settings)
                    snapshot = imported_snapshot(workspace.read()['imported'])
                else:
                    from status_updates import sync_tracker_status_tabs
                    book, _ = get_book()
                    sync_tracker_status_tabs(workspace, book)
                    snapshot = load_snapshot(force=True, tracker_only=True)
                result = save_order(workspace, book, snapshot, body)
            invalidate()
            return jsonify(**result, publication=queue_publication())
        except ValueError as exc:
            return jsonify(error=str(exc)), 409
        finally:
            gate.release()

    @app.post('/api/report-capacity')
    def save_capacity():
        body = request.get_json(silent=True) or {}
        for field in ('capacity', 'extended_capacity'):
            value = body.get(field)
            if isinstance(value, bool) or not isinstance(value, int) or not 0 <= value <= 1000000:
                raise ValueError('Capacity must be a whole number from 0 to 1,000,000.')
        workspace.save(capacity=body['capacity'], extended_capacity=body['extended_capacity'], enabled=True)
        return jsonify(**workspace.public(), publication=queue_publication())

    if workspace.read().get('publication_state') in ('pending', 'working'):
        queue_publication()

    @app.get('/api/sla-status')
    def get_sla_status():
        from sla_status import build_sla_status
        report = build_sla_status(load_snapshot())
        return jsonify(**report, sync_error=workspace.read().get('cloud_refresh_error'))

    @app.post('/api/order-status')
    def edit_order_status():
        from status_updates import status_name, rename_status, sync_tracker_status_tabs
        body = request.get_json(silent=True) or {}
        old, new = status_name(body.get('status')), status_name(body.get('replacement'))
        selected = body.get('order_numbers')
        if (not isinstance(body.get('status'), str) or not isinstance(body.get('replacement'), str)
                or not old or not new or len(new) > 1000 or old == new or not isinstance(selected, list)
                or not selected or any(not isinstance(n, str) or not n.strip() for n in selected)):
            raise ValueError('Choose a different, nonblank status and the orders to update.')
        if not gate.acquire(blocking=False):
            return jsonify(error='Wait for the current capture or Sheets update to finish, then save.'), 409
        try:
            with sync_lock(workspace.store.root):
                settings = workspace.read()
                book = None
                if settings.get('mode') == 'excel':
                    book, _ = get_book()
                    sync_imported_tabs(workspace, book, settings)
                    settings = workspace.read()
                    snapshot = imported_snapshot(settings['imported'])
                else:
                    book, _ = get_book()
                    sync_tracker_status_tabs(workspace, book)
                    snapshot = load_snapshot(force=True, tracker_only=True)
                count = rename_status(workspace, book, snapshot, old, new,
                                      [number.strip().casefold() for number in selected])
            invalidate()
            publication = queue_publication()
            return jsonify(updated=count, status=new, publication=publication)
        except ValueError as exc:
            return jsonify(error=str(exc)), 409
        finally:
            gate.release()

    @app.get('/api/export/report')
    def export_report():
        from openpyxl import Workbook
        from openpyxl.styles import Border, Font, PatternFill, Side
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.chart.series import SeriesLabel
        from tracker_formatting import REPORT_COLORS
        snapshot = load_snapshot()
        if request.args.get('sheet') == 'SLA Status':
            from sla_status import build_sla_status, excel_report
            report = build_sla_status(snapshot)
            return send_file(excel_report(report), as_attachment=True,
                download_name=f'TV_Status_{report["as_of"]}.xlsx',
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        settings = workspace.read()
        month = request.args.get('month')
        if month:
            if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
                raise ValueError('Choose a month in YYYY-MM format.')
            selected = next((r for r in snapshot['reports']['daily'] if r['Date'].startswith(month)), None)
            if not selected:
                raise ValueError('No orders are available for that month.')
            settings = dict(settings, selected_date=selected['Date'], capacity_month=month)
        frames, capacity = report_frames(snapshot, settings)
        requested_sheet = request.args.get('sheet')
        aliases = {'Daily Orders': 'Daily Status Report', 'Capacity Report': 'PR Excel'}
        requested_sheet = aliases.get(requested_sheet, requested_sheet)
        if requested_sheet:
            if requested_sheet not in frames:
                raise ValueError('Choose a report tab to download.')
            frames = {requested_sheet: frames[requested_sheet]}
        book = Workbook()
        book.remove(book.active)
        fills = {color: PatternFill('solid', fgColor=color) for color in
                 ('F6B26B', 'FFFFFF', *REPORT_COLORS.values())}
        border = Border(**{side: Side(style='thin', color='000000')
                           for side in ('left', 'right', 'top', 'bottom')})
        for title, values in frames.items():
            sheet = book.create_sheet(title[:31])
            for row in values:
                sheet.append(row)
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            if title == 'PR Excel':
                from datetime import date
                for row_index in range(2, len(capacity['monthly']) + 2):
                    raw = sheet.cell(row_index, 1).value
                    match = re.fullmatch(r'(\d{4})-(\d{2})', str(raw))
                    if match:
                        sheet.cell(row_index, 1).value = date(int(match[1]), int(match[2]), 1)
                        sheet.cell(row_index, 1).number_format = 'mmm-yy'
                daily_start = len(capacity['monthly']) + 4
                for row_index in range(daily_start, daily_start + len(capacity['daily'])):
                    raw = sheet.cell(row_index, 1).value
                    try:
                        sheet.cell(row_index, 1).value = date.fromisoformat(str(raw))
                    except ValueError:
                        pass
                    else:
                        sheet.cell(row_index, 1).number_format = 'dd-mm-yy'
            for row in sheet.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.font = Font(name='Calibri', size=11, bold=cell.row == 1)
                    cell.fill = fills['FFFFFF']
            if title == 'PR Excel':
                month_count = len(capacity['monthly'])
                daily_start = month_count + 4
                daily_total = daily_start + len(capacity['daily'])
                for cell in sheet[1]:
                    cell.fill = fills[REPORT_COLORS['header']]
                for row_index in range(2, month_count + 2):
                    for cell in sheet[row_index]:
                        cell.fill = fills[REPORT_COLORS['monthly']]
                for cell in sheet[month_count + 2]:
                    cell.fill = fills[REPORT_COLORS['ytd']]
                    cell.font = Font(name='Calibri', size=11, bold=True)
                for cell in sheet[month_count + 3]:
                    cell.fill = fills[REPORT_COLORS['header']]
                    cell.font = Font(name='Calibri', size=11, bold=True)
                for row_index in range(daily_start, daily_total):
                    for cell in sheet[row_index]:
                        if cell.column == 1 or cell.column in (8, 9):
                            cell.fill = fills[REPORT_COLORS['daily']]
                for cell in sheet[daily_total]:
                    cell.fill = fills[REPORT_COLORS['total']]
                    cell.font = Font(name='Calibri', size=11, bold=True)
                sheet.column_dimensions['A'].width = 14
                for column in range(2, len(CAPACITY_COLUMNS) + 1):
                    sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 18
            else:
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                    cell.fill = fills['F6B26B']
                if title == 'Daily Status Report':
                    from datetime import date
                    for cell in sheet[1]:
                        cell.fill = fills[REPORT_COLORS['header']]
                    for row_index in range(2, sheet.max_row):
                        raw = sheet.cell(row_index, 1).value
                        try:
                            sheet.cell(row_index, 1).value = date.fromisoformat(str(raw))
                        except ValueError:
                            pass
                        else:
                            sheet.cell(row_index, 1).number_format = 'dd-mm-yy'
                            sheet.cell(row_index, 1).fill = fills[REPORT_COLORS['daily']]
                    for cell in sheet[sheet.max_row]:
                        cell.fill = fills[REPORT_COLORS['total']]
                        cell.font = Font(name='Calibri', size=11, bold=True)
                    sheet.column_dimensions['A'].width = 14
            sheet.freeze_panes = 'A2'
        if 'PR Excel' not in book.sheetnames:
            stream = BytesIO()
            book.save(stream)
            stream.seek(0)
            filename = requested_sheet.replace(' ', '_') + '.xlsx' if requested_sheet else 'Production_data.xlsx'
            return send_file(stream, as_attachment=True, download_name=filename,
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        sheet = book['PR Excel']
        start = len(capacity['monthly']) + 4
        end = start + len(capacity['daily']) - 1
        if end >= start:
            chart, lines = BarChart(), LineChart()
            chart.title = 'Overall — ' + capacity['scope_label']
            for column, target in ((2, chart), (3, chart), (8, lines), (9, lines)):
                target.add_data(Reference(sheet, min_col=column, min_row=start, max_row=end))
                target.series[-1].title = SeriesLabel(v=CAPACITY_COLUMNS[column-1])
            chart.set_categories(Reference(sheet, min_col=1, min_row=start, max_row=end))
            chart += lines
            sheet.add_chart(chart, f'A{len(frames["PR Excel"])+3}')
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        filename = requested_sheet.replace(' ', '_') + '.xlsx' if requested_sheet else 'Production_data.xlsx'
        return send_file(stream, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    return publish

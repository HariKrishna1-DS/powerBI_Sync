"""Business queue rules, order-level comparison and selected-order corrections."""
from copy import deepcopy
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
from preview_store import PreviewStore
from report_workspace import ReportWorkspace, register_report_routes
from sla_status import preview_snapshot, build_sla_status, excel_report, queue_rule, special_product_status
from test_pdf_report_workspace import ReportBook


class SlaQueueRulesTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.store = PreviewStore(Path(folder.name) / 'previews')
        self.workspace = ReportWorkspace(self.store)
        self.book = ReportBook()
        self.app = Flask(__name__)
        register_report_routes(self.app, self.workspace, threading.Lock(), lambda: None,
                               lambda: (self.book, None), lambda: None)
        self.client = self.app.test_client()

    def capture(self, number, rows, created='2026-10-19T04:30:00Z'):
        columns = list(dict.fromkeys(['Order Number', 'SLA Expiration*'] +
                                    [column for row in rows for column in row]))
        self.store.restore(number, columns, rows, created=created)

    def test_rule_conditions_and_priority(self):
        cases = [({'ETA Comments': 'ETA: on\nHOLD - Waiting for plant date', 'Status': 'Assign to ABS'}, 'Waiting for Effective Date'),
                 ({'Comment': 'ON HOLD', 'Status': 'QC in Progress'}, 'QC in Progress'),
                 ({'Product': 'ClientReview', 'Last User': '', 'Task Status': 'Available'}, 'SearchFix'),
                 ({'Product': 'QUALITY CONTROL'}, 'SearchFix'),
                 ({'Product': 'Quailty Control'}, 'SearchFix'),
                 ({'TaskName': 'crsp2', 'LastUser': '', 'TaskStatus': 'Available', 'ETA Comments': 'ON HOLD'}, 'SearchFix'),
                 ({'Task Name': 'SearchFix', 'Last User': '', 'Task Status': 'Available'}, 'SearchFix'),
                 ({'TaskName': ' search fix ', 'TaskStatus': 'Workflow Suspended'}, 'SearchFix'),
                 ({'Product': 'Attorney Approval'}, 'Attorney Approval'),
                 ({'Product': 'Home Builders'}, 'Home Builder'),
                 ({'Product': 'Home Builder'}, 'Home Builder'),
                 ({'LastUser': ' ', 'TaskStatus': 'Available', 'Status': 'Assign to ABS'}, 'Search In Progress'),
                 ({'Last User': 'Named user', 'Task Status': 'Available', 'Status': 'Assign to ABS'}, 'Assign to ABS'),
                 ({'WorkflowSuspended': True, 'Task Status': 'Available'}, 'Awaiting for Clarification'),
                 ({'WorkflowSuspended': True, 'ETA Comments': 'ON HOLD'}, 'Waiting for Effective Date')]
        originals = [dict(fields, **{'Order Number': str(index), 'SLA Expiration*': '2h'})
                     for index, (fields, _) in enumerate(cases)]
        self.capture(1, originals)
        snapshot = preview_snapshot(self.store)
        self.assertEqual([row['Status'] for row in snapshot['sheets']['Overview']['rows']],
                         [status for _, status in cases])
        self.assertEqual(self.store.get(1)['rows'], originals)
        self.assertEqual(queue_rule({'ETA Comments': 'behold or on holder'})[1], '')

    def test_crsp2_and_searchfix_tasks_share_one_group_and_preserve_order_numbers(self):
        originals = [
            {'Order Number': '001', 'Task Name': 'CRSP2', 'Task Status': 'Task Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': '002', 'Task Name': 'SearchFix', 'Task Status': 'Available', 'SLA Expiration*': '1d'},
            {'Order Number': '003', 'TaskName': 'SEARCH FIX', 'TaskStatus': 'Workflow Suspended', 'SLA Expiration*': '-2h'},
            {'Order Number': '004', 'Task Name': 'Search', 'Task Status': 'Task Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': ' 004 ', 'Task Name': 'SearchFix', 'Task Status': 'Workflow Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': '005', 'Task Name': 'SearchFix', 'ETA Comments': 'ON HOLD', 'SLA Expiration*': '1d'},
        ]
        self.capture(1, originals)
        before = deepcopy(self.store.history())
        report = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(len(report['summary']), 1)
        self.assertEqual(report['summary'][0]['Status'], 'SearchFix')
        self.assertEqual(report['summary'][0]['Grand Total'], 5)
        self.assertEqual({row['Order Number'].strip() for row in report['rows']}, {'001', '002', '003', '004', '005'})
        self.assertEqual(sum(report['totals'][name] for name in ('Future SLA', 'Negative SLA', 'Today SLA')), 5)
        self.assertEqual(next(row['Status Reason'] for row in report['rows'] if row['Order Number'] == '002'), 'Task Name is SearchFix')
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
        self.assertEqual({row[0].strip(): row[11] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)},
                         dict.fromkeys(['001', '002', '003', '004', '005'], 'SearchFix'))
        edited = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'SearchFix',
            'replacement': 'Ready to Send', 'order_numbers': ['002'], 'edit_scope': 'orders'})
        self.assertEqual(edited.status_code, 200, edited.json)
        updated = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(next(row['Status'] for row in updated['rows'] if row['Order Number'] == '002'), 'Ready to Send')
        self.assertEqual(next(row['Grand Total'] for row in updated['summary'] if row['Status'] == 'SearchFix'), 4)
        self.assertEqual(updated['totals'], report['totals'])
        self.assertTrue(edited.json['publication']['synced'])
        self.assertEqual(self.store.history(), before)

    def test_selected_preview_counts_unique_current_orders_without_missing_or_future_orders(self):
        self.capture(1, [{'Order Number': ' 001 ', 'Task Name': 'Search', 'SLA Expiration*': '2h'},
                         {'Order Number': 'gone', 'Task Name': 'Search', 'SLA Expiration*': '2h'},
                         {'Order Number': 'GONE ', 'Task Name': 'Other task', 'SLA Expiration*': '2h'}],
                     created='2026-10-16T04:30:00Z')
        self.capture(3, [{'Order Number': '001', 'Task Name': 'Typing', 'SLA Expiration*': '2h'},
                         {'Order Number': 'new', 'Task Name': 'Search', 'SLA Expiration*': '2h'},
                         {'Order Number': 'NEW ', 'Task Name': 'CRSP2', 'SLA Expiration*': '2h'}])
        self.capture(5, [{'Order Number': 'future', 'SLA Expiration*': '2h'}])
        before = deepcopy(self.store.history())
        report = build_sla_status(preview_snapshot(self.store, preview_id=3), datetime(2026, 10, 19, 10))
        self.assertEqual(report['captured_order_count'], 2)
        self.assertEqual(report['totals']['Grand Total'], 2)
        self.assertNotIn('gone', [row['Order Number'].strip().casefold() for row in report['rows']])
        self.assertTrue(all(row['Source Preview']=='preview3' for row in report['rows']))
        self.assertEqual(sum(report['totals'][name] for name in ('Future SLA','Negative SLA','Today SLA')),2)
        self.assertEqual(next(row for row in report['rows'] if row['Order Number'].strip() == 'NEW')['Status'], 'SearchFix')
        self.assertNotIn('future', [row['Order Number'] for row in report['rows']])
        self.assertEqual(self.store.history(), before)
        first = preview_snapshot(self.store, preview_id=1)
        self.assertEqual(first['captured_order_count'], 2)

    def test_eta_assignment_and_apology_match_only_the_requested_comment_words(self):
        cases = [
            ({'ETA Comments': 'ETA: Order has been assigned to the abst.'}, 'Assign to ABS'),
            ({'ETA Comments': 'ETA: Order has been assigned to the abstractor, and we will ensure completion within SLA.'}, 'Assign to ABS'),
            ({'ETA Comments': 'ETA: ASSIGNED\nTO THE ABSTRACTOR. Completion within SLA.'}, 'Assign to ABS'),
            ({'ETAComments': 'ETA: ASSIGNED\nTO THE\tABST..', 'LastUser': '', 'TaskStatus': 'Available'}, 'Assign to ABS'),
            ({'ETA Comments': 'Hello, we apologize for the delay.'}, 'Assign to ABS'),
            ({'ETA Comments': 'Hello, we Apoloize for the delay.'}, 'Assign to ABS'),
            ({'ETA Comments': 'We APOLOGISE for the delay.'}, 'Assign to ABS'),
            ({'ETA Comments': 'We apologize. Order assigned to the abst.'}, 'Assign to ABS'),
            ({'Comment': 'assigned to the abst; we apologize', 'Status': 'QC in Progress'}, 'QC in Progress'),
            ({'ETA Comments': 'preassigned to the abst', 'Status': 'QC in Progress'}, 'QC in Progress'),
            ({'ETA Comments': 'assigned to the abstainer', 'Status': 'QC in Progress'}, 'QC in Progress'),
            ({'ETA Comments': 'assigned to the abstractors', 'Status': 'QC in Progress'}, 'Assign to ABS'),
            ({'ETA Comments': 'assigned to the abstractorElse', 'Status': 'QC in Progress'}, 'QC in Progress'),
            ({'ETA Comments': 'unapologize or apologizer', 'Status': 'QC in Progress'}, 'QC in Progress'),
        ]
        originals = [dict(fields, **{'Order Number': str(index), 'SLA Expiration*': '2h'})
                     for index, (fields, _) in enumerate(cases)]
        self.capture(1, originals)
        report = build_sla_status(preview_snapshot(self.store), datetime(2026, 10, 19, 10))
        self.assertEqual([row['Status'] for row in report['rows']], [status for _, status in cases])
        self.assertEqual(report['totals']['Grand Total'], len(cases))
        self.assertEqual(self.store.get(1)['rows'], originals)

    def test_abstractor_mentions_override_search_fallback_and_keep_rule_priority(self):
        comments = [
            'ETA: Hello, As per the abstractor, waiting for FVD and the ETA 10/22',
            'We just received the package from the abstractor.',
            'We have assigned it to the abstractor and the ETA is tomorrow.',
            'The abstractor has been unwell for the past two days.',
            'The order has been assigned to multiple abstractors.',
            "The ABSTRACTOR'S update confirms the ETA.",
            'We requested copies from the ABSTRACTORS.\nWaiting for an update.',
        ]
        originals = [{'Order Number': f'ABS-{index}', 'Task Status': 'Available', 'Last User': '',
                      'ETAComments': comment, 'SLA Expiration*': '2h'}
                     for index, comment in enumerate(comments)]
        originals += [
            {'Order Number': 'held', 'ETA Comments': 'ON HOLD; awaiting the abstractor', 'SLA Expiration*': '2h'},
            {'Order Number': 'apology', 'ETA Comments': 'We apologize for the abstractor delay', 'SLA Expiration*': '2h'},
            {'Order Number': 'review', 'Task Name': 'CRSP2', 'ETA Comments': 'Update from abstractor', 'SLA Expiration*': '2h'},
            {'Order Number': 'manual', 'Task Status': 'Available', 'ETA Comments': 'Update from abstractor', 'SLA Expiration*': '2h'},
        ]
        self.capture(1, originals)
        saved = deepcopy(self.store.history())
        report = build_sla_status(preview_snapshot(self.store, {
            'sla_preview_statuses': {'1': {'manual': 'Ready to Send'}}}))
        actual = {row['Order Number']: row['Status'] for row in report['rows']}
        self.assertTrue(all(actual[f'ABS-{index}'] == 'Assign to ABS' for index in range(len(comments))))
        self.assertEqual({key: actual[key] for key in ('held', 'apology', 'review', 'manual')}, {
            'held': 'Waiting for Effective Date', 'apology': 'Assign to ABS',
            'review': 'SearchFix', 'manual': 'Ready to Send'})
        self.assertEqual(report['totals']['Grand Total'], len(originals))
        self.assertEqual(self.store.history(), saved)
        for text in ('abstraction', 'abstainer', 'abstractorElse', 'nonabstractors'):
            self.assertEqual(queue_rule({'ETA Comments': text, 'Task Status': 'Available'})[1], 'Search In Progress')

    def test_abstract_and_dated_eod_or_sooner_comments(self):
        today = datetime(2026, 10, 21, 10)
        screenshot = ('ETA: Hello, This order was received on 10/08, and the fee was approved on 10/15, '
                      'and as per them, the ETA will be {} EOD or sooner.')
        cases = [
            ('Update from Abstract.', 'Assign to ABS'),
            ('ABSTRACT is working on the order.', 'Assign to ABS'),
            (screenshot.format('10/21'), 'Assign to ABS'),
            (screenshot.format('10/22'), 'Assign to ABS'),
            (screenshot.format('10/23'), 'Assign to ABS'),
            (screenshot.format('10/20'), 'Assign to ABS'),
            ('Received on 10/21; fee approved 10/21; ETA 10/22 EOD or sooner.', 'Assign to ABS'),
            ('ETA 10/21 sooner.', 'Assign to ABS'),
            ('ETA 10/21 or SOONER.', 'Assign to ABS'),
            ('ETA 10/21/2026\nEoD\nor sooner.', 'Assign to ABS'),
            ('ETA 10/21/26 EOD.', 'Assign to ABS'),
            ('ETA 2026-10-21 EOD.', 'Assign to ABS'),
            ('ETA 10/21/2025 EOD.', 'Assign to ABS'),
            ('As per the abstractor, ETA 10/21 EOD or sooner.', 'Assign to ABS'),
            ('ETA was 10/21 EOD; updated ETA 10/22 EOD.', 'Assign to ABS'),
            ('ETA was 10/22 EOD; updated ETA 10/21 EOD.', 'Assign to ABS'),
            ('Received 10/21; ETA EOD or sooner.', 'Assign to ABS'),
            ('Will finish sooner.', 'Assign to ABS'),
            ('ETA EOD.', 'Assign to ABS'),
            ('ETA 02/30 EOD or sooner.', 'Assign to ABS'),
            ('ETA 13/21 EOD or sooner.', 'Assign to ABS'),
            ('ETA 110/21 EOD.', 'Assign to ABS'),
        ]
        for comment, expected in cases:
            with self.subTest(comment=comment):
                self.assertEqual(queue_rule({'ETAComments': comment, 'TaskStatus': 'Available'}, today)[1], expected)
        for comment in ('abstractorElse', 'nonabstractors', 'abstraction', 'EODish', 'soonerish'):
            self.assertEqual(queue_rule({'ETA Comments': comment}, today)[1], '')
        self.assertEqual(queue_rule({'Comment': 'Abstract; ETA 10/21 EOD or sooner'}, today)[1], '')
        self.assertIn('ETA Comments contains',
                      queue_rule({'ETA Comments': screenshot.format('10/21')}, today)[2])

    def test_eta_dates_keep_one_abs_group_and_specific_rule_priority(self):
        fields = {'ETA Comments': 'ETA 10/21 EOD or sooner.'}
        self.assertEqual(queue_rule(fields, datetime(2026, 10, 20, 20, tzinfo=timezone.utc))[1], 'Assign to ABS')
        self.assertEqual(queue_rule(fields, datetime(2026, 10, 21, 19, tzinfo=timezone.utc))[1], 'Assign to ABS')
        self.assertEqual(queue_rule(fields, date(2026, 10, 21))[1], 'Assign to ABS')
        self.assertEqual(queue_rule({'ETA Comments': 'ETA 02/29 EOD'}, date(2028, 2, 29))[1], 'Assign to ABS')
        self.assertEqual(queue_rule({'ETA Comments': 'ETA 12/31 EOD'}, date(2027, 1, 1))[1], 'Assign to ABS')
        cases = [({'Out Time': '10/21/2026 09:15 AM'}, 'Completed and Delivered'),
                 ({'TaskName': 'CRSP2'}, 'SearchFix'),
                 ({'Product': 'Quality Control'}, 'SearchFix'),
                 ({'Product': 'Attorney Approval'}, 'Attorney Approval'),
                 ({'Product': 'Home Builders'}, 'Home Builder'),
                 ({'ETA Comments': 'ON HOLD; ETA 10/21 EOD'}, 'Waiting for Effective Date'),
                 ({'ETA Comments': 'We apologize; ETA 10/22 EOD'}, 'Assign to ABS')]
        for changed, expected in cases:
            with self.subTest(changed=changed):
                self.assertEqual(queue_rule(dict(fields, **changed), date(2026, 10, 21))[1], expected)

    def test_eta_date_refreshes_old_previews_deduplicates_and_keeps_manual_status(self):
        originals = [
            {'Order Number': 'same', 'ETA Comments': 'ETA 10/22 EOD'},
            {'Order Number': 'SAME ', 'ETA Comments': 'ETA 10/21 EOD'},
            {'Order Number': 'tomorrow', 'ETA Comments': 'ETA 10/22 EOD or sooner'},
            {'Order Number': 'manual', 'ETA Comments': 'ETA 10/21 sooner'},
        ]
        self.capture(1, [dict(row, **{'SLA Expiration*': '2099-10-22'}) for row in originals])
        before = deepcopy(self.store.history())
        settings = {'sla_preview_statuses': {'1': {'manual': 'QC in Progress'}}}
        for today in (datetime(2026, 10, 21, 10), datetime(2026, 10, 22, 10)):
            snapshot = preview_snapshot(self.store, settings, as_of=today)
            report = build_sla_status(snapshot, today)
            statuses = {row['Order Number'].strip().casefold(): row['Status'] for row in report['rows']}
            self.assertEqual(statuses, {'same': 'Assign to ABS', 'manual': 'QC in Progress',
                'tomorrow': 'Assign to ABS'})
            self.assertEqual(report['totals']['Grand Total'], 3)
        self.assertEqual(self.store.history(), before)

    def test_eta_statuses_reconcile_across_api_excel_and_sheets(self):
        today = datetime(2026, 10, 21, 10)
        self.capture(1, [
            {'Order Number': 'today', 'ETA Comments': 'ETA 10/21 EOD or sooner', 'SLA Expiration*': '1d'},
            {'Order Number': 'future', 'ETA Comments': 'Received 10/21; ETA 10/22 EOD or sooner', 'SLA Expiration*': '1d'},
            {'Order Number': 'plain', 'ETA Comments': 'Update from Abstract', 'SLA Expiration*': '1d'},
        ])
        before = deepcopy(self.store.history())
        with patch('sla_status.queue_rule', side_effect=lambda row, as_of=None: queue_rule(row, today)):
            report = self.client.get('/api/sla-status?preview_id=1').json
            expected = {'today': 'Assign to ABS', 'future': 'Assign to ABS', 'plain': 'Assign to ABS'}
            self.assertEqual({row['Order Number']: row['Status'] for row in report['rows']}, expected)
            self.assertEqual({row['Status']: row['Grand Total'] for row in report['summary']},
                             {'Assign to ABS': 3})
            book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
            self.assertEqual({row[0]: row[11] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)}, expected)
            edited = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'Assign to ABS',
                'replacement': 'QC in Progress', 'order_numbers': ['today'], 'edit_scope': 'orders'})
            self.assertEqual(edited.status_code, 200, edited.json)
            self.assertTrue(edited.json['publication']['synced'])
            updated = self.client.get('/api/sla-status?preview_id=1').json
            self.assertEqual({row['Order Number']: row['Status'] for row in updated['rows']},
                             dict(expected, today='QC in Progress'))
            sheet = self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')
            self.assertEqual({row[0]: row[-1] for row in sheet[2:-1]}, {'Assign to ABS': 2, 'QC in Progress': 1})
        self.assertEqual(self.store.history(), before)

    def test_full_eta_comment_is_preserved_in_lookup_export_and_saved_details(self):
        full = ('ETA: Order has been assigned to the abstractor, and we will ensure completion within SLA.\n'
                'The complete comment remains available for the team to review.')
        self.capture(1, [{'Order Number': '001', 'Task Name': 'Search', 'Task Status': 'Available',
                          'SLA Expiration*': '1d', 'ETA Comments': full}])
        report = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(report['rows'][0]['Status'], 'Assign to ABS')
        self.assertEqual(report['rows'][0]['ETA Comments'], full)
        self.assertEqual(report['rows'][0]['Status Reason'], 'ETA Comments contains assigned to the abstractor')
        lookup = self.client.get('/api/sla-order?preview_id=1&order_number=001').json
        self.assertEqual(lookup['fields']['ETA Comments'], full)
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
        self.assertEqual(book['Sheet1'].cell(2, 14).value, full)
        saved = self.client.post('/api/sla-order', json={'preview_id': 1, 'order_number': '001',
            'status': 'Assign to ABS', 'fields': {'ETA Comments': full}, 'expected_revision': lookup['revision']})
        self.assertEqual(saved.status_code, 200, saved.json)
        self.assertEqual(self.client.get('/api/sla-order?preview_id=1&order_number=001').json['fields']['ETA Comments'], full)
        self.assertEqual(self.store.get(1)['rows'][0]['ETA Comments'], full)

    def test_task_suspended_is_a_fallback_after_every_specific_queue_rule(self):
        cases = [
            ({}, 'Typing in Progress'),
            ({'TaskName': 'Search', 'TaskStatus': 'TASK\nSUSPENDED'}, 'Typing in Progress'),
            ({'Task Name': 'TypingModule'}, 'Typing in Progress'),
            ({'ETA Comments': 'ETA: ON HOLD - waiting for plant date'}, 'Waiting for Effective Date'),
            ({'Task Name': 'CRSP2', 'ETA Comments': 'assigned to the abst'}, 'SearchFix'),
            ({'Product': 'ClientReview'}, 'SearchFix'),
            ({'Product': 'Quality Control'}, 'SearchFix'),
            ({'Product': 'Attorney Appro..'}, 'Attorney Approval'),
            ({'Product': 'Home Builders'}, 'Home Builder'),
            ({'ETA Comments': 'Order has been assigned to the abst.'}, 'Assign to ABS'),
            ({'ETA Comments': 'Hello, we apologize for the delay'}, 'Assign to ABS'),
            ({'WorkflowSuspended': True}, 'Awaiting for Clarification'),
            ({'Task Status': 'Workflow Suspended'}, 'Awaiting for Clarification'),
            ({'ETA Comments': 'ON HOLD; we apologize; assigned to the abst'}, 'Waiting for Effective Date'),
        ]
        for fields, expected in cases:
            with self.subTest(fields=fields):
                # Use either the portal's spaced headers or its compact aliases.
                row = dict(fields) if 'TaskStatus' in fields else dict({'Task Status': 'Task Suspended'}, **fields)
                self.assertEqual(queue_rule(row)[1], expected)
        self.assertEqual(queue_rule({'Task Name': 'Search', 'Task Status': 'In Progress'})[1], '')
        self.assertEqual(queue_rule({'Task Name': 'TypingModule', 'Task Status': 'In Progress',
            'ETA Comments': 'assigned to the abst'})[1], 'Assign to ABS')

    def test_task_duplicates_use_the_strongest_rule_and_manual_corrections_still_win(self):
        self.capture(1, [
            {'Order Number': ' shared ', 'Task Status': 'Task Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': 'SHARED', 'ETA Comments': 'assigned to the abst', 'SLA Expiration*': '2h'},
            {'Order Number': 'held', 'Task Status': 'Task Suspended', 'ETA Comments': 'ON HOLD', 'SLA Expiration*': '2h'},
            {'Order Number': 'HELD ', 'ETA Comments': 'we apologize', 'SLA Expiration*': '2h'},
            {'Order Number': 'review', 'Task Status': 'Task Suspended', 'SLA Expiration*': '2h'},
            {'Order Number': 'REVIEW', 'Task Name': 'CRSP2', 'SLA Expiration*': '2h'},
            {'Order Number': 'unassigned', 'Task Name': 'TypingModule', 'Task Status': 'In Progress',
             'ETA Comments': 'we apologize', 'SLA Expiration*': '2h'},
            {'Order Number': 'typing', 'Task Name': 'Search', 'Task Status': 'Task Suspended', 'SLA Expiration*': '2h'},
        ])
        before = deepcopy(self.store.history())
        report = build_sla_status(preview_snapshot(self.store, {
            'sla_preview_statuses': {'1': {'shared': 'Ready to Send'}}}), datetime(2026, 10, 19, 10))
        self.assertEqual({row['Order Number'].strip().casefold(): row['Status'] for row in report['rows']}, {
            'shared': 'Ready to Send', 'held': 'Waiting for Effective Date', 'review': 'SearchFix',
            'unassigned': 'Assign to ABS', 'typing': 'Typing in Progress'})
        self.assertEqual(report['totals']['Grand Total'], 5)
        self.assertEqual(sum(report['totals'][name] for name in ('Future SLA', 'Negative SLA', 'Today SLA')), 5)
        self.assertEqual(self.store.history(), before)

    def test_new_rule_statuses_match_api_excel_and_sheets_after_an_individual_edit(self):
        self.capture(1, [
            {'Order Number': 'assigned', 'Task Status': 'Task Suspended', 'ETA Comments': 'assigned to the abst', 'SLA Expiration*': '2h'},
            {'Order Number': 'unassigned', 'Task Status': 'Task Suspended', 'ETA Comments': 'we apologize', 'SLA Expiration*': '-2h'},
            {'Order Number': 'typo', 'ETA Comments': 'we Apoloize', 'SLA Expiration*': '1d'},
            {'Order Number': 'typing', 'Task Name': 'Search', 'Task Status': 'Task Suspended', 'SLA Expiration*': '1d'},
        ])
        before = deepcopy(self.store.history())
        initial = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual({row['Order Number']: row['Status'] for row in initial['rows']}, {
            'assigned': 'Assign to ABS', 'unassigned': 'Assign to ABS', 'typo': 'Assign to ABS',
            'typing': 'Typing in Progress'})
        response = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'Assign to ABS',
            'replacement': 'Ready to Send', 'order_numbers': ['unassigned'], 'edit_scope': 'orders'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue(response.json['publication']['synced'])
        report = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(report['totals'], initial['totals'])
        self.assertEqual(next(row['Status'] for row in report['rows'] if row['Order Number'] == 'unassigned'), 'Ready to Send')
        book = load_workbook(BytesIO(self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=1').data))
        self.assertEqual({row[0]: row[11] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)},
                         {row['Order Number']: row['Status'] for row in report['rows']})
        self.assertEqual(book['Sheet2']._pivots[0].cache.recordCount, 4)
        sheet = self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')
        self.assertEqual({row[0]: row[-1] for row in sheet[2:-1]},
                         {row['Status']: row['Grand Total'] for row in report['summary']})
        self.assertEqual(self.store.history(), before)

    def test_shortened_product_names_get_separate_status_rows_and_export_counts(self):
        labels = [('Attorney Appro..', 'Attorney Approval'), ('Attorney Appro\u2026', 'Attorney Approval'),
                  ('Attorney Approach', 'Attorney Approval'), ('Attorney Approval', 'Attorney Approval'),
                  ('Home Buil..', 'Home Builder'), ('Home Builder..', 'Home Builder'),
                  ('Home Builders', 'Home Builder')]
        self.capture(1, [{'Order Number': str(index), 'Product': product, 'Last User': '',
                          'Task Status': 'Available', 'Status': 'Search In Progress',
                          'SLA Expiration*': '2099-10-19'}
                         for index, (product, _) in enumerate(labels)])
        report = self.client.get('/api/sla-status').json
        self.assertEqual([row['Status'] for row in report['rows']], [status for _, status in labels])
        self.assertEqual({row['Status']: row['Grand Total'] for row in report['summary']},
                         {'Attorney Approval': 4, 'Home Builder': 3})
        book = load_workbook(excel_report(report), data_only=True)
        self.assertEqual([(book['Sheet2'].cell(row, 1).value, book['Sheet2'].cell(row, 5).value)
                          for row in (5, 6)], [('Attorney Approval', 4), ('Home Builder', 3)])
        self.assertEqual([row[5] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)],
                         [product for product, _ in labels])
        self.assertEqual(book['Sheet2']._pivots[0].cache.recordCount, 7)
        for unrelated in ('Attorney', 'Attorney App', 'Attorney Appointment', 'Attorney Appt..',
                          'Home', 'Home Buyer', 'Home Inspection..'):
            with self.subTest(product=unrelated):
                self.assertEqual(special_product_status(unrelated), '')

    def test_manual_edit_overrides_queue_rules_and_only_changes_selected_order(self):
        self.capture(1, [{'Order Number': number, 'Task Name': 'CRSP2', 'SLA Expiration*': '2h'}
                         for number in ('001', '002')])
        before = deepcopy(self.store.history())
        body = {'preview_id': 1, 'status': 'SearchFix', 'replacement': 'Ready to Send',
                'order_numbers': ['001'], 'edit_scope': 'orders'}
        response = self.client.post('/api/order-status', json=body)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json['updated'], 1)
        report = self.client.get('/api/sla-status').json
        self.assertEqual({row['Order Number']: row['Status'] for row in report['rows']},
                         {'001': 'Ready to Send', '002': 'SearchFix'})
        self.assertIn('Manual', report['rows'][0]['Status Reason'])
        self.assertEqual(self.store.history(), before)
        stale = self.client.post('/api/order-status', json=body)
        self.assertEqual(stale.status_code, 409)
        invalid = self.client.post('/api/order-status', json=dict(body, order_numbers=['unknown']))
        self.assertEqual(invalid.status_code, 409)

    def test_missing_orders_are_excluded_and_saved_active_orders_remain_in_matching_exports(self):
        self.capture(1, [{'Order Number': 'gone', 'SLA Expiration*': '1d'}])
        self.capture(2, [{'Order Number': 'live', 'Product': 'Home Builders', 'SLA Expiration*': '2h'},
                         {'Order Number': 'already-done', 'Status': 'Ready to Send', 'SLA Expiration*': '1d'}])
        report = self.client.get('/api/sla-status').json
        self.assertEqual({row['Order Number'] for row in report['rows']}, {'live','already-done'})
        self.assertEqual(next(row['Status'] for row in report['rows'] if row['Order Number']=='already-done'), 'Ready to Send')
        downloaded = self.client.get('/api/export/report?sheet=SLA%20Status&preview_id=2')
        book = load_workbook(BytesIO(downloaded.data), data_only=True)
        self.assertEqual({row[0] for row in book['Sheet1'].iter_rows(min_row=2, values_only=True)}, {'already-done', 'live'})
        self.assertEqual(book['Sheet1'].max_row - 1, report['totals']['Grand Total'])
        self.assertEqual(book['Sheet2']._pivots[0].cache.recordCount, report['totals']['Grand Total'])
        missing = self.client.post('/api/order-status', json={'preview_id': 2,
            'status': 'Ready to Send', 'replacement': 'QC in Progress',
            'order_numbers': ['gone'], 'edit_scope': 'orders'})
        self.assertEqual(missing.status_code, 409)
        changed = self.client.post('/api/order-status', json={'preview_id': 2,
            'status': 'Home Builder', 'replacement': 'QC in Progress',
            'order_numbers': ['live'], 'edit_scope': 'orders'})
        self.assertEqual(changed.status_code, 200, changed.json)
        updated = self.client.get('/api/sla-status').json
        self.assertEqual(next(row['Status'] for row in updated['rows'] if row['Order Number'] == 'live'), 'QC in Progress')
        sheet = self.book.worksheet('SLA Status').get_all_values(value_render_option='UNFORMATTED_VALUE')
        self.assertEqual(sheet[-1][-1], updated['totals']['Grand Total'])


if __name__ == '__main__':
    unittest.main()

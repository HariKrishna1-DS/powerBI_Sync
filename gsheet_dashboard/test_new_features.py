import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path
import tempfile
import pandas as pd
from openpyxl import load_workbook

from datatrace_sync import status_report_frame
from order_reporting import monthly_orders
from preview_store import PreviewStore
from unittest.mock import patch
from test_production_sync import Book
from server import create_app


class NewFeaturesTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.store = PreviewStore(self.root / 'previews')

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_status_report_frame_with_metadata(self):
        df = pd.DataFrame([
            {'Task Status': 'Search In Progress'},
            {'Task Status': 'Available'},
            {'Task Status': 'Available'},
        ])
        # Test backwards compatibility
        df_no_meta = status_report_frame(df)
        self.assertEqual(list(df_no_meta.columns), ['Status', 'Orders', 'Share'])

        # Test with preview name and sync time
        sync_time = '2026-09-30 10:00:00 AM'
        df_meta = status_report_frame(df, preview_name='preview1', sync_time=sync_time)
        self.assertEqual(list(df_meta.columns), ['Status', 'Orders', 'Share', 'Preview', 'Sync Date & Time'])
        self.assertEqual(df_meta['Preview'].tolist(), ['preview1', 'preview1'])
        self.assertEqual(df_meta['Sync Date & Time'].tolist(), [sync_time, sync_time])

    def test_multi_sheet_excel_download(self):
        raw = pd.DataFrame([{'Order Number': '101', 'Task Status': 'Available', 'Product': 'Full Title'}])
        preview = self.store.save(raw)
        response = create_app(self.store.root).test_client().get(f"/api/previews/{preview['id']}/download/xlsx")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.data))
        self.assertEqual(wb.sheetnames, ['DataTraceQueue'])
        self.assertEqual(wb.active['B2'].value, 'Available')

    def test_sync_schedule_multiple_times(self):
        client = create_app(self.store.root).test_client()

        # Update with multiple times
        res = client.post('/api/sync-schedule', json={
            'enabled': True,
            'times': ['10:00', '09:00', '14:30']
        })
        self.assertEqual(res.status_code, 200)
        data = res.json
        self.assertTrue(data['enabled'])
        # Should be sorted and deduplicated
        self.assertEqual(data['times'], ['09:00', '10:00', '14:30'])
        self.assertEqual(data['time'], '09:00')

        # Verify state endpoint returns updated schedule
        state_res = client.get('/api/state')
        self.assertEqual(state_res.json['schedule']['times'], ['09:00', '10:00', '14:30'])

    def test_monthly_orders_and_endpoint(self):
        rows = [{'Order Number': '1', 'Date': '09/01/2026', 'Product': 'Full Title', 'Status': 'Search In Progress'},
                {'Order Number': '2', 'Date': '09/15/2026', 'Product': 'Full Title', 'Status': 'Completed and Delivered'}]
        report = monthly_orders(sheet_rows=rows)[0]
        self.assertEqual(report['Month Orders'], 2)
        self.assertEqual(report['Completed Orders'], 1)
        self.assertEqual(report['missing_ids'], [])
        book = Book([rows, []])
        with patch('server.target_worksheet', return_value=(book, book.sheets[0])):
            response = create_app(self.store.root).test_client().get('/api/monthly-orders')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['rows'][0]['Month Orders'], 2)


if __name__ == '__main__':
    unittest.main()

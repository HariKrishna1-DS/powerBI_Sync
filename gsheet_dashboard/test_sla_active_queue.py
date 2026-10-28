"""Active queue SLA classifications must never inherit production completion."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import unittest

from openpyxl import load_workbook
import test_sla_queue_rules as queue_tests
from sla_status import build_sla_status, preview_snapshot, queue_rule, excel_report, sheet_summary_values
from report_workspace import imported_snapshot


class SlaActiveQueueTests(unittest.TestCase):
    setUp = queue_tests.SlaQueueRulesTests.setUp
    capture = queue_tests.SlaQueueRulesTests.capture

    def test_saved_tracker_manual_and_completion_evidence_do_not_complete_queue_sla(self):
        rows = [
            {'Order Number': 'A', 'Task Name': 'UpdateSearch', 'Task Status': 'In Progress', 'Last User': 'Searcher'},
            {'Order Number': 'B', 'Status': 'Completed and Delivered', 'Task Name': 'Search',
             'Task Status': 'In Progress', 'Last User': 'Searcher'},
            {'Order Number': 'C', 'Task Name': 'SearchFix', 'Status': 'Completed and Delivered',
             'Out Time': '10/22/2026 09:00 AM', 'Completed Time': '10/22/2026 09:00 AM'},
        ]
        self.capture(1, [dict(row, **{'SLA Expiration*': '1d'}) for row in rows])
        self.workspace.save(sla_preview_statuses={'1': {'a': 'Completed and Delivered'}},
                            sheet_status_overrides={'b': 'Completed and Delivered'},
                            sla_preview_custom_statuses={'1': ['Completed and Delivered']})
        before = deepcopy(self.store.history())
        production = {'sheets': {'Overview': {'rows': [dict(row, Status='Completed and Delivered') for row in rows]}}}
        report = build_sla_status(preview_snapshot(self.store, self.workspace.read(), 1, production), datetime(2026, 10, 22))
        self.assertEqual({row['Order Number']: row['Status'] for row in report['rows']},
                         {'A': 'Search In Progress', 'B': 'Search In Progress', 'C': 'SearchFix'})
        self.assertIn('Task Status is In Progress', report['rows'][0]['Status Reason'])
        self.assertNotIn('Completed and Delivered', [row['Status'] for row in report['summary']])
        self.assertEqual(report['totals']['Grand Total'], 3)
        self.assertEqual(self.store.history(), before)
        self.assertEqual(queue_rule(rows[2])[1], 'Completed and Delivered')
        imported = imported_snapshot({'rows': [rows[2]], 'columns': list(rows[2]), 'files': [], 'created': '2026-10-22'})
        self.assertEqual(imported['sheets']['Overview']['rows'][0]['Status'], 'Completed and Delivered')
        workbook = load_workbook(excel_report(report), data_only=True)
        self.assertNotIn('Completed and Delivered', [cells[11] for cells in workbook['Sheet1'].iter_rows(min_row=2, values_only=True)])
        self.assertNotIn('Completed and Delivered', [values[0] for values in sheet_summary_values(report)])

    def test_manual_queue_status_creation_renaming_and_order_save_reject_completion(self):
        self.capture(1, [{'Order Number': 'A', 'Task Name': 'Search', 'Task Status': 'Available', 'SLA Expiration*': '1d'}])
        for name in ('Completed and Delivered', 'completed', ' COMPLETED  AND DELIVERED '):
            response = self.client.post('/api/sla-statuses', json={'preview_id': 1, 'status': name})
            self.assertEqual(response.status_code, 409, response.json)
            self.assertIn('not used in extracted SLA Status', response.json['error'])
        moved = self.client.post('/api/order-status', json={'preview_id': 1, 'status': 'Search In Progress',
            'replacement': 'Completed and Delivered', 'order_numbers': ['a']})
        self.assertEqual(moved.status_code, 409, moved.json)
        lookup = self.client.get('/api/sla-order?preview_id=1&order_number=A').json
        saved = self.client.post('/api/sla-order', json={'preview_id': 1, 'order_number': 'A',
            'status': 'Completed and Delivered', 'fields': {'SLA Expiration': '2099-10-22T17:00'},
            'expected_revision': lookup['revision']})
        self.assertEqual(saved.status_code, 409, saved.json)

    def test_repeat_selection_uses_workflow_condition_instead_of_completed_evidence(self):
        self.capture(1, [
            {'Order Number': 'A', 'Task Status': 'Available', 'Out Time': '10/22/2026 09:00 AM', 'SLA Expiration*': '1d'},
            {'Order Number': 'A', 'Task Status': 'Workflow Suspended', 'SLA Expiration*': '1d'},
        ])
        report = self.client.get('/api/sla-status?preview_id=1').json
        self.assertEqual(report['duplicate_groups'][0]['selected_capture_row'], 2)
        self.assertEqual(report['rows'][0]['Status'], 'Awaiting for Clarification')
        self.assertEqual(report['duplicate_groups'][0]['status_reason'], 'Workflow is suspended')


if __name__ == '__main__':
    unittest.main()

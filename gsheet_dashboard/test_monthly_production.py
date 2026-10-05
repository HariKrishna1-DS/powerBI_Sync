"""Month boundary, input fidelity, atomic move and recovery contract tests."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
import tempfile
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from requests.exceptions import Timeout

from monthly_production import (ARCHIVE, BASES, LEDGER, active_locations, apply_plan, arrival_sort,
    decode, encode, extend_headers, import_plan, local_datetime, migration_plan, monthly_workbook, next_month,
    preserve_formulas, read_import, rollover_plan, snapshot, tab_name)
from order_reporting import automatic_sync_frame
from preview_store import PreviewStore
from tracker_sync import HEADERS, sheet_reports, sync_trackers
from test_tracker_v2 import Sheet


class AtomicBook:
    """Emulate Sheets' atomic batch, including blank trailing ranges and duplicates."""
    def __init__(self):
        self.sheets, self.batches = [], []
        self.fail_before = False
        self.lose_reply = False

    def worksheets(self):
        return self.sheets

    def worksheet(self, title):
        return next(s for s in self.sheets if s.title == title)

    def add(self, title, rows=(), headers=HEADERS):
        self.sheets.append(Sheet(title, max([s.id for s in self.sheets], default=0)+1, encode(list(rows), extend_headers(headers, rows))))
        return self.sheets[-1]

    def fetch_sheet_metadata(self, params=None):
        return {'sheets': [{'properties': {'sheetId': s.id, 'title': s.title}, 'conditionalFormats': s.conditional_formats,
            'data': [{'rowData': [{'values': [{'userEnteredValue': {'formulaValue': val}} if (i,j) in getattr(s,'formula_cells',set()) else {}
                for j,val in enumerate(row)]} for i,row in enumerate(s.values)]}]} for s in self.sheets]}

    def batch_update(self, body):
        self.batches.append(deepcopy(body))
        if self.fail_before:
            raise ValueError('Permission denied before applying batch')
        working = deepcopy(self.sheets)
        for request in body['requests']:
            def find(number):
                return next(s for s in working if s.id == number)
            if 'duplicateSheet' in request:
                info = request['duplicateSheet']
                duplicate = deepcopy(find(info['sourceSheetId']))
                duplicate.id, duplicate.title = info['newSheetId'], info['newSheetName']
                working.append(duplicate)
            elif 'addSheet' in request:
                props = request['addSheet']['properties']
                if any(s.id == props['sheetId'] or s.title == props['title'] for s in working):
                    raise ValueError('Duplicate worksheet')
                working.append(Sheet(props['title'], props['sheetId']))
            elif 'updateSheetProperties' in request:
                props = request['updateSheetProperties']['properties']
                target = find(props['sheetId'])
                target.title = props.get('title', target.title)
            elif 'deleteConditionalFormatRule' in request:
                info = request['deleteConditionalFormatRule']
                find(info['sheetId']).conditional_formats.pop(info['index'])
            elif 'updateCells' in request:
                info = request['updateCells']
                region = info.get('range',info.get('start'))
                target = find(region['sheetId'])
                start=region.get('rowIndex',0)
                values = [[next(iter(c.get('userEnteredValue', {}).values()), '') for c in row['values']] for row in info['rows']]
                target.values = target.values[:start]+values
                target.formula_cells = {(i,j) for i,row in enumerate(info['rows'],start) for j,c in enumerate(row['values']) if 'formulaValue' in c.get('userEnteredValue',{})}
        self.sheets = working
        if self.lose_reply:
            self.lose_reply = False
            raise Timeout('Committed but reply lost')


def order(number, arrival='10/2/2026 12:50:45 AM', out='', **extra):
    return {'Order Number': number, 'Product': 'Full Title', 'Status': 'Search In Progress', 'In-Time': arrival, 'Out Time': out, **extra}


def source(rows, headers=HEADERS):
    return {'headers': headers, 'rows': rows, 'read': len(rows), 'skipped': 0, 'reviews': []}


class MonthlyProductionTests(unittest.TestCase):
    def setUp(self):
        self.book = AtomicBook()
        for base in BASES:
            self.book.add(tab_name(base, '2026-10'))

    def test_wall_clock_and_tolerant_dates(self):
        self.assertEqual(local_datetime(' 10/2/2026   12:50:45 AM '), datetime(2026,10,2,0,50,45))
        self.assertEqual(local_datetime('2026-10-02T23:30:00-07:00'), datetime(2026,10,2,23,30))
        self.assertIsNone(local_datetime('2/31/2026'))
        self.assertEqual(next_month('2026-12'), '2027-01')

    def test_sort_time_ties_invalid_and_serials(self):
        reviews=[]
        rows=arrival_sort([order('Z','10/2/2026 9:01 AM'),order('B','10/2/2026 9:00 AM'),order('A','10/2/2026 9:00 AM'),order('0','bad')],reviews)
        self.assertEqual([r['Order Number'] for r in rows], ['A','B','Z','0'])
        self.assertEqual([r['No'] for r in rows], [1,2,3,4])
        self.assertEqual(len(reviews), 1)

    def test_import_keeps_duplicate_headers_and_last_source_row(self):
        book=Workbook();sheet=book.active;sheet.title='TV Orders'
        headers=['No','Order number','In-Time','Out Time','SLA','SLA','Typer']
        sheet.append(headers);sheet.append([1,'001','10/2/2026 9:00 AM',None,2,3,'Jane'])
        sheet.append([2,'001','10/2/2026 10:00 AM',datetime(2026,10,2,11),4,5,'June'])
        sheet.append([3,None,'bad',None,None,None,None]);book.create_sheet('Escalations').append(['ignored'])
        stream=BytesIO();book.save(stream)
        result=read_import(stream.getvalue(),BASES[0])
        self.assertEqual(result['headers'],headers)
        self.assertEqual((result['read'],result['skipped'],len(result['rows'])),(3,2,1))
        self.assertEqual(encode(result['rows'],headers)[1][4:7],['4','5','June'])
        self.assertEqual(result['rows'][0]['Order Number'],'001')

    def test_migration_backups_and_renames_same_ids(self):
        book=AtomicBook();original=book.add(BASES[0],[order('SEPT')]);book.add(BASES[1])
        before=snapshot(book);plan=migration_plan(before)
        self.assertEqual(len(book.sheets),2)
        apply_plan(book,plan)
        self.assertEqual(book.worksheet(tab_name(BASES[0],ARCHIVE)).id,original.id)
        backup=next(s for s in book.sheets if s.title.startswith('__DataTrace_Backup') and s.values[1:])
        self.assertEqual(backup.values,original.values)
        self.assertEqual(len(book.batches),1)

    def test_stale_plan_rejected_without_backup_or_write(self):
        before=snapshot(self.book);plan=rollover_plan(before,'2026-10')
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values.append(['changed'])
        with self.assertRaisesRegex(ValueError,'changed after'):
            apply_plan(self.book,plan)
        self.assertFalse(self.book.batches)

    def test_import_rerun_updates_without_duplicate(self):
        before=snapshot(self.book)
        plan=import_plan(before,[(BASES[0],'full.xlsx',source([order('A')]))],'2026-10')
        apply_plan(self.book,plan)
        rerun=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('A',Searcher='Updated')]))],'2026-10')
        self.assertEqual((rerun['counts'][0]['added'],rerun['counts'][0]['updated']),(0,1))
        apply_plan(self.book,rerun)
        self.assertEqual(len(active_locations(snapshot(self.book))),1)

    def test_rollover_moves_once_keeps_arrival_and_archive(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A'),order('B',out='10/2/2026 1:00 AM')],HEADERS)
        archived=self.book.add(tab_name(BASES[0],ARCHIVE),[order('A')]);archive_values=deepcopy(archived.values)
        before=snapshot(self.book);plan=rollover_plan(before,'2026-10')
        self.assertEqual(len(plan['moves']),1)
        self.assertEqual(len(decode(self.book.worksheet(tab_name(BASES[0],'2026-10')).values)),2)
        apply_plan(self.book,plan)
        locations=active_locations(snapshot(self.book))
        self.assertEqual(locations['a'][0],tab_name(BASES[0],'2026-11'))
        self.assertEqual(locations['a'][1]['In-Time'],order('A')['In-Time'])
        self.assertEqual(locations['a'][1]['Carried From'],'Oct_2026')
        self.assertEqual(self.book.worksheet(archived.title).values,archive_values)
        self.assertEqual(rollover_plan(snapshot(self.book),'2026-10')['moves'],[])
        self.assertTrue(any(s.title==tab_name(BASES[1],'2026-11') for s in self.book.sheets))

    def test_lost_reply_and_repeated_confirmation_are_idempotent(self):
        plan=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('A')]))],'2026-10')
        self.book.lose_reply=True
        apply_plan(self.book,plan)
        count=len(self.book.batches)
        self.assertTrue(apply_plan(self.book,plan)['recovered'])
        self.assertEqual(len(self.book.batches),count)
        self.assertEqual(len(active_locations(snapshot(self.book))),1)

    def test_failed_batch_does_not_remove_source(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A')],HEADERS)
        before=snapshot(self.book);plan=rollover_plan(before,'2026-10');self.book.fail_before=True
        with self.assertRaises(ValueError):apply_plan(self.book,plan)
        self.assertEqual(snapshot(self.book),before)

    def test_invalid_out_time_is_retained_and_reported(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A',out='not a date')],HEADERS)
        plan=rollover_plan(snapshot(self.book),'2026-10')
        self.assertEqual(plan['moves'],[]);self.assertIn('Invalid Out Time',plan['reviews'][0]['Reason'])

    def test_duplicate_active_ownership_stops_rollover(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A')],HEADERS)
        self.book.add(tab_name(BASES[0],'2026-11'),[order('a')])
        with self.assertRaisesRegex(ValueError,'duplicate'):rollover_plan(snapshot(self.book),'2026-10')

    def test_completed_carried_order_remains_in_current_month(self):
        self.book.add(tab_name(BASES[0],'2026-11'),[order('A',out='10/31/2026 11:00 PM',**{'Carried From':'Oct_2026'})])
        self.assertFalse(rollover_plan(snapshot(self.book),'2026-11')['moves'])
        plan=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('A')]))],'2026-10')
        self.assertIn(tab_name(BASES[0],'2026-11'),plan['writes'])
        self.assertNotIn(tab_name(BASES[0],'2026-10'),plan['writes'])

    def test_formulas_preserved_and_untrusted_import_text_literal(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A',Searcher='=A2')],HEADERS)
        self.book.worksheet(tab_name(BASES[0],'2026-10')).formula_cells={(1,HEADERS.index('Searcher'))}
        before=snapshot(self.book);plan=preserve_formulas(rollover_plan(before,'2026-10'),before)
        apply_plan(self.book,plan)
        values=[cell for req in self.book.batches[-1]['requests'] if 'updateCells' in req for row in req['updateCells']['rows'] for cell in row['values']]
        self.assertIn({'userEnteredValue':{'formulaValue':'=A2'}},values)
        imported=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('X',Searcher='=HYPERLINK("https://example.org")')]))],'2026-10')
        self.assertFalse(preserve_formulas(imported,snapshot(self.book))['formula_cells'])

    def test_existing_literal_formula_looking_text_stays_literal(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A',Searcher='=literal')],HEADERS)
        before=snapshot(self.book);plan=preserve_formulas(rollover_plan(before,'2026-10'),before)
        self.assertFalse(plan['formula_cells'])
        apply_plan(self.book,plan)
        target=self.book.worksheet(tab_name(BASES[0],'2026-11'))
        self.assertEqual(decode(target.values)[0]['Searcher'],'=literal')
        self.assertFalse(target.formula_cells)

    def test_formulas_follow_their_order_when_sorting_changes_rows(self):
        target=self.book.worksheet(tab_name(BASES[0],'2026-10'))
        target.values=encode([order('Z',Searcher='=A2'),order('A',Searcher='=$A$1+A3')],HEADERS)
        target.formula_cells={(1,HEADERS.index('Searcher')),(2,HEADERS.index('Searcher'))}
        before=snapshot(self.book);plan=preserve_formulas(rollover_plan(before,'2026-10'),before)
        rows=decode(plan['writes'][tab_name(BASES[0],'2026-11')])
        self.assertEqual(rows[0]['Searcher'],'=$A$1+A2')
        self.assertEqual(rows[1]['Searcher'],'=A3')

    def test_selected_month_workbook_counts_headers_style_and_order(self):
        title=tab_name(BASES[0],'2026-10');headers=['No','Order number','In-Time','Out Time','Status','SLA','SLA']
        rows=[order('B','10/2/2026 1:00 PM'),order('A','10/2/2026 1:00 AM',Status='Awaiting for Clarification')]
        reports=sheet_reports({title:rows})['monthly'];report=reports[0]
        stream,name=monthly_workbook(report,[(BASES[0],title,encode(rows,headers))])
        self.assertEqual(name,'TV_Search_Production_Report_Oct_2026.xlsx')
        book=load_workbook(stream);self.assertEqual(book.sheetnames,['Full Search','C-O and Update','Summary'])
        sheet=book['Full Search'];self.assertEqual([c.value for c in sheet[1]][:len(headers)],headers)
        self.assertEqual(sheet['B2'].value,'A');self.assertEqual(sheet['A3'].value,2)
        self.assertEqual(book['Summary']['B3'].value,2)
        self.assertEqual(book['Summary']['C6'].value,.5)
        self.assertTrue(sheet['B2'].font.color.rgb.endswith('000000'))
        self.assertEqual(sheet['B2'].border.bottom.style,'thin')

    def test_capture_uses_month_route_and_does_not_reseed_september(self):
        with tempfile.TemporaryDirectory() as root:
            preview={'id':1,'name':'preview1','created':'2026-10-05T00:00:00+00:00','columns':['Order Number','Product','Arrival Time','Task Status'],
                     'rows':[{'Order Number':'NEW','Product':'Full Title','Arrival Time':'10/5/2026 9:00 AM','Task Status':'Available'}]}
            frame=automatic_sync_frame(preview)[0];frame.attrs['store_root']=root
            with patch('tracker_sync.load_defaults',side_effect=AssertionError('Must not seed September')):
                sync_trackers(frame,book=self.book)
                sync_trackers(frame,book=self.book)
            self.assertEqual(len(active_locations(snapshot(self.book))),1)
            self.assertEqual(active_locations(snapshot(self.book))['new'][0],tab_name(BASES[0],'2026-10'))
            self.assertEqual(len(self.book.worksheet('Sheet1').values),2)
            self.assertTrue(PreviewStore(root).get_sync_report(1)['committed'])

    def test_reimport_cannot_erase_a_later_completion(self):
        self.book.worksheet(tab_name(BASES[0],'2026-10')).values=encode([order('A',out='10/5/2026 9:00 AM',Status='Completed and Delivered')],HEADERS)
        plan=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('A')]))],'2026-10')
        actual=decode(plan['writes'][tab_name(BASES[0],'2026-10')])[0]
        self.assertEqual(actual['Out Time'],'10/5/2026 9:00 AM')
        self.assertEqual(actual['Status'],'Completed and Delivered')

    def test_api_429_retries_once_then_verifies_commit(self):
        from gspread.exceptions import APIError
        from requests import Response
        response=Response();response.status_code=429;response._content=b'{"error":{"code":429,"message":"quota"}}'
        original=self.book.batch_update;attempts=[]
        def intermittent(body):
            attempts.append(body)
            if len(attempts)==1:raise APIError(response)
            return original(body)
        plan=import_plan(snapshot(self.book),[(BASES[0],'full.xlsx',source([order('A')]))],'2026-10')
        with patch.object(self.book,'batch_update',side_effect=intermittent),patch('datatrace_sync.time.sleep') as sleep:
            apply_plan(self.book,plan)
        self.assertEqual(len(attempts),2);sleep.assert_called_once_with(1)

    def test_expired_preview_and_september_import_are_rejected(self):
        plan=rollover_plan(snapshot(self.book),'2026-10');plan['created']=time.time()-1900
        with self.assertRaisesRegex(ValueError,'expired'):apply_plan(self.book,plan)
        with self.assertRaisesRegex(ValueError,'archived'):import_plan(snapshot(self.book),[],'2026-09')
        self.assertFalse(self.book.batches)

    def test_later_capture_updates_carried_order_without_moving_it_back(self):
        self.book.add(tab_name(BASES[0],'2026-11'),[order('A',Searcher='Human',**{'Carried From':'Oct_2026'})])
        with tempfile.TemporaryDirectory() as root:
            rows=[{'Order Number':'A','Product':'Full Title','Arrival Time':'10/2/2026 12:50:45 AM','Out Time':'11/2/2026 9:00 AM','SLA Expiration':'11/2/2026 10:00 AM'}]
            preview={'id':1,'name':'preview1','created':'2026-11-02T04:30:00+00:00','columns':list(rows[0]),'rows':rows}
            frame=automatic_sync_frame(preview)[0];frame.attrs['store_root']=root
            sync_trackers(frame,book=self.book)
            title,row=active_locations(snapshot(self.book))['a']
            self.assertEqual(title,tab_name(BASES[0],'2026-11'))
            self.assertEqual(row['Searcher'],'Human');self.assertEqual(row['Free Site'],'On Time')

    def test_second_capture_appends_only_new_raw_rows(self):
        with tempfile.TemporaryDirectory() as root:
            for number in (1,2):
                rows=[{'Order Number':f'A{number}','Product':'Full Title','Arrival Time':'10/5/2026 9:00 AM'}]
                preview={'id':number,'name':f'preview{number}','created':f'2026-10-05T0{number}:00:00+00:00','columns':list(rows[0]),'rows':rows}
                frame=automatic_sync_frame(preview)[0];frame.attrs['store_root']=root
                sync_trackers(frame,book=self.book)
            history=self.book.worksheet('Sheet1')
            writes=[r['updateCells'] for r in self.book.batches[-1]['requests'] if 'updateCells' in r and r['updateCells'].get('start',{}).get('sheetId')==history.id]
            self.assertEqual(len(writes),1)
            self.assertEqual(writes[0]['start']['rowIndex'],2)
            self.assertEqual(len(writes[0]['rows']),1)
            self.assertEqual([r[0] for r in history.values[1:]],['preview1','preview2'])


class MonthlyApiTests(unittest.TestCase):
    def setUp(self):
        from server import create_app
        self.folder=tempfile.TemporaryDirectory();self.addCleanup(self.folder.cleanup)
        self.book=AtomicBook()
        self.book.add(tab_name(BASES[0],'2026-10'),[order('A')]);self.book.add(tab_name(BASES[1],'2026-10'))
        self.app=create_app(root=self.folder.name);self.client=self.app.test_client()
        patcher=patch('server.target_worksheet',return_value=(self.book,self.book.sheets[0]));patcher.start();self.addCleanup(patcher.stop)

    def test_api_requires_confirmation_and_handles_repeat(self):
        preview=self.client.post('/api/monthly-maintenance/preview',json={'kind':'rollover','month':'2026-10'})
        self.assertEqual(preview.status_code,200,preview.json)
        self.assertNotIn('writes',preview.json)
        identity=preview.json['id']
        self.assertEqual(self.client.post('/api/monthly-maintenance/apply',json={'id':identity}).status_code,422)
        for _ in range(2):
            result=self.client.post('/api/monthly-maintenance/apply',json={'id':identity,'confirmed':True})
            self.assertEqual(result.status_code,200,result.json)
        self.assertEqual(len(active_locations(snapshot(self.book))),1)
        activity=self.client.get('/api/monthly-maintenance').json['operations'][0]
        self.assertEqual(activity['status'],'applied');self.assertEqual(len(activity['moves']),1)

    def test_export_specific_month_no_data_and_snapshot_source(self):
        response=self.client.get('/api/export/monthly?month=2026-09');self.assertEqual(response.status_code,404)
        response=self.client.get('/api/export/monthly?month=2026-10');self.assertEqual(response.status_code,200)
        self.assertIn('Oct_2026.xlsx',response.headers['Content-Disposition'])
        snapshot=self.client.get('/api/live-sheets').json
        self.assertEqual(len(snapshot['sheets']['Overview']['rows']),1)
        self.assertEqual(snapshot['reports']['monthly'][0]['Month'],'2026-10')

    def test_archive_rollover_refused(self):
        response=self.client.post('/api/monthly-maintenance/preview',json={'kind':'rollover','month':'2026-09'})
        self.assertEqual(response.status_code,422)
        self.assertFalse(self.book.batches)

    def test_backup_restore_accepts_monthly_receipts(self):
        from workspace_backup import make_backup,restore_backup
        store=PreviewStore(self.folder.name)
        self.client.post('/api/monthly-maintenance/preview',json={'kind':'rollover','month':'2026-10'})
        backup=make_backup(store);restore_backup(store,backup.getvalue())
        self.assertEqual(len(self.client.get('/api/monthly-maintenance').json['operations']),1)

    def test_month_end_prepares_counts_without_moving_any_data(self):
        before=snapshot(self.book)
        class November(datetime):
            @classmethod
            def now(cls,tz=None):return cls(2026,11,1,0,1,tzinfo=tz)
        with patch('monthly_api.datetime',November):self.app.extensions['monthly_tick']()
        after=snapshot(self.book)
        for title in before:self.assertEqual(after[title],before[title])
        for base in BASES:self.assertEqual(decode(after[tab_name(base,'2026-11')]['values']),[])
        pending=[op for op in self.client.get('/api/monthly-maintenance').json['operations'] if op['status']=='preview']
        self.assertEqual(len(pending),1);self.assertTrue(pending[0]['automatic'])
        self.assertEqual(len(pending[0]['moves']),1)


if __name__=='__main__':unittest.main()
